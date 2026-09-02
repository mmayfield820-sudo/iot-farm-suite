from __future__ import annotations

import base64
import json
import hashlib
import hmac
import logging
import os
import re
import secrets
import threading
from io import BytesIO
from contextlib import asynccontextmanager
from datetime import UTC, date, datetime, time, timedelta
from typing import Any
from urllib.error import HTTPError, URLError
from urllib.parse import quote, urlencode
from urllib.request import Request, urlopen
from uuid import UUID

import bcrypt
from fastapi import Depends, FastAPI, File, HTTPException, Query, UploadFile, status
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from pydantic import BaseModel, Field
from sqlalchemy import bindparam, text
from sqlalchemy.exc import IntegrityError, SQLAlchemyError

from .auth import (
    authenticate,
    authenticate_initial_password,
    create_access_token,
    create_profile_setup_token,
    current_user,
    profile_setup_user_id,
)
from .db import check_database, engine, parse_id, utc_isoformat, utc_naive
from .irrigation_recommendations import build_irrigation_recommendation
from .moisture_alerts import current_moisture_rule, evaluate_moisture_alerts
from .mqtt import MqttService
from .plan_adjustment_reminders import (
    consecutive_rainy_days,
    purge_non_today_reminders,
    shanghai_day_utc_bounds,
    sync_weather_reminders,
)
from .weather import WeatherProviderError, open_meteo_forecast


mqtt_service = MqttService()
logger = logging.getLogger(__name__)
weather_reminder_stop = threading.Event()
weather_reminder_thread: threading.Thread | None = None

@asynccontextmanager
async def lifespan(_app: FastAPI):
    global weather_reminder_thread
    mqtt_service.start()
    weather_reminder_stop.clear()
    weather_reminder_thread = threading.Thread(
        target=weather_reminder_loop,
        name="weather-reminder-scheduler",
        daemon=True,
    )
    weather_reminder_thread.start()
    try:
        yield
    finally:
        weather_reminder_stop.set()
        if weather_reminder_thread is not None:
            weather_reminder_thread.join(timeout=2)
        mqtt_service.stop()

app = FastAPI(
    title="Smart Farm API",
    version="0.1.0",
    description="Smart agriculture backend service",
    lifespan=lifespan,
)


class LoginRequest(BaseModel):
    username: str
    password: str


class RegistrationRequest(BaseModel):
    username: str = Field(min_length=4, max_length=50, pattern=r"^[A-Za-z0-9_]+$")
    password: str = Field(min_length=6, max_length=72)
    real_name: str = Field(min_length=1, max_length=50)
    email: str | None = Field(default=None, max_length=100)
    phone: str | None = Field(default=None, max_length=20)
    role_codes: list[str] = Field(min_length=1)
    farm_ids: list[str] = Field(default_factory=list)
    plot_ids: list[str] = Field(default_factory=list)


class FirstLoginCompleteRequest(BaseModel):
    profileSetupToken: str = Field(min_length=1)
    password: str = Field(min_length=6, max_length=72)
    phone: str | None = Field(default=None, max_length=20)
    email: str | None = Field(default=None, max_length=100)


class UserCreateRequest(BaseModel):
    employeeNo: str = Field(min_length=1, max_length=50)
    idCardNo: str = Field(min_length=15, max_length=18)
    name: str = Field(min_length=1, max_length=50)
    email: str | None = Field(default=None, max_length=100)
    phone: str | None = Field(default=None, max_length=20)
    roleCodes: list[str] = Field(min_length=1)
    farmIds: list[str] = Field(default_factory=list)
    plotIds: list[str] = Field(default_factory=list)


class UserUpdateRequest(BaseModel):
    idCardNo: str | None = Field(default=None, min_length=15, max_length=18)
    name: str | None = Field(default=None, min_length=1, max_length=50)
    email: str | None = Field(default=None, max_length=100)
    phone: str | None = Field(default=None, max_length=20)
    roleCodes: list[str] | None = None
    farmIds: list[str] | None = None
    plotIds: list[str] | None = None


class ProfileUpdateRequest(BaseModel):
    name: str = Field(min_length=1, max_length=50)
    email: str | None = Field(default=None, max_length=100)
    phone: str | None = Field(default=None, max_length=20)
    avatarData: str | None = Field(default=None, max_length=1_500_000)
    currentPassword: str | None = Field(default=None, max_length=72)
    newPassword: str | None = Field(default=None, min_length=6, max_length=72)


class UserStatusUpdateRequest(BaseModel):
    status: str


class FarmSaveRequest(BaseModel):
    name: str = Field(min_length=1, max_length=100)
    province: str = Field(min_length=1, max_length=50)
    city: str = Field(min_length=1, max_length=50)
    address: str | None = Field(default=None, max_length=255)
    areaMu: float | None = Field(default=None, ge=0)
    managerName: str | None = Field(default=None, max_length=50)
    note: str | None = Field(default=None, max_length=500)


class FarmUpdateRequest(FarmSaveRequest):
    status: str | None = None


class PlotSaveRequest(BaseModel):
    farmId: str
    code: str = Field(min_length=1, max_length=50)
    name: str = Field(min_length=1, max_length=100)
    areaMu: float | None = Field(default=None, ge=0)
    location: str | None = Field(default=None, max_length=255)
    irrigationMethod: str | None = None
    status: str = "active"


class PlotUpdateRequest(BaseModel):
    farmId: str | None = None
    code: str | None = Field(default=None, min_length=1, max_length=50)
    name: str | None = Field(default=None, min_length=1, max_length=100)
    areaMu: float | None = Field(default=None, ge=0)
    location: str | None = Field(default=None, max_length=255)
    irrigationMethod: str | None = None
    status: str | None = None


class TemplateCreateRequest(BaseModel):
    name: str = Field(min_length=1, max_length=100)
    crop: str = Field(min_length=1, max_length=50)
    province: str | None = Field(default=None, max_length=50)
    city: str | None = Field(default=None, max_length=50)
    versionDefinition: dict[str, Any]


class TemplateVersionCreateRequest(BaseModel):
    versionDefinition: dict[str, Any]
    changeNote: str | None = Field(default=None, max_length=500)


class TemplateDraftCreateRequest(BaseModel):
    sourceVersionId: str | None = None
    changeNote: str | None = Field(default=None, max_length=500)


class TemplateDraftUpdateRequest(BaseModel):
    versionDefinition: dict[str, Any]
    changeNote: str | None = Field(default=None, max_length=500)


class TemplateDraftPublishRequest(BaseModel):
    productionYear: int = Field(ge=2000, le=2100)
    extendYoungTrack: bool = False
    trackOverrideReason: str | None = Field(default=None, max_length=500)


class FarmPlanGenerateRequest(BaseModel):
    stageCode: str = Field(min_length=1, max_length=50)
    regenerate: bool = False


class FarmStageStartRequest(BaseModel):
    stageCode: str = Field(min_length=1, max_length=50)
    actualStartedAt: datetime
    shiftFollowingStages: bool = True


class PlotProductionOffsetRequest(BaseModel):
    offsetDays: int = Field(ge=-60, le=60)


class TaskUpdateRequest(BaseModel):
    scheduledDate: str | None = None
    note: str | None = Field(default=None, max_length=1000)
    status: str | None = None


class TaskChangeOrderRequest(BaseModel):
    scheduledDate: str | None = None
    note: str | None = Field(default=None, max_length=1000)
    reason: str = Field(min_length=1, max_length=500)


class TemporaryTaskCreateRequest(BaseModel):
    plotId: str
    taskType: str = Field(min_length=1, max_length=50)
    title: str = Field(min_length=1, max_length=150)
    scheduledDate: str
    note: str | None = Field(default=None, max_length=1000)
    reason: str = Field(min_length=1, max_length=500)


class TaskCompletionRequest(BaseModel):
    result: str = Field(min_length=1, max_length=255)
    executedAt: datetime
    executorName: str = Field(min_length=1, max_length=50)
    note: str | None = Field(default=None, max_length=1000)
    irrigationRecordId: str | None = None


class TaskCompletionCorrectionRequest(TaskCompletionRequest):
    reason: str = Field(min_length=1, max_length=500)


class AdHocTaskRequest(BaseModel):
    taskType: str = Field(min_length=1, max_length=50)
    title: str = Field(min_length=1, max_length=150)
    executedAt: datetime
    executorName: str = Field(min_length=1, max_length=50)
    note: str | None = Field(default=None, max_length=1000)
    irrigationRecordId: str | None = None


class FarmProductionProfileRequest(BaseModel):
    templateVersionId: str
    plantingDate: str | None = None
    treeAgeYears: float | None = Field(default=None, ge=0)
    note: str | None = Field(default=None, max_length=1000)


class FarmMilestoneCompletionRequest(BaseModel):
    completedAt: datetime


class FarmAnnualProductionRequest(BaseModel):
    productionYear: int | None = Field(default=None, ge=2000, le=2100)
    templateVersionId: str | None = None
    orchardTrack: str | None = None
    versionDefinition: dict[str, Any] | None = None
    changeNote: str | None = Field(default=None, max_length=500)
    extendYoungTrack: bool = False
    trackOverrideReason: str | None = Field(default=None, max_length=500)


class FarmPlantingScheduleRequest(BaseModel):
    templateVersionId: str
    plannedPlantingDate: str


class UserSaveRequest(BaseModel):
    username: str = Field(min_length=4, max_length=50, pattern=r"^[A-Za-z0-9_]+$")
    name: str | None = Field(default=None, max_length=50)
    email: str | None = Field(default=None, max_length=100)
    phone: str | None = Field(default=None, max_length=20)
    active: bool = True
    role_codes: list[str] = Field(min_length=1)
    farm_ids: list[str] = Field(default_factory=list)
    plot_ids: list[str] = Field(default_factory=list)


class UserActiveRequest(BaseModel):
    active: bool


class RejectRegistrationRequest(BaseModel):
    reason: str = Field(min_length=1, max_length=255)


class CommandCreateRequest(BaseModel):
    action: str
    activationType: str = "manual"
    durationMinutes: int | None = Field(default=None, ge=1, le=1440)


class DeviceCreateRequest(BaseModel):
    deviceNo: str = Field(min_length=1, max_length=100)
    deviceUuid: str
    type: str
    plotId: str | None = None


class DeviceBindRequest(BaseModel):
    plotId: str


class AlertRuleRequest(BaseModel):
    plotId: str
    soilMoistureThreshold: float = Field(ge=0, le=100)
    irrigationLength: int = Field(default=0, ge=0)
    irrigationInterval: int = Field(default=0, ge=0)


class IrrigationSettingsRequest(BaseModel):
    soilMoistureThreshold: float = Field(ge=0, le=100)
    irrigationDuration: int = Field(ge=0, le=1440)
    irrigationInterval: int = Field(ge=0, le=10080)


class ChatMessageRequest(BaseModel):
    message: str = Field(min_length=1, max_length=4000)
    plotId: str | None = None


def business_id(prefix: str, value: int) -> str:
    return f"{prefix}{value:03d}"


def iso(value: datetime | None) -> str | None:
    return utc_isoformat(value)


def value_of(row: Any, key: str) -> Any:
    return row._mapping[key] if hasattr(row, "_mapping") else row[key]


ROLE_CODES = {"farmer", "farm_admin", "system_admin"}
FRONTEND_ROLE_CODES = {
    "farmer": "farmer",
    "farm_admin": "farmAdmin",
    "system_admin": "systemAdmin",
}


def canonical_role_codes(role_codes: list[str]) -> list[str]:
    aliases = {"farmAdmin": "farm_admin", "systemAdmin": "system_admin"}
    normalized = list(dict.fromkeys(aliases.get(role, role) for role in role_codes))
    if not normalized or any(role not in ROLE_CODES for role in normalized):
        raise HTTPException(status_code=400, detail="INVALID_ROLE_CODES")
    return normalized


def resource_ids(values: list[str], prefix: str) -> list[int]:
    return list(dict.fromkeys(parse_resource_id(value, prefix) for value in values))


def json_values(value: Any) -> list[str]:
    if isinstance(value, str):
        try:
            value = json.loads(value)
        except json.JSONDecodeError:
            return []
    return [str(item) for item in value] if isinstance(value, list) else []


def password_hash(password: str) -> str:
    return bcrypt.hashpw(password.encode("utf-8"), bcrypt.gensalt()).decode("utf-8")


def normalized_id_card(value: str) -> str:
    normalized = value.strip().upper()
    if re.fullmatch(r"\d{15}", normalized):
        return normalized
    if not re.fullmatch(r"\d{17}[\dX]", normalized):
        raise HTTPException(status_code=400, detail="INVALID_ID_CARD_NO")
    weights = (7, 9, 10, 5, 8, 4, 2, 1, 6, 3, 7, 9, 10, 5, 8, 4, 2)
    checks = "10X98765432"
    expected = checks[sum(int(item) * weight for item, weight in zip(normalized[:17], weights)) % 11]
    if normalized[-1] != expected:
        raise HTTPException(status_code=400, detail="INVALID_ID_CARD_NO")
    return normalized


def id_card_hash(value: str) -> str:
    secret = os.getenv("PII_HASH_SECRET") or os.getenv("JWT_SECRET")
    if not secret:
        raise RuntimeError("PII_HASH_SECRET or JWT_SECRET must be configured")
    return hmac.new(secret.encode("utf-8"), value.encode("utf-8"), hashlib.sha256).hexdigest()


def id_card_masked(value: str) -> str:
    return f"{value[:6]}{'*' * (len(value) - 10)}{value[-4:]}"


def ensure_system_admin(user: dict[str, Any]) -> None:
    if user["role"] != "system_admin":
        raise HTTPException(status_code=403, detail="FORBIDDEN")


def device_payload(row: Any) -> dict[str, Any]:
    return {
        "id": business_id("D", value_of(row, "device_id")),
        "deviceNo": value_of(row, "device_code"),
        "deviceUuid": value_of(row, "device_uuid"),
        "type": value_of(row, "device_type"),
        "plotId": (
            business_id("P", value_of(row, "plot_id"))
            if value_of(row, "plot_id") is not None
            else None
        ),
        "status": value_of(row, "status"),
        "switchState": value_of(row, "switch_state"),
        "switchChangedAt": iso(value_of(row, "switch_changed_at")),
        "lastSeenAt": iso(value_of(row, "last_heartbeat")),
    }


def farm_payload(row: Any) -> dict[str, Any]:
    return {"id": business_id("F", row["farm_id"]), "name": row["farm_name"], "province": row["province"], "city": row["city"], "address": row["address"], "areaMu": float(row["area_mu"]) if row["area_mu"] is not None else None, "managerName": row["manager_name"], "note": row["note"], "plantingDate": row.get("planting_date").isoformat() if value_of(row, "planting_date") else None, "plannedPlantingDate": row.get("planned_planting_date").isoformat() if value_of(row, "planned_planting_date") else None, "status": row["farm_status"], "createdAt": iso(row["created_at"]), "updatedAt": iso(row["updated_at"])}


def plot_payload(row: Any) -> dict[str, Any]:
    farmer_user_id = value_of(row, "scoped_farmer_user_id")
    return {"id": business_id("P", row["plot_id"]), "farmId": business_id("F", row["farm_id"]), "code": row["plot_code"], "name": row["plot_name"], "areaMu": float(row["area_mu"]) if row["area_mu"] is not None else None, "location": row["location"], "irrigationMethod": row["irrigation_method"], "productionDateOffsetDays": int(value_of(row, "production_date_offset_days") or 0), "status": row["plot_status"], "farmerUserId": business_id("U", farmer_user_id) if farmer_user_id is not None else None, "farmerName": value_of(row, "farmer_name"), "createdAt": iso(row["created_at"]), "updatedAt": iso(row["updated_at"])}


def accessible_farm_ids(user: dict[str, Any]) -> set[int] | None:
    if user["role"] == "system_admin":
        return None

    with engine.connect() as connection:
        if user["role"] == "farm_admin":
            rows = connection.execute(
                text(
                    "SELECT farm_id FROM farm WHERE admin_user_id = :user_id "
                    "UNION SELECT farm_id FROM farm_admin_scope WHERE user_id = :user_id"
                ),
                {"user_id": user["user_id"]},
            ).fetchall()
        else:
            rows = connection.execute(
                text(
                    "SELECT DISTINCT farm_id FROM plot WHERE farmer_user_id = :user_id "
                    "UNION SELECT DISTINCT p.farm_id FROM farmer_plot_scope AS s "
                    "JOIN plot AS p ON p.plot_id = s.plot_id WHERE s.user_id = :user_id"
                ),
                {"user_id": user["user_id"]},
            ).fetchall()
    return {row.farm_id for row in rows}


def ensure_farm_access(user: dict[str, Any], farm_id: int) -> None:
    allowed_farms = accessible_farm_ids(user)
    if allowed_farms is not None and farm_id not in allowed_farms:
        raise HTTPException(status_code=403, detail="FORBIDDEN")


def ensure_farm_admin_access(user: dict[str, Any], farm_id: int) -> None:
    if user["role"] != "farm_admin":
        raise HTTPException(status_code=403, detail="FORBIDDEN")
    ensure_farm_access(user, farm_id)


def record_audit_log(
    connection: Any,
    user: dict[str, Any],
    module: str,
    action: str,
    *,
    target_type: str | None = None,
    target_id: str | None = None,
    detail: str | None = None,
) -> None:
    """Write a successful key operation in the caller's transaction."""
    connection.execute(
        text(
            """
            INSERT INTO audit_log (
                actor_user_id, module_code, action_code, target_type, target_id, detail
            ) VALUES (
                :actor_user_id, :module_code, :action_code, :target_type, :target_id, :detail
            )
            """
        ),
        {
            "actor_user_id": user["user_id"],
            "module_code": module,
            "action_code": action,
            "target_type": target_type,
            "target_id": target_id,
            "detail": detail,
        },
    )


def parse_resource_id(value: str, prefix: str) -> int:
    try:
        return parse_id(value, prefix)
    except ValueError as error:
        raise HTTPException(status_code=400, detail=str(error)) from error


def normalized_uuid(value: str) -> str:
    try:
        return str(UUID(value))
    except (ValueError, AttributeError) as error:
        raise HTTPException(status_code=400, detail="deviceUuid must be a valid UUID") from error


def plot_scope(plot_id: int, user: dict[str, Any]) -> dict[str, Any]:
    with engine.connect() as connection:
        plot = connection.execute(
            text(
                """
                SELECT p.plot_id, p.farm_id, p.farmer_user_id, f.admin_user_id
                FROM plot AS p
                JOIN farm AS f ON f.farm_id = p.farm_id
                WHERE p.plot_id = :plot_id
                """
            ),
            {"plot_id": plot_id},
        ).mappings().first()
    if plot is None:
        raise HTTPException(status_code=404, detail="PLOT_NOT_FOUND")
    if user["role"] == "system_admin":
        return dict(plot)
    if user["role"] == "farm_admin" and plot["farm_id"] in accessible_farm_ids(user):
        return dict(plot)
    if user["role"] == "farmer":
        with engine.connect() as connection:
            allowed = connection.execute(
                text(
                    "SELECT 1 FROM farmer_plot_scope "
                    "WHERE user_id = :user_id AND plot_id = :plot_id"
                ),
                {"user_id": user["user_id"], "plot_id": plot_id},
            ).scalar()
        if plot["farmer_user_id"] == user["user_id"] or allowed:
            return dict(plot)
    raise HTTPException(status_code=403, detail="FORBIDDEN")


def ensure_plot_manage_access(plot_id: int, user: dict[str, Any]) -> dict[str, Any]:
    plot = plot_scope(plot_id, user)
    if user["role"] == "farmer":
        raise HTTPException(status_code=403, detail="FORBIDDEN")
    return plot


def assistant_plot_context(plot_id: int) -> dict[str, Any]:
    """Build the least-privilege plot snapshot that may leave the business API."""
    with engine.connect() as connection:
        plot = connection.execute(
            text(
                """
                SELECT p.plot_name, p.crop_name, p.location,
                       t.soil_moisture_threshold
                FROM plot AS p
                LEFT JOIN threshold_config AS t ON t.plot_id = p.plot_id
                WHERE p.plot_id = :plot_id
                """
            ),
            {"plot_id": plot_id},
        ).mappings().one()
        report = connection.execute(
            text(
                """
                SELECT plot_moisture, temperature, report_datetime
                FROM plot_report
                WHERE plot_id = :plot_id
                ORDER BY report_datetime DESC
                LIMIT 1
                """
            ),
            {"plot_id": plot_id},
        ).mappings().first()
        irrigation_device = connection.execute(
            text(
                """
                SELECT status, switch_state FROM device
                WHERE plot_id = :plot_id AND device_type = 'irrigation'
                LIMIT 1
                """
            ),
            {"plot_id": plot_id},
        ).mappings().first()
        last_command = connection.execute(
            text(
                """
                SELECT action_type, status
                FROM irrigation_record
                WHERE plot_id = :plot_id
                ORDER BY requested_at DESC, record_id DESC
                LIMIT 1
                """
            ),
            {"plot_id": plot_id},
        ).mappings().first()

    context: dict[str, Any] = {
        "plotName": plot["plot_name"],
        "cropName": plot["crop_name"],
        "location": plot["location"],
        "soilMoistureThreshold": (
            float(plot["soil_moisture_threshold"])
            if plot["soil_moisture_threshold"] is not None
            else None
        ),
        "irrigationDeviceOnline": (
            irrigation_device["status"] == "online"
            if irrigation_device is not None
            else False
        ),
    }
    if report is not None:
        context.update(
            {
                "soilMoisture": float(report["plot_moisture"]),
                "temperature": float(report["temperature"]),
                "dataUpdatedAt": iso(report["report_datetime"]),
            }
        )
    if irrigation_device is not None and irrigation_device["switch_state"] in {"on", "off"}:
        context["irrigationState"] = irrigation_device["switch_state"]
    elif last_command is not None and last_command["status"] == "success":
        context["irrigationState"] = last_command["action_type"]
    return context


def device_scope(device_id: int, user: dict[str, Any], *, allow_unbound_admin: bool = False) -> dict[str, Any]:
    with engine.connect() as connection:
        device = connection.execute(
            text(
                """
                SELECT d.device_id, d.device_code, d.device_uuid, d.device_type,
                       d.plot_id, d.status, d.switch_state, d.switch_changed_at, d.last_heartbeat,
                       p.farmer_user_id, f.admin_user_id
                FROM device AS d
                LEFT JOIN plot AS p ON p.plot_id = d.plot_id
                LEFT JOIN farm AS f ON f.farm_id = p.farm_id
                WHERE d.device_id = :device_id
                """
            ),
            {"device_id": device_id},
        ).mappings().first()
    if device is None:
        raise HTTPException(status_code=404, detail="DEVICE_NOT_FOUND")
    if user["role"] == "system_admin":
        return dict(device)
    if device["plot_id"] is None:
        if allow_unbound_admin and user["role"] == "farm_admin":
            return dict(device)
        raise HTTPException(status_code=403, detail="FORBIDDEN")
    if user["role"] == "farm_admin" and device["plot_id"] is not None:
        with engine.connect() as connection:
            allowed = connection.execute(
                text(
                    "SELECT 1 FROM farm_admin_scope AS s "
                    "JOIN plot AS p ON p.farm_id = s.farm_id "
                    "WHERE s.user_id = :user_id AND p.plot_id = :plot_id"
                ),
                {"user_id": user["user_id"], "plot_id": device["plot_id"]},
            ).scalar()
        if device["admin_user_id"] == user["user_id"] or allowed:
            return dict(device)
    if user["role"] == "farmer" and device["plot_id"] is not None:
        with engine.connect() as connection:
            allowed = connection.execute(
                text(
                    "SELECT 1 FROM farmer_plot_scope "
                    "WHERE user_id = :user_id AND plot_id = :plot_id"
                ),
                {"user_id": user["user_id"], "plot_id": device["plot_id"]},
            ).scalar()
        if device["farmer_user_id"] == user["user_id"] or allowed:
            return dict(device)
    raise HTTPException(status_code=403, detail="FORBIDDEN")


def ensure_device_manage_access(device_id: int, user: dict[str, Any], *, allow_unbound_admin: bool = False) -> dict[str, Any]:
    device = device_scope(device_id, user, allow_unbound_admin=allow_unbound_admin)
    if user["role"] == "farmer":
        raise HTTPException(status_code=403, detail="FORBIDDEN")
    return device


def command_payload(row: Any) -> dict[str, Any]:
    return {
        "commandId": value_of(row, "command_id"),
        "deviceId": business_id("D", value_of(row, "device_id")),
        "deviceUuid": value_of(row, "device_uuid"),
        "action": f"irrigation_{value_of(row, 'action_type')}",
        "activationType": value_of(row, "activation_type"),
        "status": value_of(row, "status"),
        "requestedAt": iso(value_of(row, "requested_at")),
        "executedAt": iso(value_of(row, "executed_at")),
        "durationMinutes": value_of(row, "planned_duration_minutes"),
        "scheduledStopAt": iso(value_of(row, "scheduled_stop_at")),
    }


def alert_payload(row: Any) -> dict[str, Any]:
    return {
        "id": business_id("A", value_of(row, "alert_id")),
        "plotId": business_id("P", value_of(row, "plot_id")),
        "deviceId": (
            business_id("D", value_of(row, "device_id"))
            if value_of(row, "device_id") is not None
            else None
        ),
        "type": value_of(row, "alert_type"),
        "description": value_of(row, "alert_description"),
        "value": float(value_of(row, "alert_value")) if value_of(row, "alert_value") is not None else None,
        "threshold": float(value_of(row, "alert_threshold")) if value_of(row, "alert_threshold") is not None else None,
        "status": value_of(row, "alert_status"),
        "triggeredAt": iso(value_of(row, "alert_datetime")),
        "handledAt": iso(value_of(row, "handled_at")),
    }


def plan_adjustment_reminder_payload(row: Any, tasks: list[Any]) -> dict[str, Any]:
    return {
        "id": business_id("PR", value_of(row, "reminder_id")),
        "plotId": business_id("P", value_of(row, "plot_id")),
        "plotName": value_of(row, "plot_name"),
        "type": value_of(row, "reminder_type"),
        "description": value_of(row, "reminder_description"),
        "triggeredAt": iso(value_of(row, "triggered_at")),
        "affectedTasks": [
            {
                "id": business_id("MT", value_of(task, "task_id")),
                "title": value_of(task, "task_title"),
                "type": value_of(task, "task_type"),
                "scheduledDate": value_of(task, "scheduled_date").isoformat(),
            }
            for task in tasks
        ],
    }


def evaluate_latest_moisture(
    connection: Any, *, plot_id: int, threshold: float
) -> None:
    latest_report = connection.execute(
        text(
            """
            SELECT plot_moisture, soil_moisture_device_id
            FROM plot_report
            WHERE plot_id = :plot_id
            ORDER BY report_datetime DESC, report_id DESC
            LIMIT 1
            """
        ),
        {"plot_id": plot_id},
    ).mappings().first()
    if latest_report is None:
        return
    evaluate_moisture_alerts(
        connection,
        plot_id=plot_id,
        device_id=value_of(latest_report, "soil_moisture_device_id"),
        moisture=float(value_of(latest_report, "plot_moisture")),
        fallback_threshold=threshold,
        timestamp=datetime.now(UTC).replace(tzinfo=None),
    )


def alert_rule_payload(row: Any) -> dict[str, Any]:
    return {
        "id": business_id("R", value_of(row, "config_id")),
        "plotId": business_id("P", value_of(row, "plot_id")),
        "soilMoistureThreshold": float(value_of(row, "soil_moisture_threshold")),
        "irrigationLength": value_of(row, "irrigation_length"),
        "irrigationInterval": value_of(row, "irrigation_interval"),
        "updatedAt": iso(value_of(row, "updated_at")),
    }


def irrigation_record_payload(row: Any) -> dict[str, Any]:
    return {
        "recordId": business_id("IR", value_of(row, "record_id")),
        "commandId": value_of(row, "command_id"),
        "plotId": business_id("P", value_of(row, "plot_id")),
        "deviceId": business_id("D", value_of(row, "device_id")),
        "action": f"irrigation_{value_of(row, 'action_type')}",
        "activationType": value_of(row, "activation_type"),
        "status": value_of(row, "status"),
        "requestedAt": iso(value_of(row, "requested_at")),
        "executedAt": iso(value_of(row, "executed_at")),
        "durationMinutes": value_of(row, "planned_duration_minutes"),
        "scheduledStopAt": iso(value_of(row, "scheduled_stop_at")),
        "operatorId": business_id("U", value_of(row, "operator_id")) if value_of(row, "operator_id") is not None else None,
        "operatorName": value_of(row, "operator_name") or "系统",
    }


def validate_avatar_data(value: str | None) -> str | None:
    if not value:
        return None
    match = re.fullmatch(r"data:image/(png|jpeg|webp);base64,([A-Za-z0-9+/]+={0,2})", value)
    if not match:
        raise HTTPException(status_code=400, detail="INVALID_AVATAR")
    try:
        image = base64.b64decode(match.group(2), validate=True)
    except ValueError as error:
        raise HTTPException(status_code=400, detail="INVALID_AVATAR") from error
    if len(image) > 1024 * 1024:
        raise HTTPException(status_code=400, detail="AVATAR_TOO_LARGE")
    signatures = {
        "png": image.startswith(b"\x89PNG\r\n\x1a\n"),
        "jpeg": image.startswith(b"\xff\xd8\xff"),
        "webp": len(image) >= 12 and image.startswith(b"RIFF") and image[8:12] == b"WEBP",
    }
    if not signatures[match.group(1)]:
        raise HTTPException(status_code=400, detail="INVALID_AVATAR")
    return value


def user_payload(connection: Any, user_id: int) -> dict[str, Any]:
    user = connection.execute(
        text(
            """
            SELECT user_id, employee_no, id_card_masked, username, real_name, email, phone, avatar_data,
                   user_status, activated_at, created_at, updated_at
            FROM users WHERE user_id = :user_id
            """
        ),
        {"user_id": user_id},
    ).mappings().one()
    roles = connection.execute(
        text(
            """
            SELECT r.role_code FROM user_role_rel AS ur
            JOIN roles AS r ON r.role_id = ur.role_id
            WHERE ur.user_id = :user_id AND ur.status = 'active'
            ORDER BY FIELD(r.role_code, 'system_admin', 'farm_admin', 'farmer')
            """
        ),
        {"user_id": user_id},
    ).scalars().all()
    farm_scopes = connection.execute(
        text(
            """
            SELECT farm.farm_id, farm.farm_name
            FROM farm_admin_scope AS scope
            JOIN farm ON farm.farm_id = scope.farm_id
            WHERE scope.user_id = :user_id
            ORDER BY farm.farm_id
            """
        ),
        {"user_id": user_id},
    ).mappings().all()
    plot_scopes = connection.execute(
        text(
            """
            SELECT plot.plot_id, plot.plot_name, farm.farm_id, farm.farm_name
            FROM farmer_plot_scope AS scope
            JOIN plot ON plot.plot_id = scope.plot_id
            JOIN farm ON farm.farm_id = plot.farm_id
            WHERE scope.user_id = :user_id
            ORDER BY plot.plot_id
            """
        ),
        {"user_id": user_id},
    ).mappings().all()
    return {
        "id": business_id("U", user["user_id"]),
        "employeeNo": user["employee_no"],
        "idCardMasked": user["id_card_masked"],
        "username": user["username"],
        "name": user["real_name"] or user["employee_no"],
        "displayName": user["real_name"] or user["employee_no"],
        "email": user["email"],
        "phone": user["phone"],
        "avatarData": user["avatar_data"],
        "active": user["user_status"] == "active",
        "status": user["user_status"],
        "roles": [FRONTEND_ROLE_CODES[role] for role in roles],
        "roleCodes": roles,
        "farmIds": [business_id("F", value["farm_id"]) for value in farm_scopes],
        "plotIds": [business_id("P", value["plot_id"]) for value in plot_scopes],
        "farmScopes": [
            {"id": business_id("F", value["farm_id"]), "name": value["farm_name"]}
            for value in farm_scopes
        ],
        "plotScopes": [
            {
                "id": business_id("P", value["plot_id"]),
                "name": value["plot_name"],
                "farmId": business_id("F", value["farm_id"]),
                "farmName": value["farm_name"],
            }
            for value in plot_scopes
        ],
        "activatedAt": iso(user["activated_at"]),
        "createdAt": iso(user["created_at"]),
        "updatedAt": iso(user["updated_at"]),
    }


def registration_payload(row: Any) -> dict[str, Any]:
    roles = json_values(value_of(row, "role_codes"))
    return {
        "id": business_id("REG", value_of(row, "application_id")),
        "applicationId": business_id("REG", value_of(row, "application_id")),
        "username": value_of(row, "username"),
        "realName": value_of(row, "real_name"),
        "email": value_of(row, "email"),
        "phone": value_of(row, "phone"),
        "roles": [FRONTEND_ROLE_CODES[role] for role in roles],
        "roleCodes": roles,
        "farmIds": [business_id("F", int(value)) for value in json_values(value_of(row, "farm_ids"))],
        "plotIds": [business_id("P", int(value)) for value in json_values(value_of(row, "plot_ids"))],
        "status": value_of(row, "application_status"),
        "createdAt": iso(value_of(row, "created_at")),
        "reviewedAt": iso(value_of(row, "reviewed_at")),
        "rejectReason": value_of(row, "reject_reason"),
    }


def replace_user_assignments(
    connection: Any,
    user_id: int,
    role_codes: list[str],
    farm_ids: list[int],
    plot_ids: list[int],
) -> None:
    connection.execute(text("DELETE FROM user_role_rel WHERE user_id = :user_id"), {"user_id": user_id})
    connection.execute(text("DELETE FROM farm_admin_scope WHERE user_id = :user_id"), {"user_id": user_id})
    connection.execute(text("DELETE FROM farmer_plot_scope WHERE user_id = :user_id"), {"user_id": user_id})
    connection.execute(
        text(
            """
            INSERT INTO user_role_rel (user_id, role_id, status)
            SELECT :user_id, role_id, 'active' FROM roles
            WHERE role_code IN :role_codes
            """
        ).bindparams(bindparam("role_codes", expanding=True)),
        {"user_id": user_id, "role_codes": role_codes},
    )
    if "farm_admin" in role_codes and farm_ids:
        connection.execute(
            text("INSERT INTO farm_admin_scope (user_id, farm_id) VALUES (:user_id, :farm_id)"),
            [{"user_id": user_id, "farm_id": farm_id} for farm_id in farm_ids],
        )
    if "farmer" in role_codes and plot_ids:
        connection.execute(
            text("INSERT INTO farmer_plot_scope (user_id, plot_id) VALUES (:user_id, :plot_id)"),
            [{"user_id": user_id, "plot_id": plot_id} for plot_id in plot_ids],
        )


def ensure_existing_resources(connection: Any, farm_ids: list[int], plot_ids: list[int]) -> None:
    if farm_ids:
        count = connection.execute(
            text("SELECT COUNT(*) FROM farm WHERE farm_id IN :farm_ids").bindparams(
                bindparam("farm_ids", expanding=True)
            ),
            {"farm_ids": farm_ids},
        ).scalar_one()
        if count != len(farm_ids):
            raise HTTPException(status_code=404, detail="FARM_NOT_FOUND")
    if plot_ids:
        count = connection.execute(
            text("SELECT COUNT(*) FROM plot WHERE plot_id IN :plot_ids").bindparams(
                bindparam("plot_ids", expanding=True)
            ),
            {"plot_ids": plot_ids},
        ).scalar_one()
        if count != len(plot_ids):
            raise HTTPException(status_code=404, detail="PLOT_NOT_FOUND")


def ensure_plots_available_for_farmer(
    connection: Any,
    plot_ids: list[int],
    *,
    excluding_user_id: int | None = None,
) -> None:
    if not plot_ids:
        return
    occupied_plot_ids = connection.execute(
        text(
            """
            SELECT p.plot_id
            FROM plot AS p
            WHERE p.plot_id IN :plot_ids
              AND (
                EXISTS (
                    SELECT 1 FROM farmer_plot_scope AS s
                    WHERE s.plot_id = p.plot_id
                      AND (:excluding_user_id IS NULL OR s.user_id <> :excluding_user_id)
                )
              )
            """
        ).bindparams(bindparam("plot_ids", expanding=True)),
        {"plot_ids": plot_ids, "excluding_user_id": excluding_user_id},
    ).scalars().all()
    if occupied_plot_ids:
        raise HTTPException(status_code=409, detail="PLOT_ALREADY_ASSIGNED")


@app.get("/health")
def health() -> dict[str, str]:
    try:
        check_database()
    except SQLAlchemyError:
        raise HTTPException(status_code=503, detail="database unavailable")
    return {"status": "ok", "database": "ok"}


@app.post("/api/auth/login")
def login(request: LoginRequest) -> dict[str, Any]:
    user = authenticate(request.username, request.password)
    if user is None:
        pending_user = authenticate_initial_password(request.username.strip(), request.password)
        if pending_user is None:
            raise HTTPException(status_code=401, detail="UNAUTHORIZED")
        with engine.connect() as connection:
            profile = user_payload(connection, pending_user["user_id"])
        return {
            "requiresProfileCompletion": True,
            "profileSetupToken": create_profile_setup_token(pending_user["user_id"]),
            "user": profile,
        }
    with engine.connect() as connection:
        profile = user_payload(connection, user["user_id"])
    return {
        "token": create_access_token(user),
        "user": {**profile, "role": user["role"]},
    }


@app.patch("/api/me")
def update_my_profile(request: ProfileUpdateRequest, user: dict[str, Any] = Depends(current_user)) -> dict[str, Any]:
    name = request.name.strip()
    if not name:
        raise HTTPException(status_code=400, detail="INVALID_REQUEST")
    email = request.email.strip() if request.email and request.email.strip() else None
    phone = request.phone.strip() if request.phone and request.phone.strip() else None
    avatar_data = validate_avatar_data(request.avatarData)
    if bool(request.currentPassword) != bool(request.newPassword):
        raise HTTPException(status_code=400, detail="PASSWORD_FIELDS_REQUIRED")
    with engine.begin() as connection:
        if request.newPassword:
            password = connection.execute(
                text("SELECT password_hash FROM users WHERE user_id=:user"),
                {"user": user["user_id"]},
            ).scalar_one()
            if not bcrypt.checkpw(request.currentPassword.encode("utf-8"), password.encode("utf-8")):
                raise HTTPException(status_code=400, detail="CURRENT_PASSWORD_INVALID")
        connection.execute(text("UPDATE users SET real_name=:name,email=:email,phone=:phone,avatar_data=:avatar_data WHERE user_id=:user"), {"name": name, "email": email, "phone": phone, "avatar_data": avatar_data, "user": user["user_id"]})
        if request.newPassword:
            connection.execute(
                text("UPDATE users SET password_hash=:password_hash WHERE user_id=:user"),
                {"password_hash": password_hash(request.newPassword), "user": user["user_id"]},
            )
        record_audit_log(connection, user, "account_access", "update_own_profile", target_type="user", target_id=business_id("U", user["user_id"]), detail="Updated own profile")
        return user_payload(connection, user["user_id"])


@app.get("/api/me")
def get_my_profile(user: dict[str, Any] = Depends(current_user)) -> dict[str, Any]:
    with engine.connect() as connection:
        return user_payload(connection, user["user_id"])


@app.get("/api/registration-resources", include_in_schema=False)
def registration_resources() -> None:
    raise HTTPException(status_code=410, detail="REGISTRATION_DISABLED")


@app.post("/api/auth/first-login/complete")
def complete_first_login(request: FirstLoginCompleteRequest) -> dict[str, Any]:
    now = datetime.now(UTC).replace(tzinfo=None)
    user_id = profile_setup_user_id(request.profileSetupToken)
    with engine.begin() as connection:
        employee = connection.execute(
            text(
                """
                SELECT user_id, password_hash
                FROM users
                WHERE user_id = :user_id AND user_status = 'pending_activation'
                FOR UPDATE
                """
            ),
            {"user_id": user_id},
        ).mappings().first()
        if employee is None:
            raise HTTPException(status_code=400, detail="FIRST_LOGIN_INVALID")
        if bcrypt.checkpw(request.password.encode("utf-8"), employee["password_hash"].encode("utf-8")):
            raise HTTPException(status_code=400, detail="PASSWORD_MUST_CHANGE")
        try:
            connection.execute(
                text(
                    """
                    UPDATE users SET password_hash = :password_hash,
                        phone = :phone, email = :email, user_status = 'active',
                        activated_at = :activated_at
                    WHERE user_id = :user_id
                    """
                ),
                {
                    "password_hash": password_hash(request.password),
                    "phone": request.phone or None,
                    "email": request.email or None,
                    "activated_at": now,
                    "user_id": employee["user_id"],
                },
            )
        except IntegrityError as error:
            raise HTTPException(status_code=409, detail="USERNAME_OR_EMAIL_ALREADY_EXISTS") from error
        return user_payload(connection, employee["user_id"])


@app.get("/api/registration-applications")
def list_registration_applications(user: dict[str, Any] = Depends(current_user)) -> list[dict[str, Any]]:
    ensure_system_admin(user)
    with engine.connect() as connection:
        rows = connection.execute(
            text("SELECT * FROM registration_application ORDER BY created_at DESC, application_id DESC")
        ).mappings().all()
    return [registration_payload(row) for row in rows]


@app.put("/api/registration-applications/{application_id}/approve")
def approve_registration(application_id: str, user: dict[str, Any] = Depends(current_user)) -> dict[str, Any]:
    ensure_system_admin(user)
    raise HTTPException(status_code=410, detail="REGISTRATION_DISABLED")
    numeric_id = parse_resource_id(application_id, "REG")
    with engine.begin() as connection:
        application = connection.execute(
            text("SELECT * FROM registration_application WHERE application_id = :application_id FOR UPDATE"),
            {"application_id": numeric_id},
        ).mappings().first()
        if application is None:
            raise HTTPException(status_code=404, detail="REGISTRATION_NOT_FOUND")
        if application["application_status"] != "pending":
            raise HTTPException(status_code=409, detail="REGISTRATION_NOT_PENDING")
        if connection.execute(
            text("SELECT 1 FROM users WHERE username = :username"),
            {"username": application["username"]},
        ).scalar():
            raise HTTPException(status_code=409, detail="USERNAME_ALREADY_EXISTS")
        result = connection.execute(
            text(
                """
                INSERT INTO users (username, password_hash, real_name, email, phone)
                VALUES (:username, :password_hash, :real_name, :email, :phone)
                """
            ),
            dict(application),
        )
        role_codes = json_values(application["role_codes"])
        plot_ids = [int(value) for value in json_values(application["plot_ids"])]
        if "farmer" in role_codes:
            ensure_plots_available_for_farmer(connection, plot_ids)
        replace_user_assignments(
            connection,
            result.lastrowid,
            role_codes,
            [int(value) for value in json_values(application["farm_ids"])],
            plot_ids,
        )
        connection.execute(
            text(
                """
                UPDATE registration_application
                SET application_status = 'approved', reviewed_at = :reviewed_at,
                    reviewed_by = :reviewed_by
                WHERE application_id = :application_id
                """
            ),
            {"application_id": numeric_id, "reviewed_at": datetime.now(UTC).replace(tzinfo=None), "reviewed_by": user["user_id"]},
        )
        row = connection.execute(
            text("SELECT * FROM registration_application WHERE application_id = :application_id"),
            {"application_id": numeric_id},
        ).mappings().one()
    return registration_payload(row)


@app.put("/api/registration-applications/{application_id}/reject")
def reject_registration(
    application_id: str,
    request: RejectRegistrationRequest,
    user: dict[str, Any] = Depends(current_user),
) -> dict[str, Any]:
    ensure_system_admin(user)
    raise HTTPException(status_code=410, detail="REGISTRATION_DISABLED")
    numeric_id = parse_resource_id(application_id, "REG")
    with engine.begin() as connection:
        result = connection.execute(
            text(
                """
                UPDATE registration_application
                SET application_status = 'rejected', reject_reason = :reason,
                    reviewed_at = :reviewed_at, reviewed_by = :reviewed_by
                WHERE application_id = :application_id AND application_status = 'pending'
                """
            ),
            {"application_id": numeric_id, "reason": request.reason.strip(), "reviewed_at": datetime.now(UTC).replace(tzinfo=None), "reviewed_by": user["user_id"]},
        )
        if result.rowcount == 0:
            raise HTTPException(status_code=409, detail="REGISTRATION_NOT_PENDING")
        row = connection.execute(
            text("SELECT * FROM registration_application WHERE application_id = :application_id"),
            {"application_id": numeric_id},
        ).mappings().one()
    return registration_payload(row)


@app.get("/api/_legacy/users", include_in_schema=False)
def list_users(user: dict[str, Any] = Depends(current_user)) -> list[dict[str, Any]]:
    ensure_system_admin(user)
    with engine.connect() as connection:
        user_ids = connection.execute(text("SELECT user_id FROM users ORDER BY user_id")).scalars().all()
        return [user_payload(connection, user_id) for user_id in user_ids]


def save_user_record(
    request: UserSaveRequest,
    user_id: int | None,
) -> tuple[dict[str, Any], str | None]:
    role_codes = canonical_role_codes(request.role_codes)
    farm_ids = resource_ids(request.farm_ids, "F")
    plot_ids = resource_ids(request.plot_ids, "P")
    if "farm_admin" in role_codes and not farm_ids:
        raise HTTPException(status_code=400, detail="FARM_SCOPE_REQUIRED")
    if "farmer" in role_codes and not plot_ids:
        raise HTTPException(status_code=400, detail="PLOT_SCOPE_REQUIRED")
    temporary_password: str | None = None
    with engine.begin() as connection:
        ensure_existing_resources(connection, farm_ids, plot_ids)
        if user_id is None:
            temporary_password = f"Tmp-{secrets.token_urlsafe(9)}"
            try:
                result = connection.execute(
                    text(
                        """
                        INSERT INTO users (username, password_hash, real_name, email, phone, user_status)
                        VALUES (:username, :password_hash, :real_name, :email, :phone, :user_status)
                        """
                    ),
                    {
                        "username": request.username,
                        "password_hash": password_hash(temporary_password),
                        "real_name": request.name or request.username,
                        "email": request.email or None,
                        "phone": request.phone or None,
                        "user_status": "active" if request.active else "disabled",
                    },
                )
            except IntegrityError as error:
                raise HTTPException(status_code=409, detail="USERNAME_OR_EMAIL_ALREADY_EXISTS") from error
            numeric_user_id = result.lastrowid
        else:
            existing = connection.execute(
                text("SELECT 1 FROM users WHERE user_id = :user_id"), {"user_id": user_id}
            ).scalar()
            if not existing:
                raise HTTPException(status_code=404, detail="USER_NOT_FOUND")
            try:
                connection.execute(
                    text(
                        """
                        UPDATE users
                        SET real_name = :real_name, email = :email, phone = :phone,
                            user_status = :user_status
                        WHERE user_id = :user_id
                        """
                    ),
                    {
                        "user_id": user_id,
                        "real_name": request.name or request.username,
                        "email": request.email or None,
                        "phone": request.phone or None,
                        "user_status": "active" if request.active else "disabled",
                    },
                )
            except IntegrityError as error:
                raise HTTPException(status_code=409, detail="EMAIL_ALREADY_EXISTS") from error
            numeric_user_id = user_id
        replace_user_assignments(connection, numeric_user_id, role_codes, farm_ids, plot_ids)
        return user_payload(connection, numeric_user_id), temporary_password


@app.post("/api/_legacy/users", status_code=status.HTTP_201_CREATED, include_in_schema=False)
def create_user(
    request: UserSaveRequest,
    user: dict[str, Any] = Depends(current_user),
) -> dict[str, Any]:
    ensure_system_admin(user)
    saved, temporary_password = save_user_record(request, None)
    return {**saved, "temporaryPassword": temporary_password}


@app.put("/api/_legacy/users/{user_id}", include_in_schema=False)
def update_user(
    user_id: str,
    request: UserSaveRequest,
    user: dict[str, Any] = Depends(current_user),
) -> dict[str, Any]:
    ensure_system_admin(user)
    numeric_user_id = parse_resource_id(user_id, "U")
    saved, _ = save_user_record(request, numeric_user_id)
    return saved


@app.put("/api/_legacy/users/{user_id}/active", include_in_schema=False)
def update_user_active(
    user_id: str,
    request: UserActiveRequest,
    user: dict[str, Any] = Depends(current_user),
) -> dict[str, Any]:
    ensure_system_admin(user)
    numeric_user_id = parse_resource_id(user_id, "U")
    if numeric_user_id == user["user_id"] and not request.active:
        raise HTTPException(status_code=400, detail="CANNOT_DISABLE_CURRENT_USER")
    with engine.begin() as connection:
        result = connection.execute(
            text("UPDATE users SET user_status = :status WHERE user_id = :user_id"),
            {"user_id": numeric_user_id, "status": "active" if request.active else "disabled"},
        )
        if result.rowcount == 0:
            raise HTTPException(status_code=404, detail="USER_NOT_FOUND")
        return user_payload(connection, numeric_user_id)


@app.post("/api/_legacy/users/{user_id}/reset-password", include_in_schema=False)
def reset_user_password(
    user_id: str,
    user: dict[str, Any] = Depends(current_user),
) -> dict[str, str]:
    ensure_system_admin(user)
    numeric_user_id = parse_resource_id(user_id, "U")
    temporary_password = f"Tmp-{secrets.token_urlsafe(9)}"
    with engine.begin() as connection:
        result = connection.execute(
            text("UPDATE users SET password_hash = :password_hash WHERE user_id = :user_id"),
            {"user_id": numeric_user_id, "password_hash": password_hash(temporary_password)},
        )
        if result.rowcount == 0:
            raise HTTPException(status_code=404, detail="USER_NOT_FOUND")
    return {"temporaryPassword": temporary_password}


def validated_assignments(
    connection: Any, role_codes: list[str], farm_values: list[str], plot_values: list[str], *, user_id: int | None = None
) -> tuple[list[str], list[int], list[int]]:
    roles = canonical_role_codes(role_codes)
    farm_ids = resource_ids(farm_values, "F")
    plot_ids = resource_ids(plot_values, "P")
    if "farm_admin" in roles and not farm_ids:
        raise HTTPException(status_code=400, detail="FARM_SCOPE_REQUIRED")
    if "farmer" in roles and not plot_ids:
        raise HTTPException(status_code=400, detail="PLOT_SCOPE_REQUIRED")
    ensure_existing_resources(connection, farm_ids, plot_ids)
    if "farmer" in roles:
        ensure_plots_available_for_farmer(connection, plot_ids, excluding_user_id=user_id)
    return roles, farm_ids, plot_ids


IMPORT_HEADERS = ["工号", "姓名", "身份证号", "手机号", "邮箱", "角色", "授权橙园编号", "授权地块编号"]


def import_values(value: Any) -> list[str]:
    return [part.strip() for part in re.split(r"[,，;；\s]+", str(value or "")) if part.strip()]


def import_job_payload(connection: Any, job_id: int) -> dict[str, Any]:
    job = connection.execute(text("SELECT * FROM user_import_job WHERE import_job_id=:id"), {"id": job_id}).mappings().one()
    rows = connection.execute(text("SELECT source_row_number, validation_errors FROM user_import_row WHERE import_job_id=:id AND row_status='invalid' ORDER BY source_row_number"), {"id": job_id}).mappings().all()
    errors = [{"rowNumber": row["source_row_number"], **error} for row in rows for error in (json.loads(row["validation_errors"]) if row["validation_errors"] else [])]
    return {"id": business_id("IMP", job_id), "status": job["job_status"], "totalRows": job["total_rows"], "validRows": job["valid_rows"], "invalidRows": job["invalid_rows"], "errors": errors, "createdAt": iso(job["created_at"]), "committedAt": iso(job["committed_at"])}


def validate_template_definition(definition: dict[str, Any]) -> list[dict[str, Any]]:
    stages = definition.get("stages")
    if not isinstance(stages, list) or not stages:
        raise HTTPException(status_code=400, detail="INVALID_TEMPLATE_DEFINITION")

    stage_codes: set[str] = set()
    total_duration_days = 0
    for stage in stages:
        try:
            code = str(stage["code"]).strip()
            name = str(stage["name"]).strip()
            duration = int(stage["durationDays"])
            tasks = stage.get("tasks", [])
            minimum = stage.get("moistureMin")
            maximum = stage.get("moistureMax")
        except (AttributeError, KeyError, TypeError, ValueError) as error:
            raise HTTPException(status_code=400, detail="INVALID_TEMPLATE_DEFINITION") from error
        if not code or not name or len(code) > 50 or len(name) > 100 or duration <= 0 or code in stage_codes or not isinstance(tasks, list):
            raise HTTPException(status_code=400, detail="INVALID_TEMPLATE_DEFINITION")
        if minimum is not None and maximum is not None and float(minimum) > float(maximum):
            raise HTTPException(status_code=400, detail="INVALID_TEMPLATE_DEFINITION")
        stage_codes.add(code)
        total_duration_days += duration
        for task in tasks:
            try:
                title = str(task.get("title", "任务")).strip()
                task_type = str(task.get("type", "other")).strip()
                offset = int(task.get("offsetDays", 0))
                repeat_days = int(task["repeatDays"]) if task.get("repeatDays") is not None else None
                repeat_until_stage_end = bool(task.get("repeatUntilStageEnd", False))
            except (AttributeError, TypeError, ValueError) as error:
                raise HTTPException(status_code=400, detail="INVALID_TEMPLATE_DEFINITION") from error
            if not title or not task_type or len(title) > 150 or len(task_type) > 50 or offset < 0 or offset >= duration:
                raise HTTPException(status_code=400, detail="INVALID_TEMPLATE_DEFINITION")
            if repeat_days is not None and (repeat_days <= 0 or repeat_days > 366):
                raise HTTPException(status_code=400, detail="INVALID_TEMPLATE_DEFINITION")
            if repeat_until_stage_end and repeat_days is None:
                raise HTTPException(status_code=400, detail="INVALID_TEMPLATE_DEFINITION")
            if task.get("track") not in {None, "young", "fruiting"}:
                raise HTTPException(status_code=400, detail="INVALID_TEMPLATE_DEFINITION")
    if total_duration_days != 365:
        raise HTTPException(status_code=400, detail="ANNUAL_STAGE_DURATION_MUST_TOTAL_365")
    return stages


def store_template_definition(connection: Any, version_id: int, definition: dict[str, Any]) -> None:
    stages = validate_template_definition(definition)
    for order, stage in enumerate(stages, start=1):
        result = connection.execute(text("INSERT INTO growth_stage_template (template_version_id,stage_code,stage_name,sort_order,duration_days,moisture_min,moisture_max,management_focus) VALUES (:version,:code,:name,:order,:duration,:minimum,:maximum,:focus)"), {"version":version_id,"code":str(stage["code"]).strip(),"name":str(stage["name"]).strip(),"order":order,"duration":int(stage["durationDays"]),"minimum":stage.get("moistureMin"),"maximum":stage.get("moistureMax"),"focus":stage.get("managementFocus")})
        for task in stage.get("tasks", []):
            connection.execute(text("INSERT INTO management_task_template (stage_template_id,task_type,title,offset_days,suggested_time,note,track,repeat_days,repeat_until_stage_end) VALUES (:stage,:type,:title,:offset,:suggested_time,:note,:track,:repeat_days,:repeat_until_stage_end)"), {"stage":result.lastrowid,"type":str(task.get("type", "other")).strip(),"title":str(task.get("title", "任务")).strip(),"offset":int(task.get("offsetDays", 0)),"suggested_time":task.get("suggestedTime"),"note":task.get("note"),"track":task.get("track"),"repeat_days":task.get("repeatDays"),"repeat_until_stage_end":bool(task.get("repeatUntilStageEnd", False))})


def template_definition(connection: Any, version_id: int) -> dict[str, Any]:
    stages = connection.execute(text("SELECT stage_template_id,stage_code,stage_name,duration_days,moisture_min,moisture_max,management_focus FROM growth_stage_template WHERE template_version_id=:version ORDER BY sort_order"), {"version": version_id}).mappings().all()
    definition: list[dict[str, Any]] = []
    for stage in stages:
        tasks = connection.execute(text("SELECT task_type,title,offset_days,suggested_time,note,track,repeat_days,repeat_until_stage_end FROM management_task_template WHERE stage_template_id=:stage ORDER BY task_template_id"), {"stage": stage["stage_template_id"]}).mappings().all()
        definition.append({"code": stage["stage_code"], "name": stage["stage_name"], "durationDays": stage["duration_days"], "moistureMin": float(stage["moisture_min"]) if stage["moisture_min"] is not None else None, "moistureMax": float(stage["moisture_max"]) if stage["moisture_max"] is not None else None, "managementFocus": stage["management_focus"], "tasks": [{"type": task["task_type"], "title": task["title"], "offsetDays": task["offset_days"], "suggestedTime": str(task["suggested_time"]) if task["suggested_time"] is not None else None, "note": task["note"], "track": task["track"], "repeatDays": task["repeat_days"], "repeatUntilStageEnd": bool(task["repeat_until_stage_end"])} for task in tasks]})
    return {"stages": definition}


def template_draft_payload(connection: Any, version: Any) -> dict[str, Any]:
    template = connection.execute(text("SELECT template_name,crop_name FROM crop_template WHERE template_id=:template"), {"template": version["template_id"]}).mappings().one()
    return {"id": business_id("CTV", version["template_version_id"]), "templateId": business_id("CT", version["template_id"]), "farmId": business_id("F", version["farm_id"]), "templateName": template["template_name"], "crop": template["crop_name"], "versionNo": version["version_no"], "status": version["version_status"], "changeNote": version["change_note"], "versionDefinition": template_definition(connection, version["template_version_id"])}


def current_farm_template_profile(connection: Any, farm_id: int, template_id: int, *, lock: bool = False) -> Any:
    current_year = production_year_for(date.today())
    statement = "SELECT p.*,COALESCE(a.template_version_id,p.template_version_id) AS effective_template_version_id,v.template_id FROM farm_production_profile p LEFT JOIN farm_annual_production a ON a.farm_id=p.farm_id AND a.production_year=:year JOIN crop_template_version v ON v.template_version_id=COALESCE(a.template_version_id,p.template_version_id) WHERE p.farm_id=:farm"
    if lock:
        statement += " FOR UPDATE"
    profile = connection.execute(text(statement), {"farm": farm_id, "year": current_year}).mappings().first()
    if profile is None or profile["template_id"] != template_id:
        raise HTTPException(status_code=409, detail="PRODUCTION_TEMPLATE_REQUIRED")
    return profile


def initial_template_customization_profile(connection: Any, farm_id: int, template_id: int, *, lock: bool = False) -> Any | None:
    current_year = production_year_for(date.today())
    statement = "SELECT p.*,v.template_id FROM farm_production_profile p JOIN crop_template_version v ON v.template_version_id=p.template_version_id WHERE p.farm_id=:farm"
    if lock:
        statement += " FOR UPDATE"
    profile = connection.execute(text(statement), {"farm": farm_id}).mappings().first()
    if profile is None or profile["template_id"] != template_id:
        return None
    annual_exists = connection.execute(text("SELECT 1 FROM farm_annual_production WHERE farm_id=:farm AND production_year=:year"), {"farm": farm_id, "year": current_year}).scalar()
    plan_exists = connection.execute(text("SELECT 1 FROM management_plan mp JOIN plot p ON p.plot_id=mp.plot_id WHERE p.farm_id=:farm LIMIT 1"), {"farm": farm_id}).scalar()
    return profile if not annual_exists and not plan_exists else None


def editable_farm_template_profile(connection: Any, farm_id: int, template_id: int, *, lock: bool = False) -> tuple[Any, str]:
    initial_profile = initial_template_customization_profile(connection, farm_id, template_id, lock=lock)
    if initial_profile is not None:
        return initial_profile, "initial"
    return current_farm_template_profile(connection, farm_id, template_id, lock=lock), "next"


def production_year_for(day: date) -> int:
    return day.year + 1 if day.month == 12 else day.year


def production_year_bounds(production_year: int) -> tuple[date, date]:
    return date(production_year - 1, 12, 1), date(production_year, 11, 30)


def annual_stage_window(stage_code: str, production_year: int, fallback_start: date, duration_days: int) -> tuple[date, date]:
    del stage_code, production_year
    return fallback_start, fallback_start + timedelta(days=duration_days - 1)


def template_stage_windows(connection: Any, template_version_id: int, production_year: int) -> list[dict[str, Any]]:
    stages = connection.execute(text("SELECT stage_code,stage_name,sort_order,duration_days,moisture_min,moisture_max,management_focus FROM growth_stage_template WHERE template_version_id=:version ORDER BY sort_order"), {"version": template_version_id}).mappings().all()
    fallback_start, production_year_end = production_year_bounds(production_year)
    result: list[dict[str, Any]] = []
    for index, stage in enumerate(stages):
        starts_on, ends_on = annual_stage_window(stage["stage_code"], production_year, fallback_start, stage["duration_days"])
        if index == len(stages) - 1 and ends_on < production_year_end:
            ends_on = production_year_end
        if ends_on > production_year_end:
            raise HTTPException(status_code=409, detail="ANNUAL_STAGE_WINDOW_OUT_OF_RANGE")
        fallback_start = ends_on + timedelta(days=1)
        result.append({**dict(stage), "planned_start_date": starts_on, "planned_end_date": ends_on})
    return result


def ensure_annual_stage_states(connection: Any, farm_id: int, production_year: int, template_version_id: int) -> None:
    for stage in template_stage_windows(connection, template_version_id, production_year):
        connection.execute(text("INSERT IGNORE INTO farm_annual_stage_state (farm_id,production_year,stage_code,sort_order,planned_start_date,planned_end_date) VALUES (:farm,:year,:code,:order,:starts,:ends)"), {"farm": farm_id, "year": production_year, "code": stage["stage_code"], "order": stage["sort_order"], "starts": stage["planned_start_date"], "ends": stage["planned_end_date"]})


def refresh_unlocked_stage_windows(connection: Any, farm_id: int, production_year: int, template_version_id: int) -> None:
    """Apply edited durations to future stages while preserving locked history."""
    windows = template_stage_windows(connection, template_version_id, production_year)
    states = {state["stage_code"]: state for state in connection.execute(text("SELECT * FROM farm_annual_stage_state WHERE farm_id=:farm AND production_year=:year FOR UPDATE"), {"farm": farm_id, "year": production_year}).mappings().all()}
    cursor, production_end = production_year_bounds(production_year)
    for index, stage in enumerate(windows):
        state = states.get(stage["stage_code"])
        if state is None:
            continue
        if state["stage_status"] in {"active", "completed", "not_dispatched"}:
            cursor = state["planned_end_date"] + timedelta(days=1)
            continue
        starts_on = cursor
        ends_on = starts_on + timedelta(days=stage["duration_days"] - 1)
        if index == len(windows) - 1 and ends_on < production_end:
            ends_on = production_end
        if ends_on > production_end:
            raise HTTPException(status_code=409, detail="ANNUAL_STAGE_WINDOW_OUT_OF_RANGE")
        connection.execute(text("UPDATE farm_annual_stage_state SET planned_start_date=:starts,planned_end_date=:ends WHERE farm_id=:farm AND production_year=:year AND stage_code=:stage"), {"starts": starts_on, "ends": ends_on, "farm": farm_id, "year": production_year, "stage": stage["stage_code"]})
        cursor = ends_on + timedelta(days=1)


def locked_stage_codes(connection: Any, farm_id: int, production_year: int) -> set[str]:
    return set(connection.execute(text("SELECT stage_code FROM farm_annual_stage_state WHERE farm_id=:farm AND production_year=:year AND stage_status IN ('active','completed')"), {"farm": farm_id, "year": production_year}).scalars().all())


def ensure_locked_stages_unchanged(connection: Any, farm_id: int, production_year: int, current_version_id: int, new_definition: dict[str, Any]) -> None:
    locked = locked_stage_codes(connection, farm_id, production_year)
    if not locked:
        return
    current_by_code = {stage["code"]: stage for stage in template_definition(connection, current_version_id)["stages"]}
    incoming_by_code = {str(stage.get("code", "")): stage for stage in new_definition.get("stages", [])}
    for code in locked:
        current_stage = current_by_code.get(code)
        incoming_stage = incoming_by_code.get(code)
        if current_stage is None or incoming_stage is None or json.dumps(current_stage, sort_keys=True, ensure_ascii=False, default=str) != json.dumps(incoming_stage, sort_keys=True, ensure_ascii=False, default=str):
            raise HTTPException(status_code=409, detail={"code": "LOCKED_STAGE_MODIFICATION", "stageCode": code})


def orchard_track(planting_date: date, production_year: int) -> tuple[int, str]:
    tree_age_years = max(1, production_year - planting_date.year + 1)
    return tree_age_years, "young" if tree_age_years <= 3 else "fruiting"


def recurring_task_dates(first_date: date, stage_end: date, repeat_days: int | None, repeat_until_stage_end: bool) -> list[date]:
    if first_date > stage_end:
        return []
    if not repeat_until_stage_end:
        return [first_date]
    if repeat_days is None or repeat_days <= 0:
        raise ValueError("a recurring task requires a positive repeat interval")
    occurrences: list[date] = []
    scheduled_date = first_date
    while scheduled_date <= stage_end:
        occurrences.append(scheduled_date)
        scheduled_date += timedelta(days=repeat_days)
    return occurrences


def farm_plan_task_rules(connection: Any, template_version_id: int, production_year: int, track: str, window_overrides: dict[str, tuple[date, date]] | None = None) -> list[dict[str, Any]]:
    rows = connection.execute(text("SELECT s.stage_code,s.sort_order,s.duration_days,t.task_type,t.title,t.offset_days,t.note,t.track,t.repeat_days,t.repeat_until_stage_end FROM growth_stage_template s LEFT JOIN management_task_template t ON t.stage_template_id=s.stage_template_id WHERE s.template_version_id=:version ORDER BY s.sort_order,t.offset_days,t.task_template_id"), {"version": template_version_id}).mappings().all()
    rules: list[dict[str, Any]] = []
    stage_windows = {stage["stage_code"]: (stage["planned_start_date"], stage["planned_end_date"]) for stage in template_stage_windows(connection, template_version_id, production_year)}
    for row in rows:
        if row["task_type"] is None:
            continue
        if row["track"] not in {None, track}:
            continue
        stage_start, stage_end = window_overrides.get(row["stage_code"], stage_windows[row["stage_code"]]) if window_overrides else stage_windows[row["stage_code"]]
        first_date = stage_start + timedelta(days=row["offset_days"])
        for scheduled_date in recurring_task_dates(first_date, stage_end, row["repeat_days"], bool(row["repeat_until_stage_end"])):
            rules.append({"stageCode": row["stage_code"], "taskType": row["task_type"], "title": row["title"], "scheduledDate": scheduled_date, "note": row["note"]})
    return rules


def farm_production_context(connection: Any, plot: dict[str, Any]) -> dict[str, Any]:
    farm = connection.execute(text("SELECT planting_date,planned_planting_date,planting_template_version_id FROM farm WHERE farm_id=:farm"), {"farm": plot["farm_id"]}).mappings().one()
    if farm["planting_date"] is None:
        template = None
        if farm["planting_template_version_id"] is not None:
            template = connection.execute(text("SELECT t.template_name,t.crop_name,v.version_no FROM crop_template_version v JOIN crop_template t ON t.template_id=v.template_id WHERE v.template_version_id=:version"), {"version": farm["planting_template_version_id"]}).mappings().first()
        return {"plotId": business_id("P", plot["plot_id"]), "farmId": business_id("F", plot["farm_id"]), "plantingStatus": "planned" if farm["planned_planting_date"] else "not_scheduled", "plantingDate": None, "plannedPlantingDate": farm["planned_planting_date"].isoformat() if farm["planned_planting_date"] else None, "templateVersionId": business_id("CTV", farm["planting_template_version_id"]) if farm["planting_template_version_id"] else None, "templateName": template["template_name"] if template else None, "crop": template["crop_name"] if template else None, "versionNo": template["version_no"] if template else None, "stage": None, "stages": [], "suggestions": []}
    return annual_production_context(connection, plot["farm_id"], plot_id=plot["plot_id"])


def annual_production_context(connection: Any, farm_id: int, *, production_year: int | None = None, plot_id: int | None = None) -> dict[str, Any]:
    today = date.today()
    production_year = production_year or production_year_for(today)
    farm = connection.execute(text("SELECT farm_id,farm_name,planting_date FROM farm WHERE farm_id=:farm"), {"farm": farm_id}).mappings().first()
    if farm is None:
        raise HTTPException(status_code=404, detail="FARM_NOT_FOUND")
    if farm["planting_date"] is None:
        raise HTTPException(status_code=409, detail="ORCHARD_PLANTING_DATE_REQUIRED")
    profile = connection.execute(text("SELECT p.template_version_id,t.template_name,t.crop_name,v.template_id,v.version_no FROM farm_production_profile p JOIN crop_template_version v ON v.template_version_id=p.template_version_id JOIN crop_template t ON t.template_id=v.template_id WHERE p.farm_id=:farm"), {"farm": farm_id}).mappings().first()
    annual = connection.execute(text("SELECT * FROM farm_annual_production WHERE farm_id=:farm AND production_year=:year"), {"farm": farm_id, "year": production_year}).mappings().first()
    if annual is None and profile is None:
        raise HTTPException(status_code=404, detail="PRODUCTION_TEMPLATE_REQUIRED")
    version_id = annual["template_version_id"] if annual else profile["template_version_id"]
    template = connection.execute(text("SELECT t.template_id,t.template_name,t.crop_name,v.version_no,v.farm_id FROM crop_template_version v JOIN crop_template t ON t.template_id=v.template_id WHERE v.template_version_id=:version"), {"version": version_id}).mappings().one()
    tree_age_years, automatic_track = orchard_track(farm["planting_date"], production_year)
    track = annual["orchard_track"] if annual else automatic_track
    stages = connection.execute(text("SELECT stage_code,stage_name,sort_order,duration_days,moisture_min,moisture_max,management_focus FROM growth_stage_template WHERE template_version_id=:version ORDER BY sort_order"), {"version": version_id}).mappings().all()
    stage_states = {row["stage_code"]: row for row in connection.execute(text("SELECT * FROM farm_annual_stage_state WHERE farm_id=:farm AND production_year=:year"), {"farm": farm_id, "year": production_year}).mappings().all()}
    default_stage_windows = {stage["stage_code"]: stage for stage in template_stage_windows(connection, version_id, production_year)}
    current_stage = None
    timeline: list[dict[str, Any]] = []
    for stage in stages:
        state = stage_states.get(stage["stage_code"])
        default_starts_on = default_stage_windows[stage["stage_code"]]["planned_start_date"]
        default_ends_on = default_stage_windows[stage["stage_code"]]["planned_end_date"]
        starts_on = state["planned_start_date"] if state else default_starts_on
        ends_on = state["planned_end_date"] if state else default_ends_on
        stage_status = state["stage_status"] if state else "planned"
        unresolved_previous = any(
            previous["stage_status"] == "planned" and previous["planned_end_date"] >= today
            for previous in stage_states.values()
            if previous["sort_order"] < stage["sort_order"]
        )
        item = {"code": stage["stage_code"], "name": stage["stage_name"], "sortOrder": stage["sort_order"], "startsOn": starts_on.isoformat(), "endsOn": ends_on.isoformat(), "actualStartedAt": iso(state["actual_started_at"]) if state else None, "status": stage_status, "locked": stage_status in {"active", "completed"}, "dispatched": bool(state and state["dispatched_at"]), "canStart": stage_status == "planned" and not unresolved_previous, "overdueConfirmation": stage_status == "planned" and starts_on < today, "managementFocus": stage["management_focus"], "moistureMin": float(stage["moisture_min"]) if stage["moisture_min"] is not None else None, "moistureMax": float(stage["moisture_max"]) if stage["moisture_max"] is not None else None}
        timeline.append(item)
        if stage_status == "active" or (current_stage is None and starts_on <= today <= ends_on):
            current_stage = item
    plot_offset_days = 0
    if plot_id is not None:
        plot_offset_days = connection.execute(
            text("SELECT production_date_offset_days FROM plot WHERE plot_id=:plot AND farm_id=:farm"),
            {"plot": plot_id, "farm": farm_id},
        ).scalar_one_or_none()
        if plot_offset_days is None:
            raise HTTPException(status_code=404, detail="PLOT_NOT_FOUND")
    selected_rules = farm_plan_task_rules(connection, version_id, production_year, track)
    plan_items = [{**rule, "scheduledDate": (rule["scheduledDate"] + timedelta(days=int(plot_offset_days or 0))).isoformat()} for rule in selected_rules]
    suggestions = [item for item in plan_items if current_stage and item["stageCode"] == current_stage["code"]]
    active_plan_versions = connection.execute(text("SELECT DISTINCT mp.template_version_id FROM management_plan mp JOIN plot p ON p.plot_id=mp.plot_id WHERE p.farm_id=:farm AND mp.plan_status='active'"), {"farm": farm_id}).scalars().all()
    active_plan_count = connection.execute(text("SELECT COUNT(*) FROM management_plan mp JOIN plot p ON p.plot_id=mp.plot_id WHERE p.farm_id=:farm AND mp.plan_status='active'"), {"farm": farm_id}).scalar()
    active_state = next((state for state in stage_states.values() if state["stage_status"] == "active"), None)
    has_undispatched_changes = annual is not None and active_state is not None and active_state["dispatched_template_version_id"] is None
    overdue_stages = [item for item in timeline if item["overdueConfirmation"]]
    return {"plotId": business_id("P", plot_id) if plot_id is not None else None, "farmId": business_id("F", farm_id), "farmName": farm["farm_name"], "productionYear": production_year, "productionYearStartsOn": production_year_bounds(production_year)[0].isoformat(), "productionYearEndsOn": production_year_bounds(production_year)[1].isoformat(), "templateId": business_id("CT", template["template_id"]), "templateVersionId": business_id("CTV", version_id), "templateName": template["template_name"], "templateScope": "farm" if template["farm_id"] is not None else "standard", "crop": template["crop_name"], "versionNo": template["version_no"], "versionDefinition": template_definition(connection, version_id), "annualConfigured": annual is not None, "annualStatus": annual["annual_status"] if annual else None, "activePlanCount": active_plan_count, "hasUndispatchedChanges": has_undispatched_changes, "canCustomizeInitialTemplate": annual is None and not active_plan_count, "plantingDate": farm["planting_date"].isoformat(), "treeAgeYears": tree_age_years, "track": track, "trackSource": annual["track_source"] if annual else "auto", "trackOverrideReason": annual["track_override_reason"] if annual else None, "stage": current_stage, "stages": timeline, "overdueStages": overdue_stages, "planItems": plan_items, "suggestions": suggestions}


def ensure_farmer_task_access(plot_id: int, user: dict[str, Any]) -> dict[str, Any]:
    if user["role"] != "farmer":
        raise HTTPException(status_code=403, detail="FORBIDDEN")
    return plot_scope(plot_id, user)


def task_row(connection: Any, task_id: int, *, lock: bool = False) -> Any:
    statement = """SELECT t.*,p.plot_name,p.farm_id,mp.plan_status,c.task_completion_id,c.result,c.executed_at,c.executor_user_id,c.executor_name,c.note AS completion_note,c.irrigation_record_id,c.record_type,c.corrects_completion_id FROM management_task t JOIN plot p ON p.plot_id=t.plot_id LEFT JOIN management_plan mp ON mp.plan_id=t.plan_id LEFT JOIN management_task_completion c ON c.task_completion_id=(SELECT c2.task_completion_id FROM management_task_completion c2 WHERE c2.task_id=t.task_id ORDER BY c2.task_completion_id DESC LIMIT 1) WHERE t.task_id=:task"""
    if lock:
        statement += " FOR UPDATE"
    row = connection.execute(text(statement), {"task": task_id}).mappings().first()
    if row is None:
        raise HTTPException(status_code=404, detail="TASK_NOT_FOUND")
    return row


def task_status_value(row: Any) -> str:
    if row["task_status"] == "pending" and row["scheduled_date"] < date.today():
        return "overdue"
    return row["task_status"]


def task_payload(row: Any, *, include_completion: bool = True) -> dict[str, Any]:
    payload = {"id": business_id("MT", row["task_id"]), "planId": business_id("MP", row["plan_id"]) if row["plan_id"] is not None else None, "plotId": business_id("P", row["plot_id"]), "plotName": row["plot_name"], "stageCode": row["stage_code"], "taskType": row["task_type"], "title": row["title"], "source": row["source"], "scheduledDate": row["scheduled_date"].isoformat(), "status": task_status_value(row), "note": row["note"], "updatedAt": iso(row["updated_at"])}
    if include_completion:
        payload["completion"] = None if row["task_completion_id"] is None else {"id": business_id("TC", row["task_completion_id"]), "result": row["result"], "executedAt": iso(row["executed_at"]), "executorUserId": business_id("U", row["executor_user_id"]), "executorName": row["executor_name"], "note": row["completion_note"], "irrigationRecordId": business_id("IR", row["irrigation_record_id"]) if row["irrigation_record_id"] is not None else None, "recordType": row["record_type"], "correctsCompletionId": business_id("TC", row["corrects_completion_id"]) if row["corrects_completion_id"] is not None else None}
    return payload


def finalize_farm_planting_if_ready(connection: Any, row: Any, completed_on: date, user: dict[str, Any]) -> None:
    if row["source"] != "planting":
        return
    farm = connection.execute(text("SELECT planting_date,planned_planting_date,planting_template_version_id FROM farm WHERE farm_id=:farm FOR UPDATE"), {"farm": row["farm_id"]}).mappings().one()
    if farm["planting_date"] is not None or farm["planned_planting_date"] is None or farm["planting_template_version_id"] is None:
        return
    pending = connection.execute(text("SELECT COUNT(*) FROM management_task t JOIN plot p ON p.plot_id=t.plot_id WHERE p.farm_id=:farm AND t.source='planting' AND t.scheduled_date=:scheduled AND t.task_status IN ('pending','in_progress')"), {"farm": row["farm_id"], "scheduled": farm["planned_planting_date"]}).scalar()
    if pending:
        return
    connection.execute(text("UPDATE farm SET planting_date=:actual,planned_planting_date=NULL,planting_template_version_id=NULL WHERE farm_id=:farm"), {"actual": completed_on, "farm": row["farm_id"]})
    connection.execute(text("INSERT INTO farm_production_profile (farm_id,template_version_id,planting_date,tree_age_years,note,updated_by) VALUES (:farm,:version,:planting,NULL,:note,:user) ON DUPLICATE KEY UPDATE template_version_id=VALUES(template_version_id),planting_date=VALUES(planting_date),updated_by=VALUES(updated_by)"), {"farm": row["farm_id"], "version": farm["planting_template_version_id"], "planting": completed_on, "note": "定植任务全部完成后自动建立", "user": user["user_id"]})
    record_audit_log(connection, user, "production_plan", "confirm_farm_planting", target_type="farm", target_id=business_id("F", row["farm_id"]), detail=f"Confirmed actual planting date: {completed_on.isoformat()}")


def record_task_adjustment(connection: Any, task_id: int, user_id: int, adjustment_type: str, before: dict[str, Any], after: dict[str, Any]) -> None:
    connection.execute(text("INSERT INTO management_task_adjustment (task_id,changed_by,adjustment_type,before_value,after_value) VALUES (:task,:user,:kind,:before,:after)"), {"task": task_id, "user": user_id, "kind": adjustment_type, "before": json.dumps(before), "after": json.dumps(after)})


def task_irrigation_record_id(connection: Any, value: str | None, plot_id: int) -> int | None:
    if value is None:
        return None
    record_id = parse_resource_id(value, "IR")
    if not connection.execute(text("SELECT 1 FROM irrigation_record WHERE record_id=:record AND plot_id=:plot"), {"record": record_id, "plot": plot_id}).scalar():
        raise HTTPException(status_code=400, detail="INVALID_REQUEST")
    return record_id


@app.get("/api/users")
def list_phase2_users(
    page: int = Query(default=1, ge=1), page_size: int = Query(default=20, ge=1, le=100, alias="pageSize"),
    keyword: str | None = None, role: str | None = None, status_value: str | None = Query(default=None, alias="status"),
    user: dict[str, Any] = Depends(current_user),
) -> dict[str, Any]:
    ensure_system_admin(user)
    clauses, params = [], {"limit": page_size, "offset": (page - 1) * page_size}
    if keyword:
        clauses.append("(employee_no LIKE :keyword OR real_name LIKE :keyword OR phone LIKE :keyword)"); params["keyword"] = f"%{keyword}%"
    if status_value:
        clauses.append("user_status = :status"); params["status"] = status_value
    if role:
        clauses.append("EXISTS (SELECT 1 FROM user_role_rel ur JOIN roles r ON r.role_id = ur.role_id WHERE ur.user_id = users.user_id AND r.role_code = :role)"); params["role"] = role
    where = f"WHERE {' AND '.join(clauses)}" if clauses else ""
    with engine.connect() as connection:
        total = connection.execute(text(f"SELECT COUNT(*) FROM users {where}"), params).scalar_one()
        ids = connection.execute(text(f"SELECT user_id FROM users {where} ORDER BY user_id LIMIT :limit OFFSET :offset"), params).scalars().all()
        return {"items": [user_payload(connection, item) for item in ids], "page": page, "pageSize": page_size, "total": total}


@app.post("/api/users", status_code=status.HTTP_201_CREATED)
def create_phase2_user(request: UserCreateRequest, user: dict[str, Any] = Depends(current_user)) -> dict[str, Any]:
    ensure_system_admin(user)
    employee_no = request.employeeNo.strip()
    id_card_no = normalized_id_card(request.idCardNo)
    initial_password = id_card_no[-6:]
    with engine.begin() as connection:
        roles, farms, plots = validated_assignments(connection, request.roleCodes, request.farmIds, request.plotIds)
        try:
            initial_password_hash = password_hash(initial_password)
            result = connection.execute(text("INSERT INTO users (employee_no, username, password_hash, initial_password_hash, id_card_hash, id_card_masked, real_name, email, phone, user_status) VALUES (:employee_no, :username, :password_hash, :initial_password_hash, :id_card_hash, :id_card_masked, :name, :email, :phone, 'pending_activation')"), {"employee_no": employee_no, "username": employee_no, "password_hash": initial_password_hash, "initial_password_hash": initial_password_hash, "id_card_hash": id_card_hash(id_card_no), "id_card_masked": id_card_masked(id_card_no), "name": request.name.strip(), "email": request.email or None, "phone": request.phone or None})
        except IntegrityError as error:
            raise HTTPException(status_code=409, detail="EMPLOYEE_NO_OR_ID_CARD_EXISTS") from error
        replace_user_assignments(connection, result.lastrowid, roles, farms, plots)
        record_audit_log(connection, user, "account_access", "create_employee", target_type="user", target_id=business_id("U", result.lastrowid), detail=f"Pre-provisioned employee {employee_no}")
        return user_payload(connection, result.lastrowid)


@app.patch("/api/users/{user_id}")
def update_phase2_user(user_id: str, request: UserUpdateRequest, user: dict[str, Any] = Depends(current_user)) -> dict[str, Any]:
    ensure_system_admin(user); numeric_id = parse_resource_id(user_id, "U")
    with engine.begin() as connection:
        current = user_payload(connection, numeric_id)
        roles, farms, plots = validated_assignments(connection, request.roleCodes or current["roleCodes"], request.farmIds if request.farmIds is not None else current["farmIds"], request.plotIds if request.plotIds is not None else current["plotIds"], user_id=numeric_id)
        fields = request.model_fields_set
        values = {"id": numeric_id, "name": request.name if request.name is not None else current["name"], "email": request.email if request.email is not None else current["email"], "phone": request.phone if request.phone is not None else current["phone"]}
        assignments = ["real_name = :name", "email = :email", "phone = :phone"]
        if "idCardNo" in fields:
            if request.idCardNo is None:
                raise HTTPException(status_code=400, detail="INVALID_ID_CARD_NO")
            id_card_no = normalized_id_card(request.idCardNo)
            initial_password_hash = password_hash(id_card_no[-6:])
            values.update({"id_card_hash": id_card_hash(id_card_no), "id_card_masked": id_card_masked(id_card_no), "initial_password_hash": initial_password_hash, "password_hash": initial_password_hash})
            assignments.extend(["id_card_hash = :id_card_hash", "id_card_masked = :id_card_masked", "initial_password_hash = :initial_password_hash", "password_hash = :password_hash", "user_status = 'pending_activation'", "activated_at = NULL"])
        try:
            result = connection.execute(text("UPDATE users SET " + ", ".join(assignments) + " WHERE user_id = :id"), values)
        except IntegrityError as error:
            raise HTTPException(status_code=409, detail="EMPLOYEE_NO_OR_ID_CARD_EXISTS") from error
        if result.rowcount == 0:
            raise HTTPException(status_code=404, detail="USER_NOT_FOUND")
        replace_user_assignments(connection, numeric_id, roles, farms, plots)
        record_audit_log(connection, user, "account_access", "update_employee", target_type="user", target_id=business_id("U", numeric_id), detail="Updated employee profile, role, or authorization scope")
        return user_payload(connection, numeric_id)


@app.patch("/api/users/{user_id}/status")
def update_phase2_user_status(user_id: str, request: UserStatusUpdateRequest, user: dict[str, Any] = Depends(current_user)) -> dict[str, Any]:
    ensure_system_admin(user); numeric_id = parse_resource_id(user_id, "U")
    if request.status not in {"active", "disabled"}: raise HTTPException(status_code=400, detail="INVALID_REQUEST")
    if numeric_id == user["user_id"] and request.status == "disabled": raise HTTPException(status_code=400, detail="CANNOT_DISABLE_CURRENT_USER")
    with engine.begin() as connection:
        result = connection.execute(text("UPDATE users SET user_status = :status WHERE user_id = :id"), {"status": request.status, "id": numeric_id})
        if result.rowcount == 0: raise HTTPException(status_code=404, detail="USER_NOT_FOUND")
        record_audit_log(connection, user, "account_access", "change_employee_status", target_type="user", target_id=business_id("U", numeric_id), detail=f"Set employee status to {request.status}")
        return user_payload(connection, numeric_id)


@app.post("/api/users/{user_id}/initial-password/reset")
def reset_initial_password(user_id: str, user: dict[str, Any] = Depends(current_user)) -> dict[str, Any]:
    ensure_system_admin(user); numeric_id = parse_resource_id(user_id, "U")
    with engine.begin() as connection:
        result = connection.execute(text("UPDATE users SET password_hash = initial_password_hash, user_status = 'pending_activation', activated_at = NULL WHERE user_id = :id AND initial_password_hash IS NOT NULL"), {"id": numeric_id})
        if result.rowcount == 0: raise HTTPException(status_code=404, detail="USER_NOT_FOUND")
        record_audit_log(connection, user, "account_access", "reset_initial_password", target_type="user", target_id=business_id("U", numeric_id), detail="Reset initial password and required first login")
    return {"userId": business_id("U", numeric_id), "status": "pending_activation"}


@app.get("/api/users/import-template")
def download_user_import_template(user: dict[str, Any] = Depends(current_user)) -> StreamingResponse:
    ensure_system_admin(user)
    workbook = Workbook(); sheet = workbook.active; sheet.title = "员工导入"
    sheet.append(IMPORT_HEADERS)
    sheet.append(["E2026001", "张三", "11010519491231002X", "13800138000", "zhangsan@example.com", "farmer", "", "P001,P002"])
    for column in (1, 3, 4):
        for cell in sheet.iter_cols(min_col=column, max_col=column):
            for item in cell: item.number_format = "@"
    sheet.freeze_panes = "A2"
    data = BytesIO(); workbook.save(data); data.seek(0)
    return StreamingResponse(data, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": "attachment; filename=employee-import-template.xlsx"})


@app.post("/api/users/import-jobs", status_code=status.HTTP_201_CREATED)
async def create_user_import_job(file: UploadFile = File(...), user: dict[str, Any] = Depends(current_user)) -> dict[str, Any]:
    ensure_system_admin(user)
    if not (file.filename or "").lower().endswith(".xlsx"): raise HTTPException(status_code=400, detail="INVALID_IMPORT_FILE")
    raw = await file.read()
    if not raw or len(raw) > 5 * 1024 * 1024: raise HTTPException(status_code=400, detail="INVALID_IMPORT_FILE")
    try:
        sheet = load_workbook(BytesIO(raw), read_only=True, data_only=True).active
        headers = [str(value or "").strip() for value in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))]
    except Exception as error:
        raise HTTPException(status_code=400, detail="INVALID_IMPORT_FILE") from error
    if headers != IMPORT_HEADERS: raise HTTPException(status_code=400, detail="INVALID_IMPORT_HEADERS")
    parsed_rows = [(number, dict(zip(IMPORT_HEADERS, values))) for number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2) if any(value not in (None, "") for value in values)]
    with engine.begin() as connection:
        job_id = connection.execute(text("INSERT INTO user_import_job (original_filename, created_by) VALUES (:name,:user_id)"), {"name": (file.filename or "import.xlsx")[:255], "user_id": user["user_id"]}).lastrowid
        seen_employees, seen_id_cards, seen_plots, valid_count, invalid_count = set(), set(), set(), 0, 0
        for row_number, source in parsed_rows:
            errors: list[dict[str, str]] = []
            employee_no = str(source["工号"] or "").strip(); name = str(source["姓名"] or "").strip(); roles = import_values(source["角色"]); farms = import_values(source["授权橙园编号"]); plots = import_values(source["授权地块编号"])
            try:
                id_card = normalized_id_card(str(source["身份证号"] or "")); card_hash = id_card_hash(id_card)
                if not employee_no or not name: raise HTTPException(status_code=400, detail="REQUIRED_FIELD_MISSING")
                if employee_no in seen_employees or card_hash in seen_id_cards: raise HTTPException(status_code=409, detail="DUPLICATE_IMPORT_VALUE")
                if connection.execute(text("SELECT 1 FROM users WHERE employee_no=:employee OR id_card_hash=:card LIMIT 1"), {"employee": employee_no, "card": card_hash}).scalar(): raise HTTPException(status_code=409, detail="EMPLOYEE_NO_OR_ID_CARD_EXISTS")
                role_codes, farm_ids, plot_ids = validated_assignments(connection, roles, farms, plots)
                if any(plot_id in seen_plots for plot_id in plot_ids): raise HTTPException(status_code=409, detail="PLOT_ALREADY_ASSIGNED")
                seen_employees.add(employee_no); seen_id_cards.add(card_hash); seen_plots.update(plot_ids)
            except HTTPException as error:
                errors.append({"field": "导入行", "code": str(error.detail), "message": str(error.detail)}); card_hash = None; id_card = None; role_codes = []; farm_ids = []; plot_ids = []
            source_data = {"employeeNo": employee_no, "name": name, "phone": str(source["手机号"] or "").strip() or None, "email": str(source["邮箱"] or "").strip() or None, "roleCodes": role_codes, "farmIds": farm_ids, "plotIds": plot_ids, "idCardHash": card_hash, "idCardMasked": id_card_masked(id_card) if id_card else None, "initialPasswordHash": password_hash(id_card[-6:]) if id_card else None}
            row_status = "invalid" if errors else "valid"; valid_count += not errors; invalid_count += bool(errors)
            connection.execute(text("INSERT INTO user_import_row (import_job_id, source_row_number, source_data, validation_errors, row_status) VALUES (:job,:row,:data,:errors,:status)"), {"job": job_id, "row": row_number, "data": json.dumps(source_data), "errors": json.dumps(errors) if errors else None, "status": row_status})
        job_status = "ready" if parsed_rows and not invalid_count else "failed"
        connection.execute(text("UPDATE user_import_job SET job_status=:status,total_rows=:total,valid_rows=:valid,invalid_rows=:invalid WHERE import_job_id=:id"), {"id": job_id, "status": job_status, "total": len(parsed_rows), "valid": valid_count, "invalid": invalid_count})
        return import_job_payload(connection, job_id)


@app.get("/api/users/import-jobs/{job_id}")
def get_user_import_job(job_id: str, user: dict[str, Any] = Depends(current_user)) -> dict[str, Any]:
    ensure_system_admin(user); numeric_id = parse_resource_id(job_id, "IMP")
    with engine.connect() as connection:
        if not connection.execute(text("SELECT 1 FROM user_import_job WHERE import_job_id=:id"), {"id": numeric_id}).scalar(): raise HTTPException(status_code=404, detail="IMPORT_JOB_NOT_FOUND")
        return import_job_payload(connection, numeric_id)


@app.post("/api/users/import-jobs/{job_id}/commit")
def commit_user_import_job(job_id: str, user: dict[str, Any] = Depends(current_user)) -> dict[str, Any]:
    ensure_system_admin(user); numeric_id = parse_resource_id(job_id, "IMP"); now = datetime.now(UTC).replace(tzinfo=None)
    with engine.begin() as connection:
        job = connection.execute(text("SELECT job_status FROM user_import_job WHERE import_job_id=:id FOR UPDATE"), {"id": numeric_id}).mappings().first()
        if job is None: raise HTTPException(status_code=404, detail="IMPORT_JOB_NOT_FOUND")
        if job["job_status"] == "committed": return import_job_payload(connection, numeric_id)
        if job["job_status"] != "ready": raise HTTPException(status_code=422, detail="IMPORT_VALIDATION_FAILED")
        rows = connection.execute(text("SELECT import_row_id, source_data FROM user_import_row WHERE import_job_id=:id AND row_status='valid' ORDER BY source_row_number FOR UPDATE"), {"id": numeric_id}).mappings().all()
        for row in rows:
            data = json.loads(row["source_data"])
            result = connection.execute(text("INSERT INTO users (employee_no, username, password_hash, initial_password_hash, id_card_hash, id_card_masked, real_name, email, phone, user_status) VALUES (:employeeNo,:employeeNo,:initialPasswordHash,:initialPasswordHash,:idCardHash,:idCardMasked,:name,:email,:phone,'pending_activation')"), data)
            replace_user_assignments(connection, result.lastrowid, data["roleCodes"], data["farmIds"], data["plotIds"])
            connection.execute(text("UPDATE user_import_row SET row_status='imported', imported_user_id=:user_id WHERE import_row_id=:id"), {"id": row["import_row_id"], "user_id": result.lastrowid})
        connection.execute(text("UPDATE user_import_job SET job_status='committed', committed_by=:user, committed_at=:now WHERE import_job_id=:id"), {"id": numeric_id, "user": user["user_id"], "now": now})
        record_audit_log(connection, user, "account_access", "commit_employee_import", target_type="import_job", target_id=business_id("IMP", numeric_id), detail=f"Imported {len(rows)} employee accounts")
        return import_job_payload(connection, numeric_id)


@app.get("/api/audit-logs")
def list_audit_logs(
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=20, ge=1, le=100, alias="pageSize"),
    module: str | None = None,
    keyword: str | None = None,
    user: dict[str, Any] = Depends(current_user),
) -> dict[str, Any]:
    ensure_system_admin(user)
    clauses: list[str] = []
    params: dict[str, Any] = {"limit": page_size, "offset": (page - 1) * page_size}
    if module:
        clauses.append("a.module_code = :module")
        params["module"] = module
    if keyword:
        clauses.append("(u.real_name LIKE :keyword OR u.employee_no LIKE :keyword OR a.action_code LIKE :keyword OR a.detail LIKE :keyword)")
        params["keyword"] = f"%{keyword}%"
    where = " WHERE " + " AND ".join(clauses) if clauses else ""
    with engine.connect() as connection:
        total = connection.execute(
            text("SELECT COUNT(*) FROM audit_log AS a JOIN users AS u ON u.user_id = a.actor_user_id" + where),
            params,
        ).scalar_one()
        rows = connection.execute(
            text(
                """
                SELECT a.audit_log_id, a.module_code, a.action_code, a.target_type,
                       a.target_id, a.detail, a.created_at, u.real_name, u.employee_no
                FROM audit_log AS a
                JOIN users AS u ON u.user_id = a.actor_user_id
                """
                + where
                + " ORDER BY a.created_at DESC, a.audit_log_id DESC LIMIT :limit OFFSET :offset"
            ),
            params,
        ).mappings().all()
    return {
        "items": [
            {
                "id": business_id("AL", row["audit_log_id"]),
                "occurredAt": iso(row["created_at"]),
                "actorName": row["real_name"],
                "actorEmployeeNo": row["employee_no"],
                "module": row["module_code"],
                "action": row["action_code"],
                "targetType": row["target_type"],
                "targetId": row["target_id"],
                "detail": row["detail"],
            }
            for row in rows
        ],
        "page": page,
        "pageSize": page_size,
        "total": total,
    }


@app.get("/api/farms")
def list_farms(status_value: str | None = Query(default=None, alias="status"), user: dict[str, Any] = Depends(current_user)) -> list[dict[str, Any]]:
    allowed_farms = accessible_farm_ids(user)
    with engine.connect() as connection:
        if allowed_farms is None:
            rows = connection.execute(
                text("SELECT * FROM farm WHERE (:status IS NULL OR farm_status = :status) ORDER BY farm_id"), {"status": status_value}
            ).mappings().all()
        elif not allowed_farms:
            rows = []
        else:
            statement = text(
                "SELECT * FROM farm WHERE farm_id IN :farm_ids "
                "AND (:status IS NULL OR farm_status = :status) ORDER BY farm_id"
            ).bindparams(bindparam("farm_ids", expanding=True))
            rows = connection.execute(
                statement, {"farm_ids": list(allowed_farms), "status": status_value}
            ).mappings().all()

    return [
        farm_payload(row) for row in rows
    ]


@app.post("/api/crop-templates", status_code=status.HTTP_201_CREATED)
def create_crop_template(request: TemplateCreateRequest, user: dict[str, Any] = Depends(current_user)) -> dict[str, Any]:
    ensure_system_admin(user)
    now = datetime.now(UTC).replace(tzinfo=None)
    with engine.begin() as connection:
        template = connection.execute(text("INSERT INTO crop_template (template_name,crop_name,province,city,created_by) VALUES (:name,:crop,:province,:city,:user)"), {"name":request.name.strip(),"crop":request.crop.strip(),"province":request.province,"city":request.city,"user":user["user_id"]})
        version = connection.execute(text("INSERT INTO crop_template_version (template_id,version_no,version_status,published_at,published_by) VALUES (:template,1,'published',:now,:user)"), {"template":template.lastrowid,"now":now,"user":user["user_id"]})
        store_template_definition(connection, version.lastrowid, request.versionDefinition)
        return {"id":business_id("CT",template.lastrowid),"versionId":business_id("CTV",version.lastrowid),"versionNo":1,"status":"published"}


@app.get("/api/crop-templates")
def list_crop_templates(crop: str | None = None, province: str | None = None, city: str | None = None, include_all_versions: bool = Query(default=False, alias="includeAllVersions"), user: dict[str, Any] = Depends(current_user)) -> list[dict[str, Any]]:
    clauses, params = ["t.template_status='active'", "v.version_status='published'", "v.farm_id IS NULL"], {}
    if not include_all_versions:
        clauses.append("NOT EXISTS (SELECT 1 FROM crop_template_version newer WHERE newer.template_id=v.template_id AND newer.farm_id IS NULL AND newer.version_status='published' AND newer.version_no>v.version_no)")
    for field, value in (("crop_name", crop), ("province", province), ("city", city)):
        if value: clauses.append(f"t.{field}=:{field}"); params[field] = value
    with engine.connect() as connection:
        rows = connection.execute(text("SELECT t.template_id,t.template_name,t.crop_name,t.province,t.city,v.template_version_id,v.version_no FROM crop_template t JOIN crop_template_version v ON v.template_id=t.template_id WHERE " + " AND ".join(clauses) + " ORDER BY t.template_id,v.version_no DESC"), params).mappings().all()
    return [{"id":business_id("CT",r["template_id"]),"name":r["template_name"],"crop":r["crop_name"],"province":r["province"],"city":r["city"],"versionId":business_id("CTV",r["template_version_id"]),"versionNo":r["version_no"]} for r in rows]


@app.get("/api/crop-template-versions/{version_id}")
def get_crop_template_version(version_id: str, user: dict[str, Any] = Depends(current_user)) -> dict[str, Any]:
    numeric_version_id = parse_resource_id(version_id, "CTV")
    with engine.connect() as connection:
        version = connection.execute(text("SELECT v.*,t.template_name,t.crop_name,t.province,t.city FROM crop_template_version v JOIN crop_template t ON t.template_id=v.template_id WHERE v.template_version_id=:version AND v.version_status='published'"), {"version": numeric_version_id}).mappings().first()
        if version is None:
            raise HTTPException(status_code=404, detail="TEMPLATE_NOT_FOUND")
        if version["farm_id"] is not None:
            ensure_farm_access(user, version["farm_id"])
        return {"id": business_id("CT", version["template_id"]), "versionId": business_id("CTV", numeric_version_id), "versionNo": version["version_no"], "name": version["template_name"], "crop": version["crop_name"], "province": version["province"], "city": version["city"], "changeNote": version["change_note"], "versionDefinition": template_definition(connection, numeric_version_id)}


@app.post("/api/crop-templates/{template_id}/versions", status_code=status.HTTP_201_CREATED)
def create_crop_template_version(template_id: str, request: TemplateVersionCreateRequest, user: dict[str, Any] = Depends(current_user)) -> dict[str, Any]:
    ensure_system_admin(user)
    numeric_template_id = parse_resource_id(template_id, "CT")
    now = datetime.now(UTC).replace(tzinfo=None)
    with engine.begin() as connection:
        template = connection.execute(text("SELECT template_id FROM crop_template WHERE template_id=:template AND template_status='active' FOR UPDATE"), {"template": numeric_template_id}).mappings().first()
        if template is None:
            raise HTTPException(status_code=404, detail="TEMPLATE_NOT_FOUND")
        version_no = connection.execute(text("SELECT COALESCE(MAX(version_no),0)+1 FROM crop_template_version WHERE template_id=:template"), {"template": numeric_template_id}).scalar_one()
        version = connection.execute(text("INSERT INTO crop_template_version (template_id,version_no,version_status,change_note,published_at,published_by) VALUES (:template,:version,'published',:note,:now,:user)"), {"template": numeric_template_id, "version": version_no, "note": request.changeNote, "now": now, "user": user["user_id"]})
        store_template_definition(connection, version.lastrowid, request.versionDefinition)
        record_audit_log(connection, user, "production_plan", "publish_standard_template_version", target_type="crop_template", target_id=business_id("CT", numeric_template_id), detail=f"Published standard template V{version_no}")
        return {"id": business_id("CT", numeric_template_id), "versionId": business_id("CTV", version.lastrowid), "versionNo": version_no, "status": "published"}


@app.post("/api/crop-templates/{template_id}/drafts", status_code=status.HTTP_201_CREATED)
def create_crop_template_draft(template_id: str, request: TemplateDraftCreateRequest, farm_id: str = Query(alias="farmId"), user: dict[str, Any] = Depends(current_user)) -> dict[str, Any]:
    numeric_template_id = parse_resource_id(template_id, "CT")
    numeric_farm_id = parse_resource_id(farm_id, "F")
    ensure_farm_admin_access(user, numeric_farm_id)
    with engine.begin() as connection:
        template = connection.execute(text("SELECT template_id FROM crop_template WHERE template_id=:template AND template_status='active' FOR UPDATE"), {"template": numeric_template_id}).mappings().first()
        if template is None:
            raise HTTPException(status_code=404, detail="TEMPLATE_NOT_FOUND")
        profile = connection.execute(text("SELECT p.template_version_id,COALESCE(a.template_version_id,p.template_version_id) AS effective_template_version_id FROM farm_production_profile p LEFT JOIN farm_annual_production a ON a.farm_id=p.farm_id AND a.production_year=:year WHERE p.farm_id=:farm FOR UPDATE"), {"farm": numeric_farm_id, "year": production_year_for(date.today())}).mappings().first()
        if profile is None:
            raise HTTPException(status_code=409, detail="PRODUCTION_TEMPLATE_REQUIRED")
        source_version_id = parse_resource_id(request.sourceVersionId, "CTV") if request.sourceVersionId else profile["effective_template_version_id"]
        source = connection.execute(text("SELECT template_version_id,template_id FROM crop_template_version WHERE template_version_id=:version AND version_status='published' AND (farm_id IS NULL OR farm_id=:farm)"), {"version": source_version_id, "farm": numeric_farm_id}).mappings().first()
        if source is None or source["template_id"] != numeric_template_id:
            raise HTTPException(status_code=409, detail="PRODUCTION_TEMPLATE_REQUIRED")
        existing = connection.execute(text("SELECT * FROM crop_template_version WHERE farm_id=:farm AND version_status='draft' FOR UPDATE"), {"farm": numeric_farm_id}).mappings().first()
        if existing is not None:
            raise HTTPException(status_code=409, detail="TEMPLATE_DRAFT_EXISTS")
        version_no = connection.execute(text("SELECT COALESCE(MAX(version_no), 0) + 1 FROM crop_template_version WHERE template_id=:template"), {"template": numeric_template_id}).scalar_one()
        version_id = connection.execute(text("INSERT INTO crop_template_version (template_id,farm_id,version_no,version_status,change_note) VALUES (:template,:farm,:version,'draft',:note)"), {"template": numeric_template_id, "farm": numeric_farm_id, "version": version_no, "note": request.changeNote}).lastrowid
        store_template_definition(connection, version_id, template_definition(connection, source_version_id))
        version = connection.execute(text("SELECT * FROM crop_template_version WHERE template_version_id=:id"), {"id": version_id}).mappings().one()
        return template_draft_payload(connection, version)


@app.get("/api/farms/{farm_id}/cycle-draft")
def get_farm_cycle_draft(farm_id: str, user: dict[str, Any] = Depends(current_user)) -> dict[str, Any]:
    numeric_farm_id = parse_resource_id(farm_id, "F")
    ensure_farm_admin_access(user, numeric_farm_id)
    with engine.connect() as connection:
        draft = connection.execute(
            text(
                "SELECT * FROM crop_template_version "
                "WHERE farm_id=:farm AND version_status='draft' "
                "ORDER BY template_version_id DESC LIMIT 1"
            ),
            {"farm": numeric_farm_id},
        ).mappings().first()
        if draft is None:
            raise HTTPException(status_code=404, detail="TEMPLATE_DRAFT_NOT_FOUND")
        return template_draft_payload(connection, draft)


@app.put("/api/crop-template-drafts/{draft_id}")
def update_crop_template_draft(draft_id: str, request: TemplateDraftUpdateRequest, user: dict[str, Any] = Depends(current_user)) -> dict[str, Any]:
    numeric_draft_id = parse_resource_id(draft_id, "CTV")
    validate_template_definition(request.versionDefinition)
    with engine.begin() as connection:
        draft = connection.execute(text("SELECT * FROM crop_template_version WHERE template_version_id=:id AND version_status='draft' FOR UPDATE"), {"id": numeric_draft_id}).mappings().first()
        if draft is None or draft["farm_id"] is None:
            raise HTTPException(status_code=404, detail="TEMPLATE_NOT_FOUND")
        ensure_farm_admin_access(user, draft["farm_id"])
        if connection.execute(text("SELECT 1 FROM farm_production_profile WHERE farm_id=:farm FOR UPDATE"), {"farm": draft["farm_id"]}).scalar() is None:
            raise HTTPException(status_code=409, detail="PRODUCTION_TEMPLATE_REQUIRED")
        connection.execute(text("DELETE mt FROM management_task_template mt JOIN growth_stage_template st ON st.stage_template_id=mt.stage_template_id WHERE st.template_version_id=:version"), {"version": numeric_draft_id})
        connection.execute(text("DELETE FROM growth_stage_template WHERE template_version_id=:version"), {"version": numeric_draft_id})
        store_template_definition(connection, numeric_draft_id, request.versionDefinition)
        connection.execute(text("UPDATE crop_template_version SET change_note=:note WHERE template_version_id=:id"), {"note": request.changeNote, "id": numeric_draft_id})
        version = connection.execute(text("SELECT * FROM crop_template_version WHERE template_version_id=:id"), {"id": numeric_draft_id}).mappings().one()
        record_audit_log(connection, user, "production_plan", "update_cycle_draft", target_type="cycle_draft", target_id=business_id("CTV", numeric_draft_id), detail="Updated the farm's next-year management plan draft")
        return template_draft_payload(connection, version)


@app.post("/api/crop-template-drafts/{draft_id}/publish")
def publish_crop_template_draft(draft_id: str, request: TemplateDraftPublishRequest, user: dict[str, Any] = Depends(current_user)) -> dict[str, Any]:
    numeric_draft_id = parse_resource_id(draft_id, "CTV")
    now = datetime.now(UTC).replace(tzinfo=None)
    with engine.begin() as connection:
        draft = connection.execute(text("SELECT * FROM crop_template_version WHERE template_version_id=:id AND version_status='draft' FOR UPDATE"), {"id": numeric_draft_id}).mappings().first()
        if draft is None or draft["farm_id"] is None:
            raise HTTPException(status_code=404, detail="TEMPLATE_NOT_FOUND")
        ensure_farm_admin_access(user, draft["farm_id"])
        profile = connection.execute(text("SELECT * FROM farm_production_profile WHERE farm_id=:farm FOR UPDATE"), {"farm": draft["farm_id"]}).mappings().first()
        if profile is None:
            raise HTTPException(status_code=409, detail="PRODUCTION_TEMPLATE_REQUIRED")
        current_year = production_year_for(date.today())
        current_annual = connection.execute(text("SELECT 1 FROM farm_annual_production WHERE farm_id=:farm AND production_year=:year"), {"farm": draft["farm_id"], "year": current_year}).scalar()
        plan_exists = connection.execute(text("SELECT 1 FROM management_plan mp JOIN plot p ON p.plot_id=mp.plot_id WHERE p.farm_id=:farm LIMIT 1"), {"farm": draft["farm_id"]}).scalar()
        purpose = "initial" if not current_annual and not plan_exists else "next"
        expected_year = production_year_for(date.today()) if purpose == "initial" else production_year_for(date.today()) + 1
        if request.productionYear != expected_year:
            raise HTTPException(status_code=409, detail="NEXT_PRODUCTION_YEAR_REQUIRED")
        planting_date = connection.execute(text("SELECT planting_date FROM farm WHERE farm_id=:farm"), {"farm": draft["farm_id"]}).scalar_one()
        tree_age_years, automatic_track = orchard_track(planting_date, request.productionYear)
        can_extend = automatic_track == "fruiting" and tree_age_years == 4
        if request.extendYoungTrack and not can_extend:
            raise HTTPException(status_code=409, detail="YOUNG_TRACK_EXTENSION_NOT_ALLOWED")
        if request.extendYoungTrack and not request.trackOverrideReason:
            raise HTTPException(status_code=400, detail="YOUNG_TRACK_EXTENSION_REASON_REQUIRED")
        track = "young" if request.extendYoungTrack else automatic_track
        source = "extended_young" if request.extendYoungTrack else "auto"
        connection.execute(text("UPDATE crop_template_version SET version_status='published', published_at=:now, published_by=:user WHERE template_version_id=:id"), {"now": now, "user": user["user_id"], "id": numeric_draft_id})
        connection.execute(text("INSERT INTO farm_annual_production (farm_id,production_year,template_version_id,orchard_track,track_source,track_override_reason,annual_status,created_by,updated_by) VALUES (:farm,:year,:version,:track,:source,:reason,'saved',:user,:user) ON DUPLICATE KEY UPDATE template_version_id=VALUES(template_version_id),orchard_track=VALUES(orchard_track),track_source=VALUES(track_source),track_override_reason=VALUES(track_override_reason),annual_status='saved',updated_by=VALUES(updated_by)"), {"farm": draft["farm_id"], "year": request.productionYear, "version": numeric_draft_id, "track": track, "source": source, "reason": request.trackOverrideReason.strip() if request.trackOverrideReason else None, "user": user["user_id"]})
        ensure_annual_stage_states(connection, draft["farm_id"], request.productionYear, numeric_draft_id)
        if purpose == "initial":
            connection.execute(text("UPDATE farm_production_profile SET template_version_id=:version,updated_by=:user WHERE farm_id=:farm"), {"version": numeric_draft_id, "user": user["user_id"], "farm": draft["farm_id"]})
        record_audit_log(connection, user, "production_plan", "publish_cycle_draft", target_type="cycle_draft", target_id=business_id("CTV", numeric_draft_id), detail=f"Published the {purpose} farm version for {business_id('F', draft['farm_id'])}")
        return {"id": business_id("CTV", numeric_draft_id), "templateId": business_id("CT", draft["template_id"]), "farmId": business_id("F", draft["farm_id"]), "productionYear": request.productionYear, "track": track, "status": "published", "publishedAt": iso(now)}


@app.post("/api/farms/{farm_id}/stages/start")
def start_farm_annual_stage(farm_id: str, request: FarmStageStartRequest, user: dict[str, Any] = Depends(current_user)) -> dict[str, Any]:
    numeric_farm_id = parse_resource_id(farm_id, "F")
    ensure_farm_admin_access(user, numeric_farm_id)
    actual_started_at = request.actualStartedAt
    if actual_started_at.tzinfo is not None:
        actual_started_at = actual_started_at.astimezone().replace(tzinfo=None)
    if actual_started_at > datetime.now():
        raise HTTPException(status_code=400, detail="INVALID_STAGE_START_TIME")
    production_year = production_year_for(actual_started_at.date())
    starts_on, ends_on = production_year_bounds(production_year)
    if not starts_on <= actual_started_at.date() <= ends_on:
        raise HTTPException(status_code=400, detail="INVALID_STAGE_START_TIME")
    with engine.begin() as connection:
        annual = connection.execute(text("SELECT * FROM farm_annual_production WHERE farm_id=:farm AND production_year=:year FOR UPDATE"), {"farm": numeric_farm_id, "year": production_year}).mappings().first()
        if annual is None:
            raise HTTPException(status_code=409, detail="ANNUAL_PLAN_NOT_SAVED")
        ensure_annual_stage_states(connection, numeric_farm_id, production_year, annual["template_version_id"])
        stage = connection.execute(text("SELECT * FROM farm_annual_stage_state WHERE farm_id=:farm AND production_year=:year AND stage_code=:stage FOR UPDATE"), {"farm": numeric_farm_id, "year": production_year, "stage": request.stageCode}).mappings().first()
        if stage is None:
            raise HTTPException(status_code=404, detail="STAGE_NOT_FOUND")
        if stage["stage_status"] == "completed":
            raise HTTPException(status_code=409, detail="STAGE_ALREADY_COMPLETED")
        if stage["stage_status"] == "active":
            return annual_production_context(connection, numeric_farm_id, production_year=production_year)
        unresolved_previous = connection.execute(
            text("SELECT stage_code FROM farm_annual_stage_state WHERE farm_id=:farm AND production_year=:year AND sort_order<:order AND stage_status='planned' AND planned_end_date>=:started FOR UPDATE"),
            {"farm": numeric_farm_id, "year": production_year, "order": stage["sort_order"], "started": actual_started_at.date()},
        ).scalars().all()
        if unresolved_previous:
            raise HTTPException(status_code=409, detail={"code": "PREVIOUS_STAGE_NOT_RESOLVED", "stageCodes": unresolved_previous})
        connection.execute(text("UPDATE farm_annual_stage_state SET stage_status='not_dispatched',non_dispatch_reason='阶段开始确认时已过期' WHERE farm_id=:farm AND production_year=:year AND sort_order<:order AND stage_status='planned' AND planned_end_date<:started"), {"farm": numeric_farm_id, "year": production_year, "order": stage["sort_order"], "started": actual_started_at.date()})
        connection.execute(text("UPDATE farm_annual_stage_state SET stage_status='completed',completed_at=:completed WHERE farm_id=:farm AND production_year=:year AND sort_order<:order AND stage_status='active'"), {"farm": numeric_farm_id, "year": production_year, "order": stage["sort_order"], "completed": actual_started_at})
        delta_days = (actual_started_at.date() - stage["planned_start_date"]).days
        if request.shiftFollowingStages and delta_days:
            following = connection.execute(text("SELECT stage_code,planned_start_date,planned_end_date FROM farm_annual_stage_state WHERE farm_id=:farm AND production_year=:year AND sort_order>=:order AND stage_status IN ('planned','not_dispatched') FOR UPDATE"), {"farm": numeric_farm_id, "year": production_year, "order": stage["sort_order"]}).mappings().all()
            for item in following:
                connection.execute(text("UPDATE farm_annual_stage_state SET planned_start_date=:starts,planned_end_date=:ends WHERE farm_id=:farm AND production_year=:year AND stage_code=:stage"), {"starts": item["planned_start_date"] + timedelta(days=delta_days), "ends": item["planned_end_date"] + timedelta(days=delta_days), "farm": numeric_farm_id, "year": production_year, "stage": item["stage_code"]})
        connection.execute(text("UPDATE farm_annual_stage_state SET stage_status='active',actual_started_at=:started,started_by=:user,locked_template_version_id=:version WHERE farm_id=:farm AND production_year=:year AND stage_code=:stage"), {"started": actual_started_at, "user": user["user_id"], "version": annual["template_version_id"], "farm": numeric_farm_id, "year": production_year, "stage": request.stageCode})
        record_audit_log(connection, user, "production_plan", "start_stage", target_type="farm_stage", target_id=f"{business_id('F', numeric_farm_id)}:{production_year}:{request.stageCode}", detail=f"Confirmed stage start at {actual_started_at.isoformat()}")
        return annual_production_context(connection, numeric_farm_id, production_year=production_year)


@app.post("/api/farms/{farm_id}/plans", status_code=status.HTTP_201_CREATED)
def generate_farm_plans(farm_id: str, request: FarmPlanGenerateRequest, user: dict[str, Any] = Depends(current_user)) -> dict[str, Any]:
    """Dispatch one active production stage without rewriting execution history."""
    numeric_farm_id = parse_resource_id(farm_id, "F")
    ensure_farm_admin_access(user, numeric_farm_id)
    today = date.today()
    with engine.begin() as connection:
        profile = connection.execute(text("SELECT * FROM farm_production_profile WHERE farm_id=:farm FOR UPDATE"), {"farm": numeric_farm_id}).mappings().first()
        if profile is None:
            raise HTTPException(status_code=409, detail="PRODUCTION_TEMPLATE_REQUIRED")
        annual_context = annual_production_context(connection, numeric_farm_id)
        if not annual_context["annualConfigured"]:
            raise HTTPException(status_code=409, detail="ANNUAL_PLAN_NOT_SAVED")
        annual_version_id = parse_resource_id(annual_context["templateVersionId"], "CTV")
        ensure_annual_stage_states(connection, numeric_farm_id, annual_context["productionYear"], annual_version_id)
        stage = connection.execute(text("SELECT * FROM farm_annual_stage_state WHERE farm_id=:farm AND production_year=:year AND stage_code=:stage FOR UPDATE"), {"farm": numeric_farm_id, "year": annual_context["productionYear"], "stage": request.stageCode}).mappings().first()
        if stage is None:
            raise HTTPException(status_code=404, detail="STAGE_NOT_FOUND")
        if stage["stage_status"] != "active":
            raise HTTPException(status_code=409, detail="STAGE_NOT_ACTIVE")
        plots = connection.execute(text("""SELECT p.plot_id,p.plot_code,p.plot_name,p.production_date_offset_days,COALESCE((SELECT MIN(fps.user_id) FROM farmer_plot_scope fps WHERE fps.plot_id=p.plot_id),p.farmer_user_id) AS farmer_user_id FROM plot p WHERE p.farm_id=:farm AND p.plot_status='active' ORDER BY p.plot_id FOR UPDATE"""), {"farm": numeric_farm_id}).mappings().all()
        unassigned_plots = [{"plotId": business_id("P", row["plot_id"]), "plotName": row["plot_name"]} for row in plots if row["farmer_user_id"] is None]
        assigned_plots = [row for row in plots if row["farmer_user_id"] is not None]
        plot_ids = [row["plot_id"] for row in assigned_plots]
        if not plots:
            return {"farmId": business_id("F", numeric_farm_id), "templateVersionId": annual_context["templateVersionId"], "productionYear": annual_context["productionYear"], "stageCode": request.stageCode, "track": annual_context["track"], "regenerated": request.regenerate, "plans": [], "taskCount": 0, "skippedPlots": []}
        if not plot_ids:
            raise HTTPException(status_code=409, detail={"code": "PLOT_ASSIGNEE_REQUIRED", "plots": unassigned_plots})
        active_statement = text("SELECT plan_id,plot_id,production_year,stage_code FROM management_plan WHERE plan_status='active' AND plot_id IN :plot_ids AND ((production_year=:year AND stage_code=:stage) OR production_year IS NULL) FOR UPDATE").bindparams(bindparam("plot_ids", expanding=True))
        previous_plans = connection.execute(active_statement, {"plot_ids": plot_ids, "year": annual_context["productionYear"], "stage": request.stageCode}).mappings().all()
        if previous_plans and not request.regenerate:
            raise HTTPException(status_code=409, detail="PLAN_EXISTS")
        stage_windows = {request.stageCode: (stage["actual_started_at"].date() if stage["actual_started_at"] else stage["planned_start_date"], stage["planned_end_date"])}
        task_rules = [rule for rule in farm_plan_task_rules(connection, annual_version_id, annual_context["productionYear"], annual_context["track"], stage_windows) if rule["stageCode"] == request.stageCode]
        generated: list[dict[str, Any]] = []
        replacement_by_plot: dict[int, int] = {}
        for plot_id in plot_ids:
            plot = next(row for row in assigned_plots if row["plot_id"] == plot_id)
            plan_id = connection.execute(text("INSERT INTO management_plan (plot_id,template_version_id,production_year,stage_code,plan_status,generated_by) VALUES (:plot,:version,:year,:stage,'active',:user)"), {"plot": plot_id, "version": annual_version_id, "year": annual_context["productionYear"], "stage": request.stageCode, "user": user["user_id"]}).lastrowid
            replacement_by_plot[plot_id] = plan_id
            for task in task_rules:
                scheduled_date = task["scheduledDate"] + timedelta(days=int(plot["production_date_offset_days"] or 0))
                if scheduled_date < today:
                    continue
                connection.execute(text("INSERT INTO management_task (plan_id,plot_id,stage_code,task_type,title,source,scheduled_date,task_status,note) VALUES (:plan,:plot,:stage,:type,:title,'template',:scheduled,'pending',:note)"), {"plan": plan_id, "plot": plot_id, "stage": task["stageCode"], "type": task["taskType"], "title": task["title"], "scheduled": scheduled_date, "note": task["note"]})
            task_count = sum(1 for task in task_rules if task["scheduledDate"] + timedelta(days=int(plot["production_date_offset_days"] or 0)) >= today)
            generated.append({"id": business_id("MP", plan_id), "plotId": business_id("P", plot_id), "taskCount": task_count, "dateOffsetDays": int(plot["production_date_offset_days"] or 0)})
        for old_plan in previous_plans:
            if old_plan["production_year"] is None:
                connection.execute(text("UPDATE management_task SET task_status='cancelled' WHERE plan_id=:plan AND stage_code=:stage AND task_status='pending'"), {"plan": old_plan["plan_id"], "stage": request.stageCode})
            else:
                connection.execute(text("UPDATE management_task SET task_status='cancelled' WHERE plan_id=:plan AND task_status='pending'"), {"plan": old_plan["plan_id"]})
                connection.execute(text("UPDATE management_plan SET plan_status='superseded', superseded_by=:replacement WHERE plan_id=:old"), {"replacement": replacement_by_plot[old_plan["plot_id"]], "old": old_plan["plan_id"]})
        connection.execute(text("UPDATE farm_annual_production SET annual_status='dispatched',updated_by=:user WHERE farm_id=:farm AND production_year=:year"), {"user": user["user_id"], "farm": numeric_farm_id, "year": annual_context["productionYear"]})
        connection.execute(text("UPDATE farm_annual_stage_state SET dispatched_template_version_id=:version,dispatched_at=:now WHERE farm_id=:farm AND production_year=:year AND stage_code=:stage"), {"version": annual_version_id, "now": datetime.now(UTC).replace(tzinfo=None), "farm": numeric_farm_id, "year": annual_context["productionYear"], "stage": request.stageCode})
        total_tasks = sum(item["taskCount"] for item in generated)
        record_audit_log(connection, user, "production_plan", "dispatch_stage", target_type="farm_stage", target_id=f"{business_id('F', numeric_farm_id)}:{annual_context['productionYear']}:{request.stageCode}", detail=f"Dispatched {total_tasks} tasks for stage {request.stageCode}")
    return {"farmId": business_id("F", numeric_farm_id), "templateVersionId": annual_context["templateVersionId"], "productionYear": annual_context["productionYear"], "stageCode": request.stageCode, "track": annual_context["track"], "regenerated": request.regenerate, "plans": generated, "taskCount": total_tasks, "skippedPlots": unassigned_plots}


@app.get("/api/tasks")
def list_management_tasks(scope: str = Query(default="today"), plot_id: str | None = Query(default=None, alias="plotId"), farm_id: str | None = Query(default=None, alias="farmId"), status_value: str | None = Query(default=None, alias="status"), month: str | None = None, user: dict[str, Any] = Depends(current_user)) -> list[dict[str, Any]]:
    if scope not in {"today", "overdue", "next7days", "month", "all"}:
        raise HTTPException(status_code=400, detail="INVALID_REQUEST")
    today = date.today()
    clauses, params = [], {"today": today}
    if scope == "today":
        clauses.append("t.scheduled_date=:today")
    elif scope == "overdue":
        clauses.append("t.scheduled_date<:today AND t.task_status IN ('pending','in_progress')")
    elif scope == "next7days":
        clauses.append("t.scheduled_date>=:today AND t.scheduled_date<=:end_date"); params["end_date"] = today + timedelta(days=7)
    elif scope == "month":
        try:
            month_start = date.fromisoformat(f"{month}-01") if month else None
            if month_start is None or month_start.day != 1: raise ValueError
            month_end = date(month_start.year + (month_start.month == 12), 1 if month_start.month == 12 else month_start.month + 1, 1)
        except ValueError as error:
            raise HTTPException(status_code=400, detail="INVALID_REQUEST") from error
        clauses.append("t.scheduled_date>=:month_start AND t.scheduled_date<:month_end"); params.update({"month_start": month_start, "month_end": month_end})
    if farm_id is not None:
        numeric_farm_id = parse_resource_id(farm_id, "F"); ensure_farm_access(user, numeric_farm_id); clauses.append("p.farm_id=:farm_id"); params["farm_id"] = numeric_farm_id
    if plot_id is not None:
        numeric_plot_id = parse_resource_id(plot_id, "P"); plot_scope(numeric_plot_id, user); clauses.append("t.plot_id=:plot_id"); params["plot_id"] = numeric_plot_id
    elif user["role"] == "farmer":
        clauses.append("(p.farmer_user_id=:user_id OR EXISTS (SELECT 1 FROM farmer_plot_scope fps WHERE fps.plot_id=p.plot_id AND fps.user_id=:user_id))"); params["user_id"] = user["user_id"]
    elif user["role"] == "farm_admin":
        farm_ids = accessible_farm_ids(user)
        if not farm_ids:
            return []
        clauses.append("p.farm_id IN :farm_ids"); params["farm_ids"] = list(farm_ids)
    if status_value:
        if status_value == "overdue": clauses.append("t.scheduled_date<:today AND t.task_status IN ('pending','in_progress')")
        elif status_value in {"pending", "in_progress", "completed", "cancelled"}: clauses.append("t.task_status=:task_status"); params["task_status"] = status_value
        else: raise HTTPException(status_code=400, detail="INVALID_REQUEST")
    clauses.append("(mp.plan_status IS NULL OR mp.plan_status='active' OR t.task_status IN ('completed','in_progress'))")
    statement = text("SELECT t.*,p.plot_name,p.farm_id,mp.plan_status,c.task_completion_id,c.result,c.executed_at,c.executor_user_id,c.executor_name,c.note AS completion_note,c.irrigation_record_id,c.record_type,c.corrects_completion_id FROM management_task t JOIN plot p ON p.plot_id=t.plot_id LEFT JOIN management_plan mp ON mp.plan_id=t.plan_id LEFT JOIN management_task_completion c ON c.task_completion_id=(SELECT c2.task_completion_id FROM management_task_completion c2 WHERE c2.task_id=t.task_id ORDER BY c2.task_completion_id DESC LIMIT 1) WHERE " + " AND ".join(clauses) + " ORDER BY t.scheduled_date,t.task_id")
    if "farm_ids" in params:
        statement = statement.bindparams(bindparam("farm_ids", expanding=True))
    with engine.connect() as connection:
        rows = connection.execute(statement, params).mappings().all()
    return [task_payload(row) for row in rows]


@app.get("/api/tasks/{task_id}")
def get_management_task(task_id: str, user: dict[str, Any] = Depends(current_user)) -> dict[str, Any]:
    numeric_task_id = parse_resource_id(task_id, "MT")
    with engine.connect() as connection:
        row = task_row(connection, numeric_task_id)
        plot_scope(row["plot_id"], user)
        adjustments = connection.execute(text("SELECT task_adjustment_id,changed_by,adjustment_type,before_value,after_value,created_at FROM management_task_adjustment WHERE task_id=:task ORDER BY task_adjustment_id"), {"task": numeric_task_id}).mappings().all()
        completions = connection.execute(text("SELECT task_completion_id,result,executed_at,executor_user_id,executor_name,note,irrigation_record_id,record_type,corrects_completion_id,created_at FROM management_task_completion WHERE task_id=:task ORDER BY task_completion_id"), {"task": numeric_task_id}).mappings().all()
    payload = task_payload(row)
    payload["adjustments"] = [{"id": business_id("TA", item["task_adjustment_id"]), "changedBy": business_id("U", item["changed_by"]), "type": item["adjustment_type"], "before": json.loads(item["before_value"]) if isinstance(item["before_value"], str) else item["before_value"], "after": json.loads(item["after_value"]) if isinstance(item["after_value"], str) else item["after_value"], "createdAt": iso(item["created_at"])} for item in adjustments]
    payload["completionHistory"] = [{"id": business_id("TC", item["task_completion_id"]), "result": item["result"], "executedAt": iso(item["executed_at"]), "executorUserId": business_id("U", item["executor_user_id"]), "executorName": item["executor_name"], "note": item["note"], "irrigationRecordId": business_id("IR", item["irrigation_record_id"]) if item["irrigation_record_id"] is not None else None, "recordType": item["record_type"], "correctsCompletionId": business_id("TC", item["corrects_completion_id"]) if item["corrects_completion_id"] is not None else None, "createdAt": iso(item["created_at"])} for item in completions]
    return payload


@app.patch("/api/tasks/{task_id}")
def update_management_task(task_id: str, request: TaskUpdateRequest, user: dict[str, Any] = Depends(current_user)) -> dict[str, Any]:
    numeric_task_id = parse_resource_id(task_id, "MT")
    fields = request.model_fields_set
    if not fields:
        raise HTTPException(status_code=400, detail="INVALID_REQUEST")
    if fields.intersection({"scheduledDate", "note"}):
        raise HTTPException(status_code=409, detail="TASK_CHANGE_ORDER_REQUIRED")
    with engine.begin() as connection:
        row = task_row(connection, numeric_task_id, lock=True)
        ensure_farm_admin_access(user, row["farm_id"])
        if row["task_status"] in {"completed", "cancelled"} or row["plan_status"] == "superseded": raise HTTPException(status_code=409, detail="TASK_NOT_ADJUSTABLE")
        before, after, updates = {}, {}, []
        if "scheduledDate" in fields:
            if request.scheduledDate is None: raise HTTPException(status_code=400, detail="INVALID_REQUEST")
            try: scheduled = date.fromisoformat(request.scheduledDate)
            except ValueError as error: raise HTTPException(status_code=400, detail="INVALID_REQUEST") from error
            before["scheduledDate"] = row["scheduled_date"].isoformat(); after["scheduledDate"] = scheduled.isoformat(); updates.append("scheduled_date=:scheduled")
        else: scheduled = None
        if "note" in fields:
            before["note"] = row["note"]; after["note"] = request.note; updates.append("note=:note")
        if "status" in fields:
            if request.status not in {"pending", "in_progress", "cancelled"}: raise HTTPException(status_code=400, detail="INVALID_REQUEST")
            before["status"] = task_status_value(row); after["status"] = request.status; updates.append("task_status=:status")
        connection.execute(text("UPDATE management_task SET " + ", ".join(updates) + " WHERE task_id=:task"), {"task": numeric_task_id, "scheduled": scheduled, "note": request.note, "status": request.status})
        record_task_adjustment(connection, numeric_task_id, user["user_id"], "task_update", before, after)
        record_audit_log(connection, user, "production_plan", "adjust_task", target_type="task", target_id=business_id("MT", numeric_task_id), detail="Adjusted an uncompleted task")
        return task_payload(task_row(connection, numeric_task_id))


@app.post("/api/tasks/{task_id}/change-orders", status_code=status.HTTP_201_CREATED)
def create_task_change_order(task_id: str, request: TaskChangeOrderRequest, user: dict[str, Any] = Depends(current_user)) -> dict[str, Any]:
    numeric_task_id = parse_resource_id(task_id, "MT")
    with engine.begin() as connection:
        row = task_row(connection, numeric_task_id, lock=True)
        ensure_farm_admin_access(user, row["farm_id"])
        if row["task_status"] != "pending" or row["plan_status"] == "superseded":
            raise HTTPException(status_code=409, detail="TASK_CHANGE_ORDER_NOT_ALLOWED")
        production_year = production_year_for(date.today())
        stage = connection.execute(text("SELECT stage_status,planned_start_date,planned_end_date FROM farm_annual_stage_state WHERE farm_id=:farm AND production_year=:year AND stage_code=:stage FOR UPDATE"), {"farm": row["farm_id"], "year": production_year, "stage": row["stage_code"]}).mappings().first()
        if stage is None or stage["stage_status"] != "active":
            raise HTTPException(status_code=409, detail="ACTIVE_STAGE_REQUIRED")
        try:
            scheduled_date = date.fromisoformat(request.scheduledDate) if request.scheduledDate else row["scheduled_date"]
        except ValueError as error:
            raise HTTPException(status_code=400, detail="INVALID_REQUEST") from error
        if scheduled_date < date.today() or scheduled_date > stage["planned_end_date"]:
            raise HTTPException(status_code=400, detail="INVALID_CHANGE_ORDER_DATE")
        replacement_note = request.note if "note" in request.model_fields_set else row["note"]
        if scheduled_date == row["scheduled_date"] and replacement_note == row["note"]:
            raise HTTPException(status_code=400, detail="CHANGE_ORDER_HAS_NO_CHANGES")
        replacement_id = connection.execute(text("INSERT INTO management_task (plan_id,plot_id,stage_code,task_type,title,source,scheduled_date,task_status,note) VALUES (:plan,:plot,:stage,:type,:title,:source,:scheduled,'pending',:note)"), {"plan": row["plan_id"], "plot": row["plot_id"], "stage": row["stage_code"], "type": row["task_type"], "title": row["title"], "source": row["source"], "scheduled": scheduled_date, "note": replacement_note}).lastrowid
        connection.execute(text("UPDATE management_task SET task_status='cancelled' WHERE task_id=:task"), {"task": numeric_task_id})
        before = {"scheduledDate": row["scheduled_date"].isoformat(), "note": row["note"], "status": "pending"}
        after = {"replacementTaskId": business_id("MT", replacement_id), "scheduledDate": scheduled_date.isoformat(), "note": replacement_note, "reason": request.reason.strip()}
        record_task_adjustment(connection, numeric_task_id, user["user_id"], "change_order", before, after)
        record_audit_log(connection, user, "production_plan", "create_change_order", target_type="task", target_id=business_id("MT", numeric_task_id), detail=f"Replaced by {business_id('MT', replacement_id)}: {request.reason.strip()}")
        return {"originalTaskId": business_id("MT", numeric_task_id), "replacementTask": task_payload(task_row(connection, replacement_id)), "reason": request.reason.strip()}


@app.post("/api/farms/{farm_id}/temporary-tasks", status_code=status.HTTP_201_CREATED)
def create_temporary_management_task(farm_id: str, request: TemporaryTaskCreateRequest, user: dict[str, Any] = Depends(current_user)) -> dict[str, Any]:
    numeric_farm_id = parse_resource_id(farm_id, "F")
    numeric_plot_id = parse_resource_id(request.plotId, "P")
    ensure_farm_admin_access(user, numeric_farm_id)
    try:
        scheduled_date = date.fromisoformat(request.scheduledDate)
    except ValueError as error:
        raise HTTPException(status_code=400, detail="INVALID_REQUEST") from error
    with engine.begin() as connection:
        plot = connection.execute(text("SELECT plot_id FROM plot WHERE plot_id=:plot AND farm_id=:farm AND plot_status='active' FOR UPDATE"), {"plot": numeric_plot_id, "farm": numeric_farm_id}).mappings().first()
        if plot is None:
            raise HTTPException(status_code=404, detail="PLOT_NOT_FOUND")
        production_year = production_year_for(date.today())
        stage = connection.execute(text("SELECT stage_code,planned_end_date FROM farm_annual_stage_state WHERE farm_id=:farm AND production_year=:year AND stage_status='active' FOR UPDATE"), {"farm": numeric_farm_id, "year": production_year}).mappings().first()
        if stage is None:
            raise HTTPException(status_code=409, detail="ACTIVE_STAGE_REQUIRED")
        if scheduled_date < date.today() or scheduled_date > stage["planned_end_date"]:
            raise HTTPException(status_code=400, detail="INVALID_TEMPORARY_TASK_DATE")
        task_id = connection.execute(text("INSERT INTO management_task (plan_id,plot_id,stage_code,task_type,title,source,scheduled_date,task_status,note) VALUES (NULL,:plot,:stage,:type,:title,'temporary',:scheduled,'pending',:note)"), {"plot": numeric_plot_id, "stage": stage["stage_code"], "type": request.taskType.strip(), "title": request.title.strip(), "scheduled": scheduled_date, "note": request.note}).lastrowid
        record_task_adjustment(connection, task_id, user["user_id"], "temporary_task", {}, {"reason": request.reason.strip()})
        record_audit_log(connection, user, "production_plan", "create_temporary_task", target_type="task", target_id=business_id("MT", task_id), detail=request.reason.strip())
        return task_payload(task_row(connection, task_id))


@app.post("/api/tasks/{task_id}/completions", status_code=status.HTTP_201_CREATED)
def complete_management_task(task_id: str, request: TaskCompletionRequest, user: dict[str, Any] = Depends(current_user)) -> dict[str, Any]:
    numeric_task_id = parse_resource_id(task_id, "MT")
    try: executed_at = utc_naive(request.executedAt)
    except ValueError as error: raise HTTPException(status_code=400, detail="INVALID_REQUEST") from error
    with engine.begin() as connection:
        row = task_row(connection, numeric_task_id, lock=True)
        ensure_farmer_task_access(row["plot_id"], user)
        if row["task_status"] in {"completed", "cancelled"} or row["plan_status"] == "superseded": raise HTTPException(status_code=409, detail="TASK_NOT_ADJUSTABLE")
        irrigation_record_id = task_irrigation_record_id(connection, request.irrigationRecordId, row["plot_id"])
        connection.execute(text("UPDATE management_task SET task_status='completed' WHERE task_id=:task"), {"task": numeric_task_id})
        completion_id = connection.execute(text("INSERT INTO management_task_completion (task_id,result,executed_at,executor_user_id,executor_name,note,irrigation_record_id) VALUES (:task,:result,:executed,:user,:name,:note,:irrigation)"), {"task": numeric_task_id, "result": request.result.strip(), "executed": executed_at, "user": user["user_id"], "name": request.executorName.strip(), "note": request.note, "irrigation": irrigation_record_id}).lastrowid
        finalize_farm_planting_if_ready(connection, row, executed_at.date(), user)
        record_task_adjustment(connection, numeric_task_id, user["user_id"], "completion", {"status": task_status_value(row)}, {"status": "completed", "completionId": business_id("TC", completion_id)})
        record_audit_log(connection, user, "task_execution", "complete_task", target_type="task", target_id=business_id("MT", numeric_task_id), detail=f"Recorded completion {business_id('TC', completion_id)}")
        return task_payload(task_row(connection, numeric_task_id))


@app.post("/api/tasks/{task_id}/completion-corrections", status_code=status.HTTP_201_CREATED)
def correct_management_task_completion(task_id: str, request: TaskCompletionCorrectionRequest, user: dict[str, Any] = Depends(current_user)) -> dict[str, Any]:
    """Append a correction without overwriting the original execution record."""
    numeric_task_id = parse_resource_id(task_id, "MT")
    try:
        executed_at = utc_naive(request.executedAt)
    except ValueError as error:
        raise HTTPException(status_code=400, detail="INVALID_REQUEST") from error
    with engine.begin() as connection:
        row = task_row(connection, numeric_task_id, lock=True)
        ensure_farmer_task_access(row["plot_id"], user)
        if row["task_status"] != "completed":
            raise HTTPException(status_code=409, detail="TASK_COMPLETION_CORRECTION_NOT_ALLOWED")
        previous = connection.execute(
            text("SELECT task_completion_id FROM management_task_completion WHERE task_id=:task ORDER BY task_completion_id DESC LIMIT 1 FOR UPDATE"),
            {"task": numeric_task_id},
        ).mappings().first()
        if previous is None:
            raise HTTPException(status_code=409, detail="TASK_COMPLETION_NOT_FOUND")
        irrigation_record_id = task_irrigation_record_id(connection, request.irrigationRecordId, row["plot_id"])
        correction_id = connection.execute(
            text("INSERT INTO management_task_completion (task_id,result,executed_at,executor_user_id,executor_name,note,irrigation_record_id,record_type,corrects_completion_id) VALUES (:task,:result,:executed,:user,:name,:note,:irrigation,'correction',:previous)"),
            {"task": numeric_task_id, "result": request.result.strip(), "executed": executed_at, "user": user["user_id"], "name": request.executorName.strip(), "note": request.note, "irrigation": irrigation_record_id, "previous": previous["task_completion_id"]},
        ).lastrowid
        record_task_adjustment(connection, numeric_task_id, user["user_id"], "completion_correction", {"completionId": business_id("TC", previous["task_completion_id"])}, {"completionId": business_id("TC", correction_id), "reason": request.reason.strip()})
        record_audit_log(connection, user, "task_execution", "correct_task_completion", target_type="task", target_id=business_id("MT", numeric_task_id), detail=f"Appended correction {business_id('TC', correction_id)}: {request.reason.strip()}")
        return task_payload(task_row(connection, numeric_task_id))


@app.post("/api/plots/{plot_id}/ad-hoc-tasks", status_code=status.HTTP_201_CREATED)
def create_ad_hoc_task(plot_id: str, request: AdHocTaskRequest, user: dict[str, Any] = Depends(current_user)) -> dict[str, Any]:
    numeric_plot_id = parse_resource_id(plot_id, "P")
    try: executed_at = utc_naive(request.executedAt)
    except ValueError as error: raise HTTPException(status_code=400, detail="INVALID_REQUEST") from error
    with engine.begin() as connection:
        ensure_farmer_task_access(numeric_plot_id, user)
        irrigation_record_id = task_irrigation_record_id(connection, request.irrigationRecordId, numeric_plot_id)
        task_id = connection.execute(text("INSERT INTO management_task (plan_id,plot_id,stage_code,task_type,title,source,scheduled_date,task_status,note) VALUES (NULL,:plot,NULL,:type,:title,'ad_hoc',:scheduled,'completed',:note)"), {"plot": numeric_plot_id, "type": request.taskType.strip(), "title": request.title.strip(), "scheduled": executed_at.date(), "note": request.note}).lastrowid
        connection.execute(text("INSERT INTO management_task_completion (task_id,result,executed_at,executor_user_id,executor_name,note,irrigation_record_id) VALUES (:task,'completed',:executed,:user,:name,:note,:irrigation)"), {"task": task_id, "executed": executed_at, "user": user["user_id"], "name": request.executorName.strip(), "note": request.note, "irrigation": irrigation_record_id})
        record_audit_log(connection, user, "task_execution", "create_ad_hoc_task", target_type="task", target_id=business_id("MT", task_id), detail=f"Added completed ad-hoc task for {business_id('P', numeric_plot_id)}")
        return task_payload(task_row(connection, task_id))


@app.put("/api/farms/{farm_id}/production-profile")
def save_farm_production_profile(farm_id: str, request: FarmProductionProfileRequest, user: dict[str, Any] = Depends(current_user)) -> dict[str, Any]:
    numeric_farm_id = parse_resource_id(farm_id, "F")
    ensure_farm_admin_access(user, numeric_farm_id)
    version_id = parse_resource_id(request.templateVersionId, "CTV")
    try:
        requested_planting_date = date.fromisoformat(request.plantingDate) if request.plantingDate else None
    except ValueError as error: raise HTTPException(status_code=400, detail="INVALID_REQUEST") from error
    if requested_planting_date is not None and requested_planting_date > date.today():
        raise HTTPException(status_code=400, detail="ACTUAL_PLANTING_DATE_CANNOT_BE_FUTURE")
    with engine.begin() as connection:
        farm = connection.execute(text("SELECT planting_date FROM farm WHERE farm_id=:id"), {"id":numeric_farm_id}).mappings().first()
        if farm is None: raise HTTPException(status_code=404, detail="FARM_NOT_FOUND")
        if farm["planting_date"] is None:
            raise HTTPException(status_code=409, detail="PLANTING_TASK_REQUIRED")
        planting_date = requested_planting_date or farm["planting_date"]
        if planting_date is None: raise HTTPException(status_code=409, detail="ORCHARD_PLANTING_DATE_REQUIRED")
        if not connection.execute(text("SELECT 1 FROM crop_template_version WHERE template_version_id=:id AND version_status='published' AND (farm_id IS NULL OR farm_id=:farm)"), {"id":version_id, "farm": numeric_farm_id}).scalar(): raise HTTPException(status_code=404, detail="TEMPLATE_NOT_FOUND")
        if requested_planting_date is not None:
            connection.execute(text("UPDATE farm SET planting_date=:planting WHERE farm_id=:farm"), {"planting": planting_date, "farm": numeric_farm_id})
            annual_rows = connection.execute(text("SELECT annual_production_id,production_year,track_source FROM farm_annual_production WHERE farm_id=:farm FOR UPDATE"), {"farm": numeric_farm_id}).mappings().all()
            for annual in annual_rows:
                tree_age_years, automatic_track = orchard_track(planting_date, annual["production_year"])
                keep_extension = annual["track_source"] == "extended_young" and automatic_track == "fruiting" and tree_age_years == 4
                connection.execute(text("UPDATE farm_annual_production SET orchard_track=:track,track_source=:source,track_override_reason=CASE WHEN :keep_extension THEN track_override_reason ELSE NULL END WHERE annual_production_id=:id"), {"track": "young" if keep_extension else automatic_track, "source": "extended_young" if keep_extension else "auto", "keep_extension": keep_extension, "id": annual["annual_production_id"]})
        connection.execute(text("INSERT INTO farm_production_profile (farm_id,template_version_id,planting_date,tree_age_years,note,updated_by) VALUES (:farm,:version,:planting,:age,:note,:user) ON DUPLICATE KEY UPDATE template_version_id=VALUES(template_version_id),planting_date=VALUES(planting_date),tree_age_years=VALUES(tree_age_years),note=VALUES(note),updated_by=VALUES(updated_by)"), {"farm":numeric_farm_id,"version":version_id,"planting":planting_date,"age":None,"note":request.note,"user":user["user_id"]})
        connection.execute(text("UPDATE farm_annual_production SET template_version_id=:version,updated_by=:user WHERE farm_id=:farm AND production_year=:year"), {"version": version_id, "user": user["user_id"], "farm": numeric_farm_id, "year": production_year_for(date.today())})
    return {"farmId":business_id("F",numeric_farm_id),"templateVersionId":business_id("CTV",version_id),"plantingDate":planting_date.isoformat()}


def record_farm_production_milestone(farm_id: str, completed_at: datetime, user: dict[str, Any], *, column: str, action: str, response_key: str) -> dict[str, Any]:
    numeric_farm_id = parse_resource_id(farm_id, "F")
    ensure_farm_admin_access(user, numeric_farm_id)
    if completed_at.tzinfo is not None:
        completed_at = completed_at.astimezone().replace(tzinfo=None)
    now = datetime.now()
    if completed_at > now:
        raise HTTPException(status_code=400, detail="INVALID_PRODUCTION_MILESTONE_TIME")
    with engine.begin() as connection:
        planting_date = connection.execute(text("SELECT planting_date FROM farm WHERE farm_id=:farm FOR UPDATE"), {"farm": numeric_farm_id}).scalar()
        if planting_date is None:
            raise HTTPException(status_code=409, detail="ORCHARD_PLANTING_DATE_REQUIRED")
        if completed_at.date() < planting_date:
            raise HTTPException(status_code=400, detail="INVALID_PRODUCTION_MILESTONE_TIME")
        result = connection.execute(text(f"UPDATE farm_production_profile SET {column}=:completed,next_cycle_starts_on=NULL,updated_by=:user WHERE farm_id=:farm"), {"completed": completed_at, "user": user["user_id"], "farm": numeric_farm_id})
        if result.rowcount == 0:
            raise HTTPException(status_code=409, detail="PRODUCTION_TEMPLATE_REQUIRED")
        record_audit_log(connection, user, "production_plan", action, target_type="farm", target_id=business_id("F", numeric_farm_id), detail=f"Recorded production milestone at {completed_at.isoformat()}")
    return {"farmId": business_id("F", numeric_farm_id), response_key: iso(completed_at)}


@app.put("/api/farms/{farm_id}/harvest-completion")
def confirm_farm_harvest_completion(farm_id: str, request: FarmMilestoneCompletionRequest, user: dict[str, Any] = Depends(current_user)) -> dict[str, Any]:
    return record_farm_production_milestone(farm_id, request.completedAt, user, column="harvest_completed_at", action="record_harvest_completion", response_key="harvestCompletedAt")


@app.put("/api/farms/{farm_id}/cleanup-completion")
def confirm_farm_cleanup_completion(farm_id: str, request: FarmMilestoneCompletionRequest, user: dict[str, Any] = Depends(current_user)) -> dict[str, Any]:
    return record_farm_production_milestone(farm_id, request.completedAt, user, column="harvest_cleanup_completed_at", action="record_cleanup_completion", response_key="cleanupCompletedAt")


@app.put("/api/farms/{farm_id}/harvest-cleanup")
def confirm_farm_harvest_cleanup(farm_id: str, request: FarmMilestoneCompletionRequest, user: dict[str, Any] = Depends(current_user)) -> dict[str, Any]:
    """Backward-compatible cleanup endpoint."""
    return record_farm_production_milestone(farm_id, request.completedAt, user, column="harvest_cleanup_completed_at", action="record_cleanup_completion", response_key="cleanupCompletedAt")


@app.get("/api/farms/{farm_id}/production-profile")
def get_farm_production_profile(farm_id: str, user: dict[str, Any] = Depends(current_user)) -> dict[str, Any]:
    numeric_farm_id = parse_resource_id(farm_id, "F")
    ensure_farm_access(user, numeric_farm_id)
    with engine.connect() as connection:
        profile = connection.execute(text("SELECT p.*,f.planting_date,t.template_name,t.crop_name,v.version_no FROM farm_production_profile p JOIN farm f ON f.farm_id=p.farm_id JOIN crop_template_version v ON v.template_version_id=p.template_version_id JOIN crop_template t ON t.template_id=v.template_id WHERE p.farm_id=:farm"), {"farm": numeric_farm_id}).mappings().first()
    if profile is None:
        raise HTTPException(status_code=404, detail="PLANTING_PROFILE_NOT_FOUND")
    return {"farmId": business_id("F", profile["farm_id"]), "templateVersionId": business_id("CTV", profile["template_version_id"]), "templateName": profile["template_name"], "crop": profile["crop_name"], "versionNo": profile["version_no"], "plantingDate": profile["planting_date"].isoformat() if profile["planting_date"] else None, "harvestCompletedAt": iso(profile["harvest_completed_at"]), "cleanupCompletedAt": iso(profile["harvest_cleanup_completed_at"]), "note": profile["note"], "updatedAt": iso(profile["updated_at"])}


@app.get("/api/farms/{farm_id}/annual-production")
def get_farm_annual_production(farm_id: str, production_year: int | None = Query(default=None, alias="productionYear"), user: dict[str, Any] = Depends(current_user)) -> dict[str, Any]:
    numeric_farm_id = parse_resource_id(farm_id, "F")
    ensure_farm_access(user, numeric_farm_id)
    with engine.connect() as connection:
        return annual_production_context(connection, numeric_farm_id, production_year=production_year)


@app.put("/api/farms/{farm_id}/annual-production")
def save_farm_annual_production(farm_id: str, request: FarmAnnualProductionRequest, user: dict[str, Any] = Depends(current_user)) -> dict[str, Any]:
    numeric_farm_id = parse_resource_id(farm_id, "F")
    ensure_farm_admin_access(user, numeric_farm_id)
    production_year = request.productionYear or production_year_for(date.today())
    with engine.begin() as connection:
        farm = connection.execute(text("SELECT planting_date FROM farm WHERE farm_id=:farm FOR UPDATE"), {"farm": numeric_farm_id}).mappings().first()
        if farm is None:
            raise HTTPException(status_code=404, detail="FARM_NOT_FOUND")
        if farm["planting_date"] is None:
            raise HTTPException(status_code=409, detail="ORCHARD_PLANTING_DATE_REQUIRED")
        profile = connection.execute(text("SELECT template_version_id FROM farm_production_profile WHERE farm_id=:farm"), {"farm": numeric_farm_id}).mappings().first()
        existing_annual = connection.execute(text("SELECT template_version_id,orchard_track FROM farm_annual_production WHERE farm_id=:farm AND production_year=:year FOR UPDATE"), {"farm": numeric_farm_id, "year": production_year}).mappings().first()
        version_id = parse_resource_id(request.templateVersionId, "CTV") if request.templateVersionId else (profile["template_version_id"] if profile else None)
        selected_version = None if version_id is None else connection.execute(text("SELECT template_version_id,template_id FROM crop_template_version WHERE template_version_id=:version AND version_status='published' AND (farm_id IS NULL OR farm_id=:farm) FOR UPDATE"), {"version": version_id, "farm": numeric_farm_id}).mappings().first()
        if selected_version is None:
            raise HTTPException(status_code=409, detail="PRODUCTION_TEMPLATE_REQUIRED")
        if request.versionDefinition is not None:
            validate_template_definition(request.versionDefinition)
            if existing_annual is not None:
                ensure_locked_stages_unchanged(connection, numeric_farm_id, production_year, existing_annual["template_version_id"], request.versionDefinition)
            version_no = connection.execute(text("SELECT COALESCE(MAX(version_no),0)+1 FROM crop_template_version WHERE template_id=:template"), {"template": selected_version["template_id"]}).scalar_one()
            snapshot = connection.execute(text("INSERT INTO crop_template_version (template_id,farm_id,version_no,version_status,change_note,published_at,published_by) VALUES (:template,:farm,:version,'published',:note,:now,:user)"), {"template": selected_version["template_id"], "farm": numeric_farm_id, "version": version_no, "note": request.changeNote.strip() if request.changeNote else None, "now": datetime.now(UTC).replace(tzinfo=None), "user": user["user_id"]})
            version_id = snapshot.lastrowid
            store_template_definition(connection, version_id, request.versionDefinition)
        tree_age_years, automatic_track = orchard_track(farm["planting_date"], production_year)
        if request.orchardTrack is not None:
            if request.orchardTrack not in {"young", "fruiting"}:
                raise HTTPException(status_code=400, detail="INVALID_ORCHARD_TRACK")
            track, source, reason = request.orchardTrack, "selected_template", None
        else:
            can_extend = automatic_track == "fruiting" and tree_age_years == 4
            if request.extendYoungTrack and not can_extend:
                raise HTTPException(status_code=409, detail="YOUNG_TRACK_EXTENSION_NOT_ALLOWED")
            if request.extendYoungTrack and not request.trackOverrideReason:
                raise HTTPException(status_code=400, detail="YOUNG_TRACK_EXTENSION_REASON_REQUIRED")
            track = "young" if request.extendYoungTrack else automatic_track
            source = "extended_young" if request.extendYoungTrack else "auto"
            reason = request.trackOverrideReason.strip() if request.trackOverrideReason else None
        if existing_annual is not None and locked_stage_codes(connection, numeric_farm_id, production_year) and track != existing_annual["orchard_track"]:
            raise HTTPException(status_code=409, detail="LOCKED_STAGE_TRACK_CHANGE")
        annual_status = "saved" if request.versionDefinition is not None else "published"
        connection.execute(text("INSERT INTO farm_annual_production (farm_id,production_year,template_version_id,orchard_track,track_source,track_override_reason,annual_status,created_by,updated_by) VALUES (:farm,:year,:version,:track,:source,:reason,:status,:user,:user) ON DUPLICATE KEY UPDATE template_version_id=VALUES(template_version_id),orchard_track=VALUES(orchard_track),track_source=VALUES(track_source),track_override_reason=VALUES(track_override_reason),annual_status=VALUES(annual_status),updated_by=VALUES(updated_by)"), {"farm": numeric_farm_id, "year": production_year, "version": version_id, "track": track, "source": source, "reason": reason, "status": annual_status, "user": user["user_id"]})
        ensure_annual_stage_states(connection, numeric_farm_id, production_year, version_id)
        refresh_unlocked_stage_windows(connection, numeric_farm_id, production_year, version_id)
        if request.versionDefinition is not None:
            connection.execute(text("UPDATE farm_production_profile SET template_version_id=:version,updated_by=:user WHERE farm_id=:farm"), {"version": version_id, "user": user["user_id"], "farm": numeric_farm_id})
        record_audit_log(connection, user, "production_plan", "save_annual_production", target_type="farm", target_id=business_id("F", numeric_farm_id), detail=f"Saved the independent {production_year} farm management plan using track {track}")
        return annual_production_context(connection, numeric_farm_id, production_year=production_year)


@app.get("/api/farms/{farm_id}/planting-schedule")
def get_farm_planting_schedule(farm_id: str, user: dict[str, Any] = Depends(current_user)) -> dict[str, Any]:
    numeric_farm_id = parse_resource_id(farm_id, "F")
    ensure_farm_access(user, numeric_farm_id)
    with engine.connect() as connection:
        row = connection.execute(text("SELECT f.planting_date,f.planned_planting_date,f.planting_template_version_id,t.template_name,v.version_no FROM farm f LEFT JOIN crop_template_version v ON v.template_version_id=f.planting_template_version_id LEFT JOIN crop_template t ON t.template_id=v.template_id WHERE f.farm_id=:farm"), {"farm": numeric_farm_id}).mappings().one()
        pending_count = connection.execute(text("SELECT COUNT(*) FROM management_task mt JOIN plot p ON p.plot_id=mt.plot_id WHERE p.farm_id=:farm AND mt.source='planting' AND mt.task_status IN ('pending','in_progress')"), {"farm": numeric_farm_id}).scalar()
    return {"farmId": business_id("F", numeric_farm_id), "plantingDate": row["planting_date"].isoformat() if row["planting_date"] else None, "plannedPlantingDate": row["planned_planting_date"].isoformat() if row["planned_planting_date"] else None, "templateVersionId": business_id("CTV", row["planting_template_version_id"]) if row["planting_template_version_id"] else None, "templateName": row["template_name"], "versionNo": row["version_no"], "pendingTaskCount": int(pending_count)}


@app.put("/api/farms/{farm_id}/planting-schedule")
def save_farm_planting_schedule(farm_id: str, request: FarmPlantingScheduleRequest, user: dict[str, Any] = Depends(current_user)) -> dict[str, Any]:
    numeric_farm_id = parse_resource_id(farm_id, "F")
    ensure_farm_admin_access(user, numeric_farm_id)
    version_id = parse_resource_id(request.templateVersionId, "CTV")
    try:
        planned_date = date.fromisoformat(request.plannedPlantingDate)
    except ValueError as error:
        raise HTTPException(status_code=400, detail="INVALID_REQUEST") from error
    if planned_date < date.today():
        raise HTTPException(status_code=400, detail="PLANNED_PLANTING_DATE_CANNOT_BE_PAST")
    with engine.begin() as connection:
        farm = connection.execute(text("SELECT planting_date FROM farm WHERE farm_id=:farm FOR UPDATE"), {"farm": numeric_farm_id}).mappings().first()
        if farm is None:
            raise HTTPException(status_code=404, detail="FARM_NOT_FOUND")
        if farm["planting_date"] is not None:
            raise HTTPException(status_code=409, detail="PLANTING_ALREADY_CONFIRMED")
        if not connection.execute(text("SELECT 1 FROM crop_template_version WHERE template_version_id=:version AND version_status='published' AND farm_id IS NULL"), {"version": version_id}).scalar():
            raise HTTPException(status_code=404, detail="TEMPLATE_NOT_FOUND")
        plots = connection.execute(text("SELECT plot_id FROM plot WHERE farm_id=:farm AND plot_status='active' ORDER BY plot_id FOR UPDATE"), {"farm": numeric_farm_id}).mappings().all()
        if not plots:
            raise HTTPException(status_code=409, detail="ACTIVE_PLOT_REQUIRED")
        connection.execute(text("UPDATE management_task mt JOIN plot p ON p.plot_id=mt.plot_id SET mt.task_status='cancelled' WHERE p.farm_id=:farm AND mt.source='planting' AND mt.task_status IN ('pending','in_progress')"), {"farm": numeric_farm_id})
        connection.execute(text("UPDATE farm SET planned_planting_date=:planned,planting_template_version_id=:version WHERE farm_id=:farm"), {"planned": planned_date, "version": version_id, "farm": numeric_farm_id})
        for plot in plots:
            connection.execute(text("INSERT INTO management_task (plan_id,plot_id,stage_code,task_type,title,source,scheduled_date,task_status,note) VALUES (NULL,:plot,'planting','planting','完成脐橙定植','planting',:scheduled,'pending','按定植方案完成苗木栽植，并记录现场结果。')"), {"plot": plot["plot_id"], "scheduled": planned_date})
        record_audit_log(connection, user, "production_plan", "schedule_farm_planting", target_type="farm", target_id=business_id("F", numeric_farm_id), detail=f"Scheduled planting tasks for {planned_date.isoformat()}")
    return get_farm_planting_schedule(farm_id, user)


@app.get("/api/plots/{plot_id}/production-context")
def get_plot_production_context(plot_id: str, user: dict[str, Any] = Depends(current_user)) -> dict[str, Any]:
    """Return the authorized plot's shared farm-cycle context for task explanation."""
    numeric_plot_id = parse_resource_id(plot_id, "P")
    plot = plot_scope(numeric_plot_id, user)
    with engine.connect() as connection:
        return farm_production_context(connection, plot)


@app.post("/api/farms", status_code=status.HTTP_201_CREATED)
def create_farm(request: FarmSaveRequest, user: dict[str, Any] = Depends(current_user)) -> dict[str, Any]:
    ensure_system_admin(user)
    with engine.begin() as connection:
        result = connection.execute(text("INSERT INTO farm (farm_name, province, city, address, area_mu, manager_name, note) VALUES (:name, :province, :city, :address, :area, :manager, :note)"), {"name": request.name, "province": request.province, "city": request.city, "address": request.address, "area": request.areaMu, "manager": request.managerName, "note": request.note})
        row = connection.execute(text("SELECT * FROM farm WHERE farm_id = :id"), {"id": result.lastrowid}).mappings().one()
        record_audit_log(connection, user, "master_data", "create_farm", target_type="farm", target_id=business_id("F", result.lastrowid), detail=f"Created farm {request.name}")
        return farm_payload(row)


@app.get("/api/farms/{farm_id}")
def get_farm(farm_id: str, user: dict[str, Any] = Depends(current_user)) -> dict[str, Any]:
    numeric_id = parse_resource_id(farm_id, "F"); ensure_farm_access(user, numeric_id)
    with engine.connect() as connection:
        row = connection.execute(text("SELECT * FROM farm WHERE farm_id = :id"), {"id": numeric_id}).mappings().first()
        if row is None: raise HTTPException(status_code=404, detail="FARM_NOT_FOUND")
        return farm_payload(row)


@app.patch("/api/farms/{farm_id}")
def update_farm(farm_id: str, request: FarmUpdateRequest, user: dict[str, Any] = Depends(current_user)) -> dict[str, Any]:
    ensure_system_admin(user); numeric_id = parse_resource_id(farm_id, "F")
    if request.status is not None and request.status not in {"active", "inactive"}: raise HTTPException(status_code=400, detail="INVALID_REQUEST")
    with engine.begin() as connection:
        result = connection.execute(text("UPDATE farm SET farm_name=:name, province=:province, city=:city, address=:address, area_mu=:area, manager_name=:manager, note=:note, farm_status=COALESCE(:status, farm_status) WHERE farm_id=:id"), {"id": numeric_id, "name": request.name, "province": request.province, "city": request.city, "address": request.address, "area": request.areaMu, "manager": request.managerName, "note": request.note, "status": request.status})
        if result.rowcount == 0: raise HTTPException(status_code=404, detail="FARM_NOT_FOUND")
        record_audit_log(connection, user, "master_data", "update_farm", target_type="farm", target_id=business_id("F", numeric_id), detail="Updated farm profile or status")
        return farm_payload(connection.execute(text("SELECT * FROM farm WHERE farm_id=:id"), {"id": numeric_id}).mappings().one())


@app.get("/api/plots")
def list_plots(page: int = Query(default=1, ge=1), page_size: int = Query(default=20, ge=1, le=100, alias="pageSize"), farm_id: str | None = Query(default=None, alias="farmId"), keyword: str | None = None, status_value: str | None = Query(default=None, alias="status"), user: dict[str, Any] = Depends(current_user)) -> dict[str, Any]:
    clauses, params = [], {"limit": page_size, "offset": (page - 1) * page_size, "user_id": user["user_id"]}
    if user["role"] == "farmer": clauses.append("EXISTS (SELECT 1 FROM farmer_plot_scope s WHERE s.plot_id=p.plot_id AND s.user_id=:user_id)")
    elif user["role"] == "farm_admin": clauses.append("EXISTS (SELECT 1 FROM farm_admin_scope s WHERE s.farm_id=p.farm_id AND s.user_id=:user_id)")
    if farm_id: clauses.append("p.farm_id=:farm_id"); params["farm_id"] = parse_resource_id(farm_id, "F")
    if status_value: clauses.append("p.plot_status=:status"); params["status"] = status_value
    if keyword: clauses.append("(p.plot_code LIKE :keyword OR p.plot_name LIKE :keyword)"); params["keyword"] = f"%{keyword}%"
    where = f"WHERE {' AND '.join(clauses)}" if clauses else ""
    select = "SELECT p.*, s.user_id AS scoped_farmer_user_id, u.real_name AS farmer_name FROM plot p LEFT JOIN farmer_plot_scope s ON s.plot_id=p.plot_id LEFT JOIN users u ON u.user_id=s.user_id"
    with engine.connect() as connection:
        total = connection.execute(text(f"SELECT COUNT(*) FROM plot p {where}"), params).scalar_one()
        rows = connection.execute(text(f"{select} {where} ORDER BY p.plot_id LIMIT :limit OFFSET :offset"), params).mappings().all()
        return {"items": [plot_payload(row) for row in rows], "page": page, "pageSize": page_size, "total": total}


@app.post("/api/plots", status_code=status.HTTP_201_CREATED)
def create_plot(request: PlotSaveRequest, user: dict[str, Any] = Depends(current_user)) -> dict[str, Any]:
    ensure_system_admin(user)
    if request.status not in {"active", "inactive"} or request.irrigationMethod not in {None, "drip", "sprinkler", "flood", "manual", "other"}: raise HTTPException(status_code=400, detail="INVALID_REQUEST")
    with engine.begin() as connection:
        ensure_existing_resources(connection, [parse_resource_id(request.farmId, "F")], [])
        try: result = connection.execute(text("INSERT INTO plot (farm_id, plot_code, plot_name, area_mu, location, irrigation_method, plot_status) VALUES (:farm,:code,:name,:area,:location,:method,:status)"), {"farm": parse_resource_id(request.farmId, "F"), "code": request.code, "name": request.name, "area": request.areaMu, "location": request.location, "method": request.irrigationMethod, "status": request.status})
        except IntegrityError as error: raise HTTPException(status_code=409, detail="PLOT_CODE_EXISTS") from error
        record_audit_log(connection, user, "master_data", "create_plot", target_type="plot", target_id=business_id("P", result.lastrowid), detail=f"Created plot {request.name}")
        row = connection.execute(text("SELECT p.*, NULL AS scoped_farmer_user_id, NULL AS farmer_name FROM plot p WHERE p.plot_id=:id"), {"id": result.lastrowid}).mappings().one(); return plot_payload(row)


@app.get("/api/plots/{plot_id}")
def get_plot(plot_id: str, user: dict[str, Any] = Depends(current_user)) -> dict[str, Any]:
    numeric_id = parse_resource_id(plot_id, "P"); plot_scope(numeric_id, user)
    with engine.connect() as connection:
        row = connection.execute(text("SELECT p.*, s.user_id AS scoped_farmer_user_id, u.real_name AS farmer_name FROM plot p LEFT JOIN farmer_plot_scope s ON s.plot_id=p.plot_id LEFT JOIN users u ON u.user_id=s.user_id WHERE p.plot_id=:id"), {"id": numeric_id}).mappings().one(); return plot_payload(row)


@app.put("/api/plots/{plot_id}/production-offset")
def update_plot_production_offset(plot_id: str, request: PlotProductionOffsetRequest, user: dict[str, Any] = Depends(current_user)) -> dict[str, Any]:
    numeric_id = parse_resource_id(plot_id, "P")
    with engine.begin() as connection:
        farm_id = connection.execute(text("SELECT farm_id FROM plot WHERE plot_id=:plot FOR UPDATE"), {"plot": numeric_id}).scalar()
        if farm_id is None:
            raise HTTPException(status_code=404, detail="PLOT_NOT_FOUND")
        ensure_farm_admin_access(user, farm_id)
        connection.execute(text("UPDATE plot SET production_date_offset_days=:offset WHERE plot_id=:plot"), {"offset": request.offsetDays, "plot": numeric_id})
        row = connection.execute(text("SELECT p.*,s.user_id AS scoped_farmer_user_id,u.real_name AS farmer_name FROM plot p LEFT JOIN farmer_plot_scope s ON s.plot_id=p.plot_id LEFT JOIN users u ON u.user_id=s.user_id WHERE p.plot_id=:plot"), {"plot": numeric_id}).mappings().one()
        record_audit_log(connection, user, "production_plan", "update_plot_date_offset", target_type="plot", target_id=business_id("P", numeric_id), detail=f"Set production schedule offset to {request.offsetDays} days")
        return plot_payload(row)


@app.patch("/api/plots/{plot_id}")
def update_plot(plot_id: str, request: PlotUpdateRequest, user: dict[str, Any] = Depends(current_user)) -> dict[str, Any]:
    ensure_system_admin(user)
    numeric_id = parse_resource_id(plot_id, "P")
    fields = request.model_fields_set
    if not fields:
        raise HTTPException(status_code=400, detail="INVALID_REQUEST")
    if "status" in fields and request.status not in {"active", "inactive"}:
        raise HTTPException(status_code=400, detail="INVALID_REQUEST")
    if "irrigationMethod" in fields and request.irrigationMethod not in {None, "drip", "sprinkler", "flood", "manual", "other"}:
        raise HTTPException(status_code=400, detail="INVALID_REQUEST")

    values = {
        "farm_id": parse_resource_id(request.farmId, "F") if "farmId" in fields else None,
        "plot_code": request.code.strip() if "code" in fields else None,
        "plot_name": request.name.strip() if "name" in fields else None,
        "area_mu": request.areaMu,
        "location": request.location,
        "irrigation_method": request.irrigationMethod,
        "plot_status": request.status,
    }
    assignments = [column + " = :" + column for field, column in (("farmId", "farm_id"), ("code", "plot_code"), ("name", "plot_name"), ("areaMu", "area_mu"), ("location", "location"), ("irrigationMethod", "irrigation_method"), ("status", "plot_status")) if field in fields]
    values["id"] = numeric_id
    with engine.begin() as connection:
        if "farmId" in fields:
            ensure_existing_resources(connection, [values["farm_id"]], [])
        try:
            result = connection.execute(text("UPDATE plot SET " + ", ".join(assignments) + " WHERE plot_id = :id"), values)
        except IntegrityError as error:
            raise HTTPException(status_code=409, detail="PLOT_CODE_EXISTS") from error
        if result.rowcount == 0:
            raise HTTPException(status_code=404, detail="PLOT_NOT_FOUND")
        row = connection.execute(text("SELECT p.*, s.user_id AS scoped_farmer_user_id, u.real_name AS farmer_name FROM plot p LEFT JOIN farmer_plot_scope s ON s.plot_id=p.plot_id LEFT JOIN users u ON u.user_id=s.user_id WHERE p.plot_id=:id"), {"id": numeric_id}).mappings().one()
        record_audit_log(connection, user, "master_data", "update_plot", target_type="plot", target_id=business_id("P", numeric_id), detail="Updated plot profile, ownership, or status")
        return plot_payload(row)


@app.get("/api/plots/{plot_id}/devices")
def list_plot_devices(plot_id: str, user: dict[str, Any] = Depends(current_user)) -> list[dict[str, Any]]:
    numeric_id = parse_resource_id(plot_id, "P")
    plot_scope(numeric_id, user)
    with engine.connect() as connection:
        rows = connection.execute(text("SELECT device_id, device_code, device_uuid, device_type, plot_id, status, switch_state, switch_changed_at, last_heartbeat FROM device WHERE plot_id = :plot_id ORDER BY device_id"), {"plot_id": numeric_id}).fetchall()
    return [device_payload(row) for row in rows]


@app.get("/api/farms/{farm_id}/overview")
def farm_overview(
    farm_id: str,
    user: dict[str, Any] = Depends(current_user),
) -> list[dict[str, Any]]:
    try:
        numeric_farm_id = parse_id(farm_id, "F")
    except ValueError as error:
        raise HTTPException(status_code=400, detail=str(error))

    ensure_farm_access(user, numeric_farm_id)

    with engine.connect() as connection:
        plot_filter = "farm_id = :farm_id"
        params: dict[str, Any] = {"farm_id": numeric_farm_id}
        if user["role"] == "farmer":
            plot_filter += (
                " AND (farmer_user_id = :user_id OR EXISTS ("
                "SELECT 1 FROM farmer_plot_scope AS s "
                "WHERE s.user_id = :user_id AND s.plot_id = plot.plot_id))"
            )
            params["user_id"] = user["user_id"]

        plots = connection.execute(
            text(
                f"""
                SELECT plot_id, farm_id, plot_name, area_mu, location, irrigation_method, crop_name, description
                FROM plot
                WHERE {plot_filter}
                ORDER BY plot_id
                """
            ),
            params,
        ).fetchall()

        result = []
        for plot in plots:
            devices = connection.execute(
                text(
                    """
                           SELECT device_id, device_code, device_uuid, device_type,
                           plot_id, status, switch_state, switch_changed_at, last_heartbeat
                    FROM device
                    WHERE plot_id = :plot_id
                    ORDER BY device_id
                    """
                ),
                {"plot_id": plot.plot_id},
            ).fetchall()
            report = connection.execute(
                text(
                    """
                    SELECT plot_moisture, temperature, report_datetime
                    FROM plot_report
                    WHERE plot_id = :plot_id
                    ORDER BY report_datetime DESC
                    LIMIT 1
                    """
                ),
                {"plot_id": plot.plot_id},
            ).first()

            metrics = []
            if report:
                metrics = [
                    {
                        "type": "humidity",
                        "value": float(report.plot_moisture),
                        "unit": "%",
                        "ts": iso(report.report_datetime),
                    },
                    {
                        "type": "temperature",
                        "value": float(report.temperature),
                        "unit": "C",
                        "ts": iso(report.report_datetime),
                    },
                ]

            result.append(
                {
                    "plot": {
                        "id": business_id("P", plot.plot_id),
                        "farmId": business_id("F", plot.farm_id),
                        "name": plot.plot_name,
                        "areaMu": float(plot.area_mu) if plot.area_mu is not None else None,
                        "location": plot.location,
                        "irrigationMethod": plot.irrigation_method,
                        "crop": plot.crop_name,
                        "description": plot.description,
                    },
                    "latestMetrics": metrics,
                    "devices": [device_payload(device) for device in devices],
                }
            )
    return result


@app.get("/api/farms/{farm_id}/management-statistics")
def farm_management_statistics(
    farm_id: str,
    user: dict[str, Any] = Depends(current_user),
) -> list[dict[str, Any]]:
    """Summarize actionable reminders, active-plan completion, and irrigation by plot."""
    numeric_farm_id = parse_resource_id(farm_id, "F")
    ensure_farm_admin_access(user, numeric_farm_id)

    with engine.connect() as connection:
        rows = connection.execute(
            text(
                """
                SELECT p.plot_id, p.plot_name,
                       COALESCE(task_stats.total_count, 0) AS task_total,
                       COALESCE(task_stats.completed_count, 0) AS completed_count,
                       COALESCE(task_stats.pending_count, 0) AS pending_count,
                       COALESCE(task_stats.overdue_count, 0) AS overdue_count,
                       COALESCE(reminder_stats.pending_count, 0) AS pending_reminder_count,
                       irrigation.record_id, irrigation.action_type, irrigation.activation_type,
                       irrigation.status AS irrigation_status, irrigation.requested_at,
                       irrigation.executed_at
                FROM plot AS p
                LEFT JOIN (
                    SELECT t.plot_id,
                           COUNT(*) AS total_count,
                           SUM(CASE WHEN t.task_status = 'completed' THEN 1 ELSE 0 END) AS completed_count,
                           SUM(CASE WHEN t.task_status IN ('pending', 'in_progress') AND t.scheduled_date >= CURRENT_DATE THEN 1 ELSE 0 END) AS pending_count,
                           SUM(CASE WHEN t.task_status IN ('pending', 'in_progress') AND t.scheduled_date < CURRENT_DATE THEN 1 ELSE 0 END) AS overdue_count
                    FROM management_task AS t
                    JOIN management_plan AS mp
                      ON mp.plan_id = t.plan_id AND mp.plan_status = 'active'
                    GROUP BY t.plot_id
                ) AS task_stats ON task_stats.plot_id = p.plot_id
                LEFT JOIN (
                    SELECT plot_id, COUNT(*) AS pending_count
                    FROM plan_adjustment_reminder
                    WHERE reminder_status = 'active'
                    GROUP BY plot_id
                ) AS reminder_stats ON reminder_stats.plot_id = p.plot_id
                LEFT JOIN irrigation_record AS irrigation ON irrigation.record_id = (
                    SELECT latest.record_id
                    FROM irrigation_record AS latest
                    WHERE latest.plot_id = p.plot_id
                    ORDER BY latest.requested_at DESC, latest.record_id DESC
                    LIMIT 1
                )
                WHERE p.farm_id = :farm_id
                ORDER BY p.plot_id
                """
            ),
            {"farm_id": numeric_farm_id},
        ).mappings().all()

    result: list[dict[str, Any]] = []
    for row in rows:
        task_total = int(row["task_total"])
        completed_count = int(row["completed_count"])
        latest_irrigation = None
        if row["record_id"] is not None:
            latest_irrigation = {
                "recordId": business_id("IR", row["record_id"]),
                "action": "irrigation_on" if row["action_type"] == "on" else "irrigation_off",
                "activationType": row["activation_type"],
                "status": row["irrigation_status"],
                "requestedAt": iso(row["requested_at"]),
                "executedAt": iso(row["executed_at"]),
            }
        result.append(
            {
                "plotId": business_id("P", row["plot_id"]),
                "plotName": row["plot_name"],
                "pendingReminderCount": int(row["pending_reminder_count"]),
                "taskTotal": task_total,
                "completedTaskCount": completed_count,
                "pendingTaskCount": int(row["pending_count"]),
                "overdueTaskCount": int(row["overdue_count"]),
                "taskCompletionRate": round(completed_count * 100 / task_total, 1) if task_total else 0,
                "latestIrrigation": latest_irrigation,
            }
        )
    return result


@app.get("/api/devices")
def list_devices(
    plot_id: str | None = Query(default=None, alias="plotId"),
    user: dict[str, Any] = Depends(current_user),
) -> list[dict[str, Any]]:
    params: dict[str, Any] = {}
    clauses = []

    if plot_id is not None:
        try:
            params["plot_id"] = parse_id(plot_id, "P")
        except ValueError as error:
            raise HTTPException(status_code=400, detail=str(error))
        clauses.append("d.plot_id = :plot_id")

    if user["role"] == "farm_admin":
        clauses.append(
            "(f.admin_user_id = :user_id OR EXISTS ("
            "SELECT 1 FROM farm_admin_scope AS s "
            "WHERE s.user_id = :user_id AND s.farm_id = f.farm_id) OR d.plot_id IS NULL)"
        )
        params["user_id"] = user["user_id"]
    elif user["role"] == "farmer":
        clauses.append(
            "(p.farmer_user_id = :user_id OR EXISTS ("
            "SELECT 1 FROM farmer_plot_scope AS s "
            "WHERE s.user_id = :user_id AND s.plot_id = p.plot_id))"
        )
        params["user_id"] = user["user_id"]

    where = " WHERE " + " AND ".join(clauses) if clauses else ""
    with engine.connect() as connection:
        rows = connection.execute(
            text(
                f"""
                SELECT d.device_id, d.device_code, d.device_uuid, d.device_type,
                       d.plot_id, d.status, d.switch_state, d.switch_changed_at, d.last_heartbeat
                FROM device AS d
                LEFT JOIN plot AS p ON p.plot_id = d.plot_id
                LEFT JOIN farm AS f ON f.farm_id = p.farm_id
                {where}
                ORDER BY d.device_id
                """
            ),
            params,
        ).fetchall()
    return [device_payload(row) for row in rows]


@app.post("/api/devices", status_code=status.HTTP_201_CREATED)
def create_device(
    request: DeviceCreateRequest,
    user: dict[str, Any] = Depends(current_user),
) -> dict[str, Any]:
    device_no = request.deviceNo.strip()
    if not device_no:
        raise HTTPException(status_code=400, detail="deviceNo must not be blank")
    if request.type not in {"temperature", "humidity", "irrigation"}:
        raise HTTPException(status_code=400, detail="unsupported device type")
    plot_id = parse_resource_id(request.plotId, "P") if request.plotId else None
    if plot_id is None and user["role"] != "system_admin":
        raise HTTPException(status_code=403, detail="FORBIDDEN")
    if plot_id is not None:
        ensure_plot_manage_access(plot_id, user)

    with engine.begin() as connection:
        try:
            result = connection.execute(
                text(
                    """
                    INSERT INTO device (device_code, device_uuid, device_type, plot_id)
                    VALUES (:device_code, :device_uuid, :device_type, :plot_id)
                    """
                ),
                {
                    "device_code": device_no,
                    "device_uuid": normalized_uuid(request.deviceUuid),
                    "device_type": request.type,
                    "plot_id": plot_id,
                },
            )
        except IntegrityError as error:
            raise HTTPException(status_code=409, detail="DEVICE_ALREADY_EXISTS") from error
        device_id = result.lastrowid
        row = connection.execute(
            text(
                """
                SELECT device_id, device_code, device_uuid, device_type,
                       plot_id, status, switch_state, switch_changed_at, last_heartbeat
                FROM device WHERE device_id = :device_id
                """
            ),
            {"device_id": device_id},
        ).mappings().one()
    return device_payload(row)


@app.post("/api/devices/{device_id}/bind")
def bind_device(
    device_id: str,
    request: DeviceBindRequest,
    user: dict[str, Any] = Depends(current_user),
) -> dict[str, Any]:
    numeric_device_id = parse_resource_id(device_id, "D")
    target_plot_id = parse_resource_id(request.plotId, "P")
    ensure_device_manage_access(numeric_device_id, user, allow_unbound_admin=True)
    ensure_plot_manage_access(target_plot_id, user)
    with engine.begin() as connection:
        try:
            connection.execute(
                text("UPDATE device SET plot_id = :plot_id WHERE device_id = :device_id"),
                {"plot_id": target_plot_id, "device_id": numeric_device_id},
            )
        except IntegrityError as error:
            raise HTTPException(status_code=409, detail="PLOT_DEVICE_TYPE_EXISTS") from error
        row = connection.execute(
            text(
                """
                SELECT device_id, device_code, device_uuid, device_type,
                       plot_id, status, switch_state, switch_changed_at, last_heartbeat
                FROM device WHERE device_id = :device_id
                """
            ),
            {"device_id": numeric_device_id},
        ).mappings().one()
        record_audit_log(connection, user, "device_binding", "bind_device", target_type="device", target_id=business_id("D", numeric_device_id), detail=f"Bound device to {business_id('P', target_plot_id)}")
    return device_payload(row)


@app.delete("/api/devices/{device_id}/bind", status_code=status.HTTP_204_NO_CONTENT)
def unbind_device(
    device_id: str,
    user: dict[str, Any] = Depends(current_user),
) -> None:
    numeric_device_id = parse_resource_id(device_id, "D")
    ensure_device_manage_access(numeric_device_id, user)
    with engine.begin() as connection:
        connection.execute(
            text("UPDATE device SET plot_id = NULL, status = 'offline' WHERE device_id = :device_id"),
            {"device_id": numeric_device_id},
        )
        record_audit_log(connection, user, "device_binding", "unbind_device", target_type="device", target_id=business_id("D", numeric_device_id), detail="Unbound device and marked it offline")


@app.get("/api/devices/{device_id}/status")
def get_device_status(
    device_id: str,
    user: dict[str, Any] = Depends(current_user),
) -> dict[str, Any]:
    device = device_scope(parse_resource_id(device_id, "D"), user)
    return {
        "deviceId": business_id("D", device["device_id"]),
        "status": device["status"],
        "switchState": device["switch_state"],
        "switchChangedAt": iso(device["switch_changed_at"]),
        "lastSeenAt": iso(device["last_heartbeat"]),
    }


@app.post("/api/devices/{device_id}/commands", status_code=202)
def create_device_command(
    device_id: str,
    request: CommandCreateRequest,
    user: dict[str, Any] = Depends(current_user),
) -> dict[str, Any]:
    try:
        numeric_device_id = parse_id(device_id, "D")
    except ValueError as error:
        raise HTTPException(status_code=400, detail=str(error))

    with engine.connect() as connection:
        device = connection.execute(
            text(
                """
                SELECT d.device_id, d.device_uuid, d.device_type,
                       p.plot_id, p.farmer_user_id, f.farm_id, f.admin_user_id
                FROM device AS d
                JOIN plot AS p ON p.plot_id = d.plot_id
                JOIN farm AS f ON f.farm_id = p.farm_id
                WHERE d.device_id = :device_id
                """
            ),
            {"device_id": numeric_device_id},
        ).mappings().first()
    if device is None:
        raise HTTPException(status_code=404, detail="DEVICE_NOT_FOUND")
    plot_scope(device["plot_id"], user)
    if device["device_type"] != "irrigation":
        raise HTTPException(status_code=409, detail="DEVICE_NOT_IRRIGATION")

    action = request.action
    action = {"irrigation_on": "on", "irrigation_off": "off"}.get(action, action)
    if action not in {"on", "off"}:
        raise HTTPException(status_code=400, detail="action must be irrigation_on, irrigation_off, on, or off")
    if action == "on" and request.durationMinutes is None:
        raise HTTPException(status_code=400, detail="durationMinutes is required when starting irrigation")
    if action == "off" and request.durationMinutes is not None:
        raise HTTPException(status_code=400, detail="durationMinutes is only allowed when starting irrigation")
    try:
        command = mqtt_service.publish_control_command(
            device_uuid=device["device_uuid"],
            switch=action,
            operator_id=user["user_id"],
            activation_type=request.activationType,
            planned_duration_minutes=request.durationMinutes,
        )
        with engine.begin() as connection:
            record_audit_log(connection, user, "irrigation", "issue_irrigation_command", target_type="device", target_id=business_id("D", numeric_device_id), detail=f"Issued irrigation_{action} command {command['commandId']} for {business_id('P', device['plot_id'])}")
        return {
            "commandId": command["commandId"],
            "deviceId": business_id("D", numeric_device_id),
            "deviceUuid": device["device_uuid"],
            "action": f"irrigation_{action}",
            "activationType": request.activationType,
            "status": command["status"],
            "requestedAt": command["requestedAt"],
            "executedAt": None,
            "durationMinutes": request.durationMinutes,
            "scheduledStopAt": None,
        }
    except ValueError as error:
        raise HTTPException(status_code=400, detail=str(error))
    except RuntimeError as error:
        raise HTTPException(status_code=503, detail=str(error))


@app.get("/api/commands/{command_id}")
def get_command(
    command_id: str,
    user: dict[str, Any] = Depends(current_user),
) -> dict[str, Any]:
    with engine.connect() as connection:
        command = connection.execute(
            text(
                """
                SELECT r.command_id, r.device_id, d.device_uuid, r.plot_id,
                       r.action_type, r.activation_type, r.status,
                       r.requested_at, r.executed_at, r.planned_duration_minutes,
                       r.scheduled_stop_at
                FROM irrigation_record AS r
                JOIN device AS d ON d.device_id = r.device_id
                WHERE r.command_id = :command_id
                """
            ),
            {"command_id": command_id},
        ).mappings().first()
    if command is None:
        raise HTTPException(status_code=404, detail="COMMAND_NOT_FOUND")
    plot_scope(command["plot_id"], user)
    return command_payload(command)


@app.get("/api/plots/{plot_id}/metrics")
def list_plot_metrics(
    plot_id: str,
    metric_type: str = Query(alias="type"),
    from_value: datetime = Query(alias="from"),
    to_value: datetime = Query(alias="to"),
    user: dict[str, Any] = Depends(current_user),
) -> list[dict[str, Any]]:
    numeric_plot_id = parse_resource_id(plot_id, "P")
    plot_scope(numeric_plot_id, user)
    if metric_type not in {"temperature", "humidity"}:
        raise HTTPException(status_code=400, detail="type must be temperature or humidity")
    try:
        from_utc = utc_naive(from_value)
        to_utc = utc_naive(to_value)
    except ValueError as error:
        raise HTTPException(status_code=400, detail=str(error)) from error
    if from_utc >= to_utc:
        raise HTTPException(status_code=400, detail="from must be before to")

    value_column = "temperature" if metric_type == "temperature" else "plot_moisture"
    unit = "C" if metric_type == "temperature" else "%"
    with engine.connect() as connection:
        rows = connection.execute(
            text(
                f"""
                SELECT {value_column} AS metric_value, report_datetime
                FROM plot_report
                WHERE plot_id = :plot_id
                  AND report_datetime >= :from_utc
                  AND report_datetime < :to_utc
                ORDER BY report_datetime
                """
            ),
            {"plot_id": numeric_plot_id, "from_utc": from_utc, "to_utc": to_utc},
        ).mappings().all()
    return [
        {
            "type": metric_type,
            "value": float(row["metric_value"]),
            "unit": unit,
            "ts": iso(row["report_datetime"]),
        }
        for row in rows
    ]


@app.get("/api/plots/{plot_id}/management-history")
def plot_management_history(
    plot_id: str,
    from_value: datetime = Query(alias="from"),
    to_value: datetime = Query(alias="to"),
    user: dict[str, Any] = Depends(current_user),
) -> dict[str, Any]:
    """Return the plot's existing management records in one chronological feed."""
    numeric_plot_id = parse_resource_id(plot_id, "P")
    plot_scope(numeric_plot_id, user)
    try:
        from_utc = utc_naive(from_value)
        to_utc = utc_naive(to_value)
    except ValueError as error:
        raise HTTPException(status_code=400, detail=str(error)) from error
    if from_utc >= to_utc:
        raise HTTPException(status_code=400, detail="from must be before to")

    events: list[dict[str, Any]] = []
    with engine.connect() as connection:
        tasks = connection.execute(
            text(
                "SELECT t.task_id, t.title, t.task_type, t.scheduled_date, t.task_status, t.note, "
                "p.plan_status FROM management_task t "
                "LEFT JOIN management_plan p ON p.plan_id = t.plan_id "
                "WHERE t.plot_id = :plot AND t.scheduled_date >= :start AND t.scheduled_date <= :end "
                "ORDER BY t.scheduled_date, t.task_id"
            ),
            {"plot": numeric_plot_id, "start": from_utc.date(), "end": to_utc.date()},
        ).mappings().all()
        completions = connection.execute(
            text(
                "SELECT c.task_completion_id, c.executed_at, c.executor_name, c.result, c.note, "
                "t.task_id, t.title FROM management_task_completion c "
                "JOIN management_task t ON t.task_id = c.task_id "
                "WHERE t.plot_id = :plot AND c.executed_at >= :start AND c.executed_at < :end "
                "ORDER BY c.executed_at, c.task_completion_id"
            ),
            {"plot": numeric_plot_id, "start": from_utc, "end": to_utc},
        ).mappings().all()
        irrigation = connection.execute(
            text(
                "SELECT r.record_id, r.device_id, r.action_type, r.activation_type, r.status, r.requested_at, r.executed_at, "
                "u.real_name AS operator_name FROM irrigation_record r "
                "LEFT JOIN users u ON u.user_id = r.operator_id "
                "WHERE r.plot_id = :plot AND r.requested_at >= :start AND r.requested_at < :end "
                "ORDER BY r.requested_at, r.record_id"
            ),
            {"plot": numeric_plot_id, "start": from_utc, "end": to_utc},
        ).mappings().all()
        alerts = connection.execute(
            text(
                "SELECT a.alert_id, a.alert_type, a.alert_description, a.alert_status, a.alert_datetime, "
                "d.device_code FROM alert a LEFT JOIN device d ON d.device_id = a.device_id "
                "WHERE a.plot_id = :plot AND a.alert_datetime >= :start AND a.alert_datetime < :end "
                "ORDER BY a.alert_datetime, a.alert_id"
            ),
            {"plot": numeric_plot_id, "start": from_utc, "end": to_utc},
        ).mappings().all()

    for row in tasks:
        scheduled_at = datetime.combine(row["scheduled_date"], time.min)
        events.append({
            "id": f"task:{business_id('MT', row['task_id'])}",
            "type": "plan_task",
            "occurredAt": iso(scheduled_at),
            "title": row["title"],
            "description": row["note"] or f"计划{row['scheduled_date'].isoformat()}执行",
            "operator": "系统",
            "status": task_status_value(row),
            "referenceId": business_id("MT", row["task_id"]),
        })
    for row in completions:
        events.append({
            "id": f"completion:{business_id('TC', row['task_completion_id'])}",
            "type": "task_completion",
            "occurredAt": iso(row["executed_at"]),
            "title": f"完成：{row['title']}",
            "description": f"完成{row['title']}",
            "operator": row["executor_name"],
            "status": "completed",
            "referenceId": business_id("TC", row["task_completion_id"]),
        })
    for index, row in enumerate(irrigation):
        if row["action_type"] != "on":
            continue
        stopped_at = None
        for following in irrigation[index + 1:]:
            if following["device_id"] != row["device_id"]:
                continue
            if following["action_type"] == "on":
                break
            stopped_at = following["requested_at"]
            break

        action = "开启灌溉"
        mode = "自动" if row["activation_type"] == "auto" else "手动"
        duration_seconds = max(0, round((stopped_at - row["requested_at"]).total_seconds())) if stopped_at else None
        if row["status"] == "failed":
            description = f"{mode}灌溉失败"
        elif row["status"] == "timeout":
            description = f"{mode}灌溉超时"
        elif duration_seconds is None:
            description = f"{mode}灌溉"
        elif duration_seconds >= 60:
            description = f"{mode}灌溉 {round(duration_seconds / 60)} 分钟"
        else:
            description = f"{mode}灌溉 {duration_seconds} 秒"
        events.append({
            "id": f"irrigation:{business_id('IR', row['record_id'])}",
            "type": "irrigation",
            "occurredAt": iso(row["requested_at"]),
            "title": action,
            "description": description,
            "activationType": row["activation_type"],
            "durationSeconds": duration_seconds,
            "operator": row["operator_name"] or ("系统" if row["activation_type"] == "auto" else "农户"),
            "status": row["status"],
            "referenceId": business_id("IR", row["record_id"]),
        })
    for row in alerts:
        alert_labels = {
            "low_moisture": "土壤湿度偏低",
            "high_moisture": "土壤湿度偏高",
            "device_offline": "设备离线",
            "device_fault": "设备故障",
            "irrigation_failed": "灌溉失败",
        }
        alert_label = alert_labels.get(row["alert_type"], row["alert_type"])
        description = f"{alert_label}（{row['device_code']}）" if row["device_code"] else alert_label
        events.append({
            "id": f"alert:{business_id('A', row['alert_id'])}",
            "type": "alert",
            "occurredAt": iso(row["alert_datetime"]),
            "title": row["alert_type"],
            "description": description,
            "deviceCode": row["device_code"],
            "operator": "系统",
            "status": row["alert_status"],
            "referenceId": business_id("A", row["alert_id"]),
        })
    events.sort(key=lambda item: (item["occurredAt"] or "", item["id"]), reverse=True)
    return {
        "plotId": business_id("P", numeric_plot_id),
        "from": iso(from_utc),
        "to": iso(to_utc),
        "events": events,
    }


@app.get("/api/irrigation/records")
def list_irrigation_records(
    plot_id: str = Query(alias="plotId"),
    user: dict[str, Any] = Depends(current_user),
) -> list[dict[str, Any]]:
    numeric_plot_id = parse_resource_id(plot_id, "P")
    plot_scope(numeric_plot_id, user)
    with engine.connect() as connection:
        rows = connection.execute(
            text(
                """
                SELECT irrigation.record_id, irrigation.command_id, irrigation.plot_id,
                       irrigation.device_id, irrigation.action_type,
                       irrigation.activation_type, irrigation.status,
                       irrigation.requested_at, irrigation.executed_at,
                       irrigation.planned_duration_minutes, irrigation.scheduled_stop_at,
                       irrigation.operator_id, users.real_name AS operator_name
                FROM irrigation_record AS irrigation
                LEFT JOIN users ON users.user_id = irrigation.operator_id
                WHERE irrigation.plot_id = :plot_id
                ORDER BY irrigation.requested_at DESC, irrigation.record_id DESC
                """
            ),
            {"plot_id": numeric_plot_id},
        ).mappings().all()
    return [irrigation_record_payload(row) for row in rows]


@app.get("/api/plots/{plot_id}/irrigation-settings")
def get_irrigation_settings(
    plot_id: str,
    user: dict[str, Any] = Depends(current_user),
) -> dict[str, Any]:
    numeric_plot_id = parse_resource_id(plot_id, "P")
    plot_scope(numeric_plot_id, user)
    with engine.connect() as connection:
        row = connection.execute(
            text(
                """
                SELECT soil_moisture_threshold, irrigation_length,
                       irrigation_interval, updated_at
                FROM threshold_config
                WHERE plot_id = :plot_id
                """
            ),
            {"plot_id": numeric_plot_id},
        ).mappings().first()
    if row is None:
        return {
            "plotId": business_id("P", numeric_plot_id),
            "soilMoistureThreshold": 35,
            "irrigationDuration": 30,
            "irrigationInterval": 60,
            "updatedAt": None,
        }
    return {
        "plotId": business_id("P", numeric_plot_id),
        "soilMoistureThreshold": float(value_of(row, "soil_moisture_threshold")),
        "irrigationDuration": value_of(row, "irrigation_length"),
        "irrigationInterval": value_of(row, "irrigation_interval"),
        "updatedAt": iso(value_of(row, "updated_at")),
    }


@app.put("/api/plots/{plot_id}/irrigation-settings")
def update_irrigation_settings(
    plot_id: str,
    request: IrrigationSettingsRequest,
    user: dict[str, Any] = Depends(current_user),
) -> dict[str, Any]:
    numeric_plot_id = parse_resource_id(plot_id, "P")
    plot_scope(numeric_plot_id, user)
    with engine.begin() as connection:
        connection.execute(
            text(
                """
                INSERT INTO threshold_config (
                    plot_id, soil_moisture_threshold, irrigation_length,
                    irrigation_interval, updated_by
                ) VALUES (
                    :plot_id, :soil_moisture_threshold, :irrigation_length,
                    :irrigation_interval, :updated_by
                )
                ON DUPLICATE KEY UPDATE
                    soil_moisture_threshold = VALUES(soil_moisture_threshold),
                    irrigation_length = VALUES(irrigation_length),
                    irrigation_interval = VALUES(irrigation_interval),
                    updated_by = VALUES(updated_by)
                """
            ),
            {
                "plot_id": numeric_plot_id,
                "soil_moisture_threshold": request.soilMoistureThreshold,
                "irrigation_length": request.irrigationDuration,
                "irrigation_interval": request.irrigationInterval,
                "updated_by": user["user_id"],
            },
        )
        evaluate_latest_moisture(
            connection,
            plot_id=numeric_plot_id,
            threshold=request.soilMoistureThreshold,
        )
        row = connection.execute(
            text(
                """
                SELECT soil_moisture_threshold, irrigation_length, irrigation_interval, updated_at
                FROM threshold_config
                WHERE plot_id = :plot_id
                """
            ),
            {"plot_id": numeric_plot_id},
        ).mappings().one()
    return {
        "plotId": business_id("P", numeric_plot_id),
        "soilMoistureThreshold": float(value_of(row, "soil_moisture_threshold")),
        "irrigationDuration": value_of(row, "irrigation_length"),
        "irrigationInterval": value_of(row, "irrigation_interval"),
        "updatedAt": iso(value_of(row, "updated_at")),
    }


def load_weather_forecast(numeric_plot_id: int) -> dict[str, Any]:
    """Load the cached/provider forecast shared by weather and recommendation APIs."""
    service_url = os.getenv("WEATHER_SERVICE_URL", "").rstrip("/")
    if service_url:
        separator = "&" if "?" in service_url else "?"
        target = f"{service_url}{separator}{urlencode({'plotId': business_id('P', numeric_plot_id)})}"
        try:
            with urlopen(target, timeout=float(os.getenv("WEATHER_SERVICE_TIMEOUT_SECONDS", "10"))) as response:
                result = json.loads(response.read().decode("utf-8"))
        except (HTTPError, URLError, TimeoutError, json.JSONDecodeError, OSError) as error:
            raise HTTPException(status_code=503, detail="WEATHER_SERVICE_UNAVAILABLE") from error
    else:
        configured_location = os.getenv("WEATHER_LOCATION", "重庆市奉节县").strip()
        if configured_location:
            location = {"province": "", "city": configured_location, "address": ""}
        else:
            with engine.connect() as connection:
                location = connection.execute(
                    text(
                        "SELECT f.province, f.city, f.address FROM plot p "
                        "JOIN farm f ON f.farm_id = p.farm_id WHERE p.plot_id = :plot_id"
                    ),
                    {"plot_id": numeric_plot_id},
                ).mappings().one()
        try:
            result = open_meteo_forecast(
                province=location["province"] or "",
                city=location["city"] or "",
                address=location["address"] or "",
            )
        except (WeatherProviderError, ValueError) as error:
            raise HTTPException(status_code=503, detail="WEATHER_SERVICE_UNAVAILABLE") from error
    return result


def refresh_weather_reminders() -> None:
    """Evaluate weather reminders for every active plot and discard older rows."""
    timestamp = datetime.now(UTC).replace(tzinfo=None)
    with engine.begin() as connection:
        purge_non_today_reminders(connection, timestamp=timestamp)
        plot_ids = connection.execute(
            text("SELECT plot_id FROM plot WHERE plot_status = 'active' ORDER BY plot_id")
        ).scalars().all()
    for plot_id in plot_ids:
        try:
            forecast = load_weather_forecast(int(plot_id))
            with engine.begin() as connection:
                sync_weather_reminders(
                    connection,
                    plot_id=int(plot_id),
                    forecast=forecast,
                    timestamp=timestamp,
                )
        except Exception:
            logger.exception("Failed to refresh weather reminders for plot %s", plot_id)


def weather_reminder_loop() -> None:
    interval = max(60, int(os.getenv("WEATHER_REMINDER_REFRESH_SECONDS", "1800")))
    while not weather_reminder_stop.is_set():
        try:
            refresh_weather_reminders()
        except Exception:
            logger.exception("Failed to run weather reminder refresh")
        if weather_reminder_stop.wait(interval):
            break


@app.get("/api/plots/{plot_id}/irrigation-recommendation")
def get_irrigation_recommendation(
    plot_id: str,
    user: dict[str, Any] = Depends(current_user),
) -> dict[str, Any]:
    """Calculate a stage-, sensor-, weather-, history-, and device-aware recommendation."""
    numeric_plot_id = parse_resource_id(plot_id, "P")
    plot_scope(numeric_plot_id, user)
    now = datetime.now(UTC).replace(tzinfo=None)
    with engine.connect() as connection:
        settings = connection.execute(
            text(
                "SELECT soil_moisture_threshold, irrigation_length, irrigation_interval "
                "FROM threshold_config WHERE plot_id = :plot_id"
            ),
            {"plot_id": numeric_plot_id},
        ).mappings().first()
        fallback_threshold = float(settings["soil_moisture_threshold"]) if settings else 35.0
        rule = current_moisture_rule(
            connection,
            plot_id=numeric_plot_id,
            fallback_threshold=fallback_threshold,
            as_of=now.date(),
        )
        report = connection.execute(
            text(
                "SELECT plot_moisture, report_datetime FROM plot_report "
                "WHERE plot_id = :plot_id ORDER BY report_datetime DESC, report_id DESC LIMIT 1"
            ),
            {"plot_id": numeric_plot_id},
        ).mappings().first()
        device = connection.execute(
            text(
                "SELECT device_id, status FROM device WHERE plot_id = :plot_id "
                "AND device_type = 'irrigation' ORDER BY device_id DESC LIMIT 1"
            ),
            {"plot_id": numeric_plot_id},
        ).mappings().first()
        latest_command = connection.execute(
            text(
                "SELECT action_type, status, requested_at, scheduled_stop_at "
                "FROM irrigation_record WHERE plot_id = :plot_id AND status = 'success' "
                "ORDER BY requested_at DESC, record_id DESC LIMIT 1"
            ),
            {"plot_id": numeric_plot_id},
        ).mappings().first()
        last_irrigation = connection.execute(
            text(
                "SELECT requested_at, planned_duration_minutes FROM irrigation_record "
                "WHERE plot_id = :plot_id AND action_type = 'on' AND status = 'success' "
                "ORDER BY requested_at DESC, record_id DESC LIMIT 1"
            ),
            {"plot_id": numeric_plot_id},
        ).mappings().first()

    weather_available = True
    weather: dict[str, Any] = {}
    try:
        loaded_weather = load_weather_forecast(numeric_plot_id)
        weather = loaded_weather if isinstance(loaded_weather, dict) else {}
    except HTTPException:
        weather_available = False
    current_weather = weather.get("current", weather)
    if not isinstance(current_weather, dict):
        current_weather = {}
    hourly = weather.get("hourly", [])
    daily = weather.get("daily", [])
    probabilities: list[float] = []
    for item in hourly[:24] if isinstance(hourly, list) else []:
        if not isinstance(item, dict):
            continue
        value = item.get("rainProbability", item.get("rain_probability", item.get("precipitation_probability")))
        try:
            probabilities.append(float(value))
        except (TypeError, ValueError):
            continue
    if not probabilities:
        for item in daily[:1] if isinstance(daily, list) else []:
            if not isinstance(item, dict):
                continue
            value = item.get("rainProbability", item.get("rain_probability", item.get("precipitation_probability")))
            try:
                probabilities.append(float(value))
            except (TypeError, ValueError):
                continue
    rain_probability = max(probabilities) if probabilities else None
    scheduled_stop_at = latest_command["scheduled_stop_at"] if latest_command else None
    irrigation_active = bool(
        latest_command
        and latest_command["action_type"] == "on"
        and (scheduled_stop_at is None or scheduled_stop_at > now)
    )
    recommendation = build_irrigation_recommendation(
        now=now,
        moisture=float(report["plot_moisture"]) if report else None,
        report_at=report["report_datetime"] if report else None,
        target_min=rule["lower"],
        target_max=rule["upper"],
        stage_name=rule["stageName"],
        device_bound=device is not None,
        device_online=bool(device and device["status"] == "online"),
        irrigation_active=irrigation_active,
        base_duration_minutes=settings["irrigation_length"] if settings else 30,
        minimum_interval_minutes=settings["irrigation_interval"] if settings else 60,
        last_irrigation_at=last_irrigation["requested_at"] if last_irrigation else None,
        last_irrigation_minutes=last_irrigation["planned_duration_minutes"] if last_irrigation else None,
        rain_probability=rain_probability,
        rainy_days=consecutive_rainy_days(daily) if isinstance(daily, list) else 0,
        temperature=current_weather.get("temperature", current_weather.get("temp")),
        weather_available=weather_available,
    )
    return {"plotId": business_id("P", numeric_plot_id), **recommendation}


@app.get("/api/weather/forecast")
def weather_forecast(
    plot_id: str = Query(alias="plotId"),
    user: dict[str, Any] = Depends(current_user),
) -> dict[str, Any]:
    """Proxy weather data for an authorized plot without exposing provider credentials."""
    numeric_plot_id = parse_resource_id(plot_id, "P")
    plot_scope(numeric_plot_id, user)
    result = load_weather_forecast(numeric_plot_id)
    if not isinstance(result, dict):
        raise HTTPException(status_code=502, detail="WEATHER_SERVICE_INVALID_RESPONSE")
    current = result.get("current", result)
    if not isinstance(current, dict):
        raise HTTPException(status_code=502, detail="WEATHER_SERVICE_INVALID_RESPONSE")
    daily = result.get("daily", [])
    with engine.begin() as connection:
        sync_weather_reminders(
            connection,
            plot_id=numeric_plot_id,
            forecast=result,
            timestamp=datetime.now(UTC).replace(tzinfo=None),
        )
    return {
        "current": current,
        "hourly": result.get("hourly", []),
        "daily": daily,
    }


@app.get("/api/plan-adjustment-reminders")
def list_plan_adjustment_reminders(
    plot_id: str | None = Query(default=None, alias="plotId"),
    user: dict[str, Any] = Depends(current_user),
) -> list[dict[str, Any]]:
    timestamp = datetime.now(UTC).replace(tzinfo=None)
    start, end = shanghai_day_utc_bounds(timestamp)
    reminder_types = [
        "continuous_rain",
        "high_rain_probability",
        "high_temperature",
        "frost_risk",
        "strong_wind",
    ]
    clauses = [
        "r.triggered_at >= :start",
        "r.triggered_at < :end",
        "r.reminder_type IN :reminder_types",
    ]
    params: dict[str, Any] = {"start": start, "end": end, "reminder_types": reminder_types}
    if plot_id is not None:
        numeric_plot_id = parse_resource_id(plot_id, "P")
        plot_scope(numeric_plot_id, user)
        clauses.append("r.plot_id = :plot_id")
        params["plot_id"] = numeric_plot_id
    elif user["role"] == "farm_admin":
        clauses.append(
            "(f.admin_user_id = :user_id OR EXISTS ("
            "SELECT 1 FROM farm_admin_scope s WHERE s.user_id = :user_id AND s.farm_id = f.farm_id))"
        )
        params["user_id"] = user["user_id"]
    elif user["role"] == "farmer":
        clauses.append(
            "(p.farmer_user_id = :user_id OR EXISTS ("
            "SELECT 1 FROM farmer_plot_scope s WHERE s.user_id = :user_id AND s.plot_id = p.plot_id))"
        )
        params["user_id"] = user["user_id"]
    where = " WHERE " + " AND ".join(clauses)
    with engine.begin() as connection:
        purge_non_today_reminders(connection, timestamp=timestamp)
        reminders = connection.execute(
            text(
                "SELECT r.reminder_id, r.plot_id, p.plot_name, r.reminder_type, "
                "r.reminder_description, r.triggered_at "
                "FROM plan_adjustment_reminder r "
                "JOIN plot p ON p.plot_id = r.plot_id JOIN farm f ON f.farm_id = p.farm_id "
                f"{where} ORDER BY r.triggered_at DESC, r.reminder_id DESC"
            ).bindparams(bindparam("reminder_types", expanding=True)),
            params,
        ).mappings().all()
        result = []
        for reminder in reminders:
            tasks = connection.execute(
                text(
                    "SELECT task_id, task_title, task_type, scheduled_date "
                    "FROM plan_adjustment_reminder_task WHERE reminder_id = :reminder "
                    "ORDER BY scheduled_date, task_id"
                ),
                {"reminder": reminder["reminder_id"]},
            ).mappings().all()
            result.append(plan_adjustment_reminder_payload(reminder, tasks))
    return result


@app.get("/api/alerts")
def list_alerts(
    plot_id: str | None = Query(default=None, alias="plotId"),
    alert_status: str | None = Query(default=None, alias="status"),
    user: dict[str, Any] = Depends(current_user),
) -> list[dict[str, Any]]:
    clauses: list[str] = []
    params: dict[str, Any] = {}
    if plot_id is not None:
        numeric_plot_id = parse_resource_id(plot_id, "P")
        plot_scope(numeric_plot_id, user)
        clauses.append("a.plot_id = :plot_id")
        params["plot_id"] = numeric_plot_id
    elif user["role"] == "farm_admin":
        clauses.append(
            "(f.admin_user_id = :user_id OR EXISTS ("
            "SELECT 1 FROM farm_admin_scope AS s "
            "WHERE s.user_id = :user_id AND s.farm_id = f.farm_id))"
        )
        params["user_id"] = user["user_id"]
    elif user["role"] == "farmer":
        clauses.append(
            "(p.farmer_user_id = :user_id OR EXISTS ("
            "SELECT 1 FROM farmer_plot_scope AS s "
            "WHERE s.user_id = :user_id AND s.plot_id = p.plot_id))"
        )
        params["user_id"] = user["user_id"]
    if alert_status is not None:
        if alert_status not in {"active", "acknowledged", "recovered"}:
            raise HTTPException(status_code=400, detail="invalid alert status")
        clauses.append("a.alert_status = :alert_status")
        params["alert_status"] = alert_status
    where = " WHERE " + " AND ".join(clauses) if clauses else ""
    with engine.connect() as connection:
        rows = connection.execute(
            text(
                f"""
                SELECT a.alert_id, a.plot_id, a.device_id, a.alert_type,
                       a.alert_description, a.alert_value, a.alert_threshold,
                       a.alert_status, a.alert_datetime, a.handled_at
                FROM alert AS a
                JOIN plot AS p ON p.plot_id = a.plot_id
                JOIN farm AS f ON f.farm_id = p.farm_id
                {where}
                ORDER BY a.alert_datetime DESC, a.alert_id DESC
                """
            ),
            params,
        ).mappings().all()
    return [alert_payload(row) for row in rows]


@app.post("/api/alerts/{alert_id}/ack")
def acknowledge_alert(
    alert_id: str,
    user: dict[str, Any] = Depends(current_user),
) -> dict[str, Any]:
    numeric_alert_id = parse_resource_id(alert_id, "A")
    with engine.begin() as connection:
        alert = connection.execute(
            text("SELECT plot_id FROM alert WHERE alert_id = :alert_id"),
            {"alert_id": numeric_alert_id},
        ).mappings().first()
        if alert is None:
            raise HTTPException(status_code=404, detail="ALERT_NOT_FOUND")
        plot_scope(alert["plot_id"], user)
        connection.execute(
            text(
                """
                UPDATE alert
                SET alert_status = 'acknowledged', handled_by = :user_id,
                    handled_at = :handled_at
                WHERE alert_id = :alert_id AND alert_status = 'active'
                """
            ),
            {
                "alert_id": numeric_alert_id,
                "user_id": user["user_id"],
                "handled_at": datetime.now(UTC).replace(tzinfo=None),
            },
        )
        row = connection.execute(
            text(
                """
                SELECT alert_id, plot_id, device_id, alert_type, alert_description,
                       alert_value, alert_threshold, alert_status, alert_datetime,
                       handled_at
                FROM alert WHERE alert_id = :alert_id
                """
            ),
            {"alert_id": numeric_alert_id},
        ).mappings().one()
    return alert_payload(row)


@app.get("/api/alert-rules")
def list_alert_rules(
    plot_id: str | None = Query(default=None, alias="plotId"),
    user: dict[str, Any] = Depends(current_user),
) -> list[dict[str, Any]]:
    clauses: list[str] = []
    params: dict[str, Any] = {}
    if plot_id is not None:
        numeric_plot_id = parse_resource_id(plot_id, "P")
        plot_scope(numeric_plot_id, user)
        clauses.append("t.plot_id = :plot_id")
        params["plot_id"] = numeric_plot_id
    elif user["role"] == "farm_admin":
        clauses.append(
            "(f.admin_user_id = :user_id OR EXISTS ("
            "SELECT 1 FROM farm_admin_scope AS s "
            "WHERE s.user_id = :user_id AND s.farm_id = f.farm_id))"
        )
        params["user_id"] = user["user_id"]
    elif user["role"] == "farmer":
        clauses.append(
            "(p.farmer_user_id = :user_id OR EXISTS ("
            "SELECT 1 FROM farmer_plot_scope AS s "
            "WHERE s.user_id = :user_id AND s.plot_id = p.plot_id))"
        )
        params["user_id"] = user["user_id"]
    where = " WHERE " + " AND ".join(clauses) if clauses else ""
    with engine.connect() as connection:
        rows = connection.execute(
            text(
                f"""
                SELECT t.config_id, t.plot_id, t.soil_moisture_threshold,
                       t.irrigation_length, t.irrigation_interval, t.updated_at
                FROM threshold_config AS t
                JOIN plot AS p ON p.plot_id = t.plot_id
                JOIN farm AS f ON f.farm_id = p.farm_id
                {where}
                ORDER BY t.plot_id
                """
            ),
            params,
        ).mappings().all()
    return [alert_rule_payload(row) for row in rows]


@app.post("/api/alert-rules")
def save_alert_rule(
    request: AlertRuleRequest,
    user: dict[str, Any] = Depends(current_user),
) -> dict[str, Any]:
    numeric_plot_id = parse_resource_id(request.plotId, "P")
    # Farmers may configure the threshold of their own plots. Device binding
    # remains an administrator-only operation through ensure_plot_manage_access.
    plot_scope(numeric_plot_id, user)
    with engine.begin() as connection:
        connection.execute(
            text(
                """
                INSERT INTO threshold_config (
                    plot_id, soil_moisture_threshold, irrigation_length,
                    irrigation_interval, updated_by
                ) VALUES (
                    :plot_id, :threshold, :length, :interval, :updated_by
                )
                ON DUPLICATE KEY UPDATE
                    soil_moisture_threshold = VALUES(soil_moisture_threshold),
                    irrigation_length = VALUES(irrigation_length),
                    irrigation_interval = VALUES(irrigation_interval),
                    updated_by = VALUES(updated_by)
                """
            ),
            {
                "plot_id": numeric_plot_id,
                "threshold": request.soilMoistureThreshold,
                "length": request.irrigationLength,
                "interval": request.irrigationInterval,
                "updated_by": user["user_id"],
            },
        )
        evaluate_latest_moisture(
            connection,
            plot_id=numeric_plot_id,
            threshold=request.soilMoistureThreshold,
        )
        row = connection.execute(
            text(
                """
                SELECT config_id, plot_id, soil_moisture_threshold,
                       irrigation_length, irrigation_interval, updated_at
                FROM threshold_config WHERE plot_id = :plot_id
                """
            ),
            {"plot_id": numeric_plot_id},
        ).mappings().one()
    return alert_rule_payload(row)


@app.post("/api/chat/sessions/{session_id}/messages")
def send_chat_message(
    session_id: str,
    request: ChatMessageRequest,
    user: dict[str, Any] = Depends(current_user),
) -> dict[str, Any]:
    context: dict[str, Any] = {}
    if request.plotId is not None:
        numeric_plot_id = parse_resource_id(request.plotId, "P")
        plot_scope(numeric_plot_id, user)
        context = assistant_plot_context(numeric_plot_id)
    service_url = os.getenv("AI_SERVICE_URL", "").rstrip("/")
    if not service_url:
        raise HTTPException(status_code=503, detail="AI_SERVICE_UNAVAILABLE")
    payload = json.dumps(
        {
            "sessionId": session_id,
            "message": request.message,
            "plotId": request.plotId,
            "userId": business_id("U", user["user_id"]),
            "context": context,
        }
    ).encode("utf-8")
    target = f"{service_url}/chat/sessions/{quote(session_id, safe='')}/messages"
    try:
        with urlopen(
            Request(target, data=payload, headers={"Content-Type": "application/json"}),
            timeout=float(os.getenv("AI_SERVICE_TIMEOUT_SECONDS", "10")),
        ) as response:
            result = json.loads(response.read().decode("utf-8"))
    except (HTTPError, URLError, TimeoutError, json.JSONDecodeError, OSError) as error:
        raise HTTPException(status_code=503, detail="AI_SERVICE_UNAVAILABLE") from error
    if not isinstance(result, dict) or not isinstance(result.get("answer"), str):
        raise HTTPException(status_code=502, detail="AI_SERVICE_INVALID_RESPONSE")
    sources = result.get("sources", [])
    if not isinstance(sources, list) or not all(isinstance(source, str) for source in sources):
        raise HTTPException(status_code=502, detail="AI_SERVICE_INVALID_RESPONSE")
    return {"sessionId": session_id, "answer": result["answer"], "sources": sources}
